#!/usr/bin/env python3
"""Validate JSON/app/Excel synchronization and the curated source links."""
from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TARGET_IDS = {"S096", "S942", *(f"S{i:03d}" for i in range(950, 959))}
FIELDS = [
    "Source_ID", "Title", "Title_EN", "Date_Gregorian", "Date_Persian", "Year", "Language",
    "Source_Type", "Publisher", "Author", "Interviewee", "Person", "Position", "Organization",
    "Province", "Location", "Project", "Basin_Source", "Basin_Destination", "Topic",
    "Position_Stance", "Summary", "Key_Claim", "Quote", "URL", "Canonical_URL", "Archive_URL",
    "DOI", "Document_ID", "Reliability", "Primary_or_Secondary", "Full_Text", "PDF_Status",
    "PDF_URL", "Local_PDF", "SHA256", "Notes",
]


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def norm(value):
    return " ".join(clean(value).replace("‌", " ").replace("ي", "ی").replace("ك", "ک").lower().split())


def extract_fallback(app_text):
    prefix = "const OFFLINE_FALLBACK = "
    start = app_text.index(prefix) + len(prefix)
    end = app_text.index(";\n\nconst $", start)
    return json.loads(app_text[start:end])


def main():
    failures = []
    warnings = []
    data = json.loads((ROOT / "data.json").read_text(encoding="utf-8"))
    sources = data["sources"]
    by_id = {source["Source_ID"]: source for source in sources}

    expected_ids = [f"S{i:03d}" for i in range(1, 959)]
    actual_ids = [source["Source_ID"] for source in sources]
    if actual_ids != expected_ids:
        failures.append("Source IDs are not continuous S001-S958 in order")
    if len(sources) != 958 or data["meta"]["totalSources"] != 958 or data["meta"]["lastSourceId"] != "S958":
        failures.append("Meta/source count mismatch")
    timeline_ids = [row["Source_ID"] for row in data["timeline"]]
    if timeline_ids != expected_ids:
        failures.append("Timeline IDs/count/order mismatch")
    if set(by_id) != set(actual_ids):
        failures.append("Duplicate Source_ID in sources")

    app_data = extract_fallback((ROOT / "app.js").read_text(encoding="utf-8"))
    if app_data != data:
        failures.append("app.js OFFLINE_FALLBACK differs from data.json")

    wb = load_workbook(ROOT / "Data.xlsx", data_only=False, read_only=False)
    all_ws = wb["ALL_SOURCES"]
    headers = [clean(cell.value) for cell in all_ws[1]][: len(FIELDS)]
    if headers != FIELDS:
        failures.append("ALL_SOURCES headers differ from canonical fields")
    excel_ids = [clean(all_ws.cell(row, 1).value) for row in range(2, all_ws.max_row + 1)]
    if excel_ids != expected_ids:
        failures.append("Excel ALL_SOURCES IDs/count/order mismatch")
    excel_row = {clean(all_ws.cell(row, 1).value): row for row in range(2, all_ws.max_row + 1)}
    for source_id in sorted(TARGET_IDS):
        row = excel_row.get(source_id)
        if not row:
            failures.append(f"{source_id} absent from ALL_SOURCES")
            continue
        for col, field in enumerate(FIELDS, 1):
            if clean(all_ws.cell(row, col).value) != clean(by_id[source_id].get(field)):
                failures.append(f"Excel mismatch {source_id}.{field}")

    required_membership = {
        "ACADEMIC": {"S096", "S942", "S950", "S951", "S952", "S956"},
        "NEWS_MEDIA": {"S953", "S954", "S955", "S958"},
        "OFFICIAL_LEGAL": {"S936", "S937", "S945", "S957"},
        "PARLIAMENT": {"S936", "S945"},
        "PDF_INDEX": {"S096", "S950", "S951", "S952"},
    }
    sheet_membership = {}
    for sheet, required in required_membership.items():
        ids = {clean(wb[sheet].cell(row, 1).value) for row in range(2, wb[sheet].max_row + 1)}
        sheet_membership[sheet] = sorted(required & ids)
        missing = required - ids
        if missing:
            failures.append(f"Missing from {sheet}: {sorted(missing)}")

    forbidden = "10.1007/s10040-020-02213-y"
    if forbidden.lower() in (ROOT / "data.json").read_text(encoding="utf-8").lower():
        failures.append("Old out-of-scope S942 DOI remains in data.json")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            if any(forbidden.lower() in clean(cell.value).lower() for cell in row):
                failures.append(f"Old out-of-scope S942 DOI remains in Excel sheet {ws.title}")
                break

    for key in ("DOI", "Canonical_URL"):
        values = [norm(source.get(key)) for source in sources if source.get(key)]
        duplicates = sorted(value for value, count in Counter(values).items() if count > 1)
        # Existing collections can legitimately contain legacy duplicates, but none may involve this pass.
        for value in duplicates:
            ids = [source["Source_ID"] for source in sources if norm(source.get(key)) == value]
            if TARGET_IDS.intersection(ids):
                failures.append(f"Duplicate {key} involving curated rows: {ids}")
    titles = [norm(source.get("Title")) for source in sources if source.get("Title")]
    for title, count in Counter(titles).items():
        if count > 1:
            ids = [source["Source_ID"] for source in sources if norm(source.get("Title")) == title]
            if TARGET_IDS.intersection(ids):
                failures.append(f"Duplicate title involving curated rows: {ids}")

    link_results = []
    checked = set()
    for source_id in sorted(TARGET_IDS):
        source = by_id[source_id]
        for field in ("URL", "Archive_URL", "PDF_URL"):
            url = source.get(field)
            if not url or url in checked:
                continue
            checked.add(url)
            try:
                response = requests.get(url, timeout=35, allow_redirects=True, headers={"User-Agent": "Mozilla/5.0"})
                result = {
                    "Source_ID": source_id, "field": field, "url": url,
                    "status": response.status_code, "final_url": response.url,
                    "bytes": len(response.content), "content_type": response.headers.get("content-type", ""),
                }
                if field == "PDF_URL" and not (response.content[:4] == b"%PDF" or "pdf" in result["content_type"].lower()):
                    failures.append(f"PDF_URL is not a PDF: {source_id}")
                if response.status_code >= 400:
                    if source_id == "S957" and response.status_code == 403:
                        warnings.append("S957 Nezamat blocks scripted requests with 403; indexed web page is publicly discoverable")
                    elif source_id == "S096" and field == "Archive_URL" and response.status_code == 403:
                        warnings.append("S096 OuluREPO archive mirror blocks scripted requests with 403; DOI and publisher PDF both return 200")
                    else:
                        failures.append(f"HTTP {response.status_code}: {source_id} {field}")
                link_results.append(result)
            except Exception as exc:
                failures.append(f"Link error {source_id} {field}: {type(exc).__name__}: {exc}")

    report = {
        "date": "2026-09-13",
        "passed": not failures,
        "source_count": len(sources),
        "timeline_count": len(data["timeline"]),
        "pdf_index_count": len(data["pdfIndex"]),
        "excel_all_sources_rows": all_ws.max_row - 1,
        "target_ids": sorted(TARGET_IDS),
        "required_sheet_membership": sheet_membership,
        "link_checks": link_results,
        "warnings": warnings,
        "failures": failures,
    }
    (ROOT / "reports/deep_persian_challenges_validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(1 if failures else 0)


if __name__ == "__main__":
    main()
