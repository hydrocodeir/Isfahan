#!/usr/bin/env python3
"""Add the non-duplicate, basin-scoped records from the targeted document search.

This pass is deliberately small and auditable: every row below has a stable
publisher/archival URL (or a registered DOI), while the report also records
items that were found only as secondary references or whose primary legal text
could not be recovered.
"""
from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
TODAY = "2026-09-13"
TODAY_FA = "1405/06/22"
FIELDS = [
    "Source_ID", "Title", "Title_EN", "Date_Gregorian", "Date_Persian", "Year", "Language",
    "Source_Type", "Publisher", "Author", "Interviewee", "Person", "Position", "Organization",
    "Province", "Location", "Project", "Basin_Source", "Basin_Destination", "Topic",
    "Position_Stance", "Summary", "Key_Claim", "Quote", "URL", "Canonical_URL", "Archive_URL",
    "DOI", "Document_ID", "Reliability", "Primary_or_Secondary", "Full_Text", "PDF_Status",
    "PDF_URL", "Local_PDF", "SHA256", "Notes",
]

def rec(**kw):
    row = {f: None for f in FIELDS}
    row.update(kw)
    return row

ADDED = [
rec(Source_ID="S936", Title="قانون برنامه هفتم پیشرفت جمهوری اسلامی ایران (۱۴۰۷-۱۴۰۳) ـ بند «ج» ماده ۴۰", Date_Gregorian="2024-06-20", Year=1403, Language="فارسی", Source_Type="قانون بالادستی", Publisher="سامانه قوانین و مقررات کشور / متن قانون", Province="ملی", Project="انتقال بین‌حوضه‌ای / زاینده‌رود / کارون", Basin_Source="حوضه‌های آبریز داخلی ایران", Basin_Destination="حوضه‌های آبریز داخلی ایران", Topic="قانون‌گذاری، تخصیص آب، ممنوعیت انتقال غیرشرب", Position_Stance="تنظیمی", Summary="بند ج ماده ۴۰ انتقال آب بین شش حوضه آبریز داخلی برای مصارف غیرشرب را ممنوع می‌کند و انتقال شرب را به مراحل فنی، تصویب شورای عالی آب، نیازهای زیست‌محیطی و رعایت حقابه مبدأ مشروط می‌سازد.", Key_Claim="قاعده عام و به‌روز حقوقی برای ارزیابی طرح‌های انتقال کارون به زاینده‌رود.", URL="https://nezamat.ir/%D9%82%D8%A7%D9%86%D9%88%D9%86-%D8%A8%D8%B1%D9%86%D8%A7%D9%85%D9%87-%D9%87%D9%81%D8%AA%D9%85-%D9%BE%DB%8C%D8%B4%D8%B1%D9%81%D8%AA-%D8%AC%D9%85%D9%87%D9%88%D8%B1%DB%8C-%D8%A7%D8%B3%D9%84%D8%A7%D9%85/", Canonical_URL="https://nezamat.ir/%D9%82%D8%A7%D9%86%D9%88%D9%86-%D8%A8%D8%B1%D9%86%D8%A7%D9%85%D9%87-%D9%87%D9%81%D8%AA%D9%85-%D9%BE%DB%8C%D8%B4%D8%B1%D9%81%D8%AA-%D8%AC%D9%85%D9%87%D9%88%D8%B1%DB%8C-%D8%A7%D8%B3%D9%84%D8%A7%D9%85/", Archive_URL="https://nezaratmali.ir/wp-content/uploads/2024/07/GoB7.pdf", Document_ID="Program7-Art40-J", Reliability="A", Primary_or_Secondary="Primary legal text; publisher mirror", Full_Text="Yes", PDF_Status="Direct Link Available", PDF_URL="https://nezaratmali.ir/wp-content/uploads/2024/07/GoB7.pdf", Notes="متن بند ج ماده ۴۰ در PDF قانون/گزارش تفکیک وظایف دیوان محاسبات کنترل شد؛ پیوند HTML سامانه قوانین در برخی درخواست‌ها 403 می‌دهد.") ,
rec(Source_ID="S937", Title="Gavkhouni Lake and marshes of the lower Zaindeh Rud — Ramsar Information Sheet", Date_Gregorian="1997-09-15", Year=1997, Language="English", Source_Type="سند معاهده‌ای/برگه رامسر", Publisher="Ramsar Sites Information Service", Province="اصفهان", Project="زاینده‌رود / گاوخونی", Basin_Source="زاینده‌رود", Basin_Destination="گاوخونی", Topic="رامسر، تالاب، تعهد بین‌المللی، حقابه محیط‌زیست", Position_Stance="حفاظتی", Summary="برگه رسمی سایت رامسر شماره ۵۳ برای گاوخونی و نیزارهای پایین‌دست زاینده‌رود؛ تاریخ ثبت ۲۳ ژوئن ۱۹۷۵ و مساحت ثبت‌شده ۴۳٬۰۰۰ هکتار.", Key_Claim="کاهش جریان زاینده‌رود مستقیماً به یک تالاب دارای وضعیت بین‌المللی رامسر مربوط است.", URL="https://rsis.ramsar.org/RISapp/files/RISrep/IR53RIS.pdf", Canonical_URL="https://rsis.ramsar.org/RISapp/files/RISrep/IR53RIS.pdf", Document_ID="Ramsar-IR53", Reliability="A", Primary_or_Secondary="Primary treaty record", Full_Text="Yes", PDF_Status="Direct Link Available", PDF_URL="https://rsis.ramsar.org/RISapp/files/RISrep/IR53RIS.pdf", Notes="با فهرست رسمی Annotated List رامسر و صفحه RSIS سایت 53 تطبیق شد؛ منبع علمی S745 تکرار محسوب نمی‌شود چون این سند، رکورد معاهده‌ای است.") ,
rec(Source_ID="S938", Title="Water supply and demand forecasting in the Zayandeh Rud Basin, Iran", Title_EN="Water supply and demand forecasting in the Zayandeh Rud Basin, Iran", Date_Gregorian="2002", Year=2002, Language="English", Source_Type="گزارش فنی/هیدرولوژیک", Publisher="IAERI–EARC–IWMI", Author="H. R. Salemi; H. Murray-Rust", Province="اصفهان", Project="زاینده‌رود / بهشت‌آباد / کوهرنگ ۳", Basin_Source="کارون/کوهرنگ", Basin_Destination="زاینده‌رود", Topic="عرضه و تقاضا، سناریو، انتقال بین‌حوضه‌ای", Position_Stance="تحلیلی", Summary="گزارش پژوهشی ۱۳ IWMI نشان می‌دهد تکمیل انتقال سوم حدود ۲۸٪ به عرضه می‌افزاید، اما با رشد تقاضا تراز تا ۲۰۱۰ منفی و در ۲۰۲۰ بحرانی می‌شود؛ بخش کشاورزی نخستین محل فشار است.", Key_Claim="افزایش عرضه انتقالی بدون مهار تقاضا، کسری ساختاری حوضه را حل نمی‌کند.", URL="https://ageconsearch.umn.edu/record/158345/?ln=en", Canonical_URL="https://ageconsearch.umn.edu/record/158345/?ln=en", Archive_URL="https://ageconsearch.umn.edu/record/158345/files/H040773.pdf", DOI="10.22004/ag.econ.158345", Document_ID="IAERI-EARC-IWMI-RR13", Reliability="A", Primary_or_Secondary="Primary research report", Full_Text="Yes", PDF_Status="Direct Link Available", PDF_URL="https://ageconsearch.umn.edu/record/158345/files/H040773.pdf", Notes="رکورد، عنوان، نویسندگان، سال، شماره گزارش و DOI در AgeconSearch/IWMI تأیید شد؛ در بانک فعلی عنوان مشابه نبود.") ,
rec(Source_ID="S939", Title="Water Resources Development and Water Utilization in the Zayandeh Rud basin, Iran", Title_EN="Water Resources Development and Water Utilization in the Zayandeh Rud basin, Iran", Date_Gregorian="2002", Year=2002, Language="English", Source_Type="گزارش فنی/هیدرولوژیک", Publisher="IAERI–EARC–IWMI", Author="H. Murray-Rust; H. R. Salemi; P. Droogers", Province="اصفهان / چهارمحال‌وبختیاری / خوزستان", Project="زاینده‌رود / کوهرنگ", Basin_Source="کارون/کوهرنگ", Basin_Destination="زاینده‌رود", Topic="توسعه منابع، عرضه و تقاضا، خشکسالی، تخصیص", Position_Stance="تحلیلی/انتقادی", Summary="گزارش ۱۴ IWMI نشان می‌دهد دو انتقال و مخزن عرضه را تقریباً دو برابر کردند، اما ظرفیت برداشت هم هم‌زمان افزایش یافت؛ در کسری، جریان پایین‌دست و گاوخونی حذف می‌شود.", Key_Claim="اثر خالص توسعه سازه‌ای بدون مدیریت تقاضا، افزایش آسیب‌پذیری در خشکسالی است.", URL="https://archive.iwmi.org/assessment/files/word/ProjectDocuments/Zayandeh%20Rud/Zayandeh_14.PDF", Canonical_URL="https://archive.iwmi.org/assessment/files/word/ProjectDocuments/Zayandeh%20Rud/Zayandeh_14.PDF", Document_ID="IAERI-EARC-IWMI-RR14", Reliability="A", Primary_or_Secondary="Primary research report", Full_Text="Yes", PDF_Status="Direct Link Available", PDF_URL="https://archive.iwmi.org/assessment/files/word/ProjectDocuments/Zayandeh%20Rud/Zayandeh_14.PDF", Notes="رکورد رسمی آرشیو IWMI و فهرست پروژه Zayandeh Rud تطبیق شد؛ در بانک فعلی رکورد هم‌عنوان نبود.") ,
rec(Source_ID="S940", Title="The 1999–2001 drought in the Zayandeh Rud basin, Iran, and its impact on water allocation and agriculture", Title_EN="The 1999–2001 drought in the Zayandeh Rud basin, Iran, and its impact on water allocation and agriculture", Date_Gregorian="2004", Year=2004, Language="English", Source_Type="گزارش فنی/هیدرولوژیک", Publisher="IWMI / Comprehensive Assessment", Author="François Molle; Alireza Mamanpoush", Province="اصفهان", Project="زاینده‌رود / کوهرنگ", Basin_Source="کارون/کوهرنگ", Basin_Destination="زاینده‌رود", Topic="خشکسالی، مخزن چادگان، تخصیص، کشاورزی", Position_Stance="تحلیلی", Summary="تحلیل ۱۱ صفحه‌ای IWMI از خشکسالی ۱۹۹۹–۲۰۰۱؛ دو تونل کوهرنگ در سه سال ۱٫۱ میلیارد مترمکعب وارد مخزن کردند، اما عدم تطبیق رهاسازی ۱۹۹۹ بحران را تشدید و کشاورزی را فشرده کرد.", Key_Claim="انتقال، تاب‌آوری کوتاه‌مدت عرضه را بالا می‌برد اما مدیریت مخزن و تقاضا تعیین‌کننده شدت بحران است.", URL="https://archive.iwmi.org/assessment/files_new/research_projects/The%201999_2001%20crisis%20in%20Zayandeh%20Rud.pdf?redirected=yes", Canonical_URL="https://archive.iwmi.org/assessment/files_new/research_projects/The%201999_2001%20crisis%20in%20Zayandeh%20Rud.pdf?redirected=yes", Document_ID="IWMI-Drought-1999-2001", Reliability="A", Primary_or_Secondary="Primary research report", Full_Text="Yes", PDF_Status="Direct Link Available", PDF_URL="https://archive.iwmi.org/assessment/files_new/research_projects/The%201999_2001%20crisis%20in%20Zayandeh%20Rud.pdf?redirected=yes", Notes="PDF رسمی IWMI باز شد و متن/صفحات ۰–۹ برای ارقام تونل، رهاسازی و پیامد کشاورزی بررسی شد.") ,
rec(Source_ID="S941", Title="The Dangerous Condition of Ground during High Overburden Tunneling (A Case Study in Iran)", Title_EN="The Dangerous Condition of Ground during High Overburden Tunneling (A Case Study in Iran)", Date_Gregorian="2016", Year=2016, Language="English", Source_Type="مقاله زمین‌شناسی/ژئوتکنیک", Publisher="Periodica Polytechnica Civil Engineering", Author="Raheb Bagherpour; Mohammad Javad Rahimdel", Province="چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد", Basin_Source="بهشت‌آباد/کارون", Basin_Destination="زاینده‌رود", Topic="زمین‌شناسی، فشار زمین، تونل عمیق", Position_Stance="فنی/هشداردهنده", Summary="مطالعه ژئوتکنیکی موردی تونل انتقال بهشت‌آباد با طول حدود ۶۴٫۹ کیلومتر؛ خطرهای squeezing و rock burst و نقش پوشش سنگ و روباره را ارزیابی می‌کند.", Key_Claim="شرایط زمین‌شناسی و روباره زیاد، ریسک اجرایی و هزینه‌ای تونل را بالا می‌برد.", URL="https://www.pp.bme.hu/ci/article/download/7923/6826/16046", Canonical_URL="https://doi.org/10.3311/PPci.7923", DOI="10.3311/PPci.7923", Document_ID="DOI:10.3311/PPci.7923", Reliability="A", Primary_or_Secondary="Primary peer-reviewed article", Full_Text="Yes", PDF_Status="Direct Link Available", PDF_URL="https://www.pp.bme.hu/ci/article/download/7923/6826/16046", Notes="DOI در Crossref ثبت و عنوان/نویسندگان کنترل شد؛ عنوان و DOI در منابع موجود تکرار نشد.") ,
rec(Source_ID="S942", Title="Assessment of groundwater ingress to a partially pressurized water-conveyance tunnel using a conduit-flow process model: a case study in Iran", Title_EN="Assessment of groundwater ingress to a partially pressurized water-conveyance tunnel using a conduit-flow process model: a case study in Iran", Date_Gregorian="2020", Year=2020, Language="English", Source_Type="مقاله هیدروژئولوژی", Publisher="Hydrogeology Journal / Springer", Author="Hossein Gholizadeh; Ahmad Behrouj Peely; Bryan W. Karney; Ahmad Malekpour", Province="چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد", Basin_Source="بهشت‌آباد/کارون", Basin_Destination="زاینده‌رود", Topic="آب زیرزمینی، نفوذ به تونل، مدل‌سازی", Position_Stance="فنی", Summary="مدل فرایند جریان مجرا برای برآورد ورود آب زیرزمینی به تونل انتقال آب در ایران ارائه می‌کند و حساسیت شرایط فشار و تراوایی را می‌سنجد.", Key_Claim="نفوذ آب زیرزمینی، هم مسئله پایداری/ایمنی تونل و هم مسئله اثر بر آبخوان مبدأ است.", URL="https://doi.org/10.1007/s10040-020-02213-y", Canonical_URL="https://doi.org/10.1007/s10040-020-02213-y", DOI="10.1007/s10040-020-02213-y", Document_ID="DOI:10.1007/s10040-020-02213-y", Reliability="A", Primary_or_Secondary="Primary peer-reviewed article", Full_Text="Unknown", PDF_Status="Landing Page Only", Notes="DOI و کتابشناسی در Crossref و Hydrogeology Journal/OUCI تأیید شد؛ نسخه فعلی بانک تکرار DOI/عنوان نداشت.") ,
rec(Source_ID="S943", Title="Evaluation of required thrust force based on advance rates in shielded TBMs under squeezing conditions", Title_EN="Evaluation of required thrust force based on advance rates in shielded TBMs under squeezing conditions", Date_Gregorian="2019", Year=2019, Language="English", Source_Type="مقاله ژئوتکنیک/تونل", Publisher="Journal of Geophysics and Engineering / Oxford Academic", Author="Danial Mohammadzamani; Saeed Mahdevari; Raheb Bagherpour", Province="چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد", Basin_Source="بهشت‌آباد/کارون", Basin_Destination="زاینده‌رود", Topic="TBM، فشار رانش، squeezing، امکان‌سنجی", Position_Stance="فنی", Summary="مدل عددی نیروی رانش TBM در شرایط squeezing را با تونل بهشت‌آباد به‌عنوان مطالعه موردی بررسی می‌کند و اثر نرخ پیشروی، over-boring و توقف دستگاه را می‌سنجد.", Key_Claim="رفتار خزشی و توقف TBM می‌تواند نیروی رانش، زمان اجرا و ریسک هزینه را به‌طور معنادار تغییر دهد.", URL="https://academic.oup.com/jge/article/16/5/842/5556564", Canonical_URL="https://doi.org/10.1093/jge/gxz050", DOI="10.1093/jge/gxz050", Document_ID="DOI:10.1093/jge/gxz050", Reliability="A", Primary_or_Secondary="Primary peer-reviewed article", Full_Text="Yes", PDF_Status="Landing Page Only", Notes="عنوان، DOI، نویسندگان و چکیده در Oxford Academic و Crossref بررسی شد؛ تکرار در بانک وجود نداشت.") ,
rec(Source_ID="S944", Title="طرحی که نه ارزیابی محیط زیستی دارد و نه تعریف تأمین آب شرب", Date_Gregorian="2020-06-20", Year=1399, Language="فارسی", Source_Type="گزارش بازرسی/تحقیقی", Publisher="خبرگزاری مهر", Province="چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد / کوهرنگ ۳", Basin_Source="کارون/بهشت‌آباد", Basin_Destination="زاینده‌رود", Topic="سازمان بازرسی، ارزیابی محیط‌زیستی، آب شرب", Position_Stance="انتقادی", Summary="گزارش تحقیقی مهر به نامه شهریور ۱۳۹۰ سازمان بازرسی کل کشور استناد می‌کند که با مدیریت مصرف، تخصیص‌های اصفهان را بدون کوهرنگ ۳ و بهشت‌آباد برای افق بررسی‌شده کافی دانسته و به خلأ ارزیابی محیط‌زیستی و تعریف نیاز شرب اشاره دارد.", Key_Claim="سند بازرسیِ نقل‌شده، ضرورت انتقال را مشروط به مدیریت مصرف و ارزیابی مستقل می‌داند.", URL="https://www.mehrnews.com/news/4954007/%D8%B7%D8%B1%D8%AD%DB%8C-%DA%A9%D9%87-%D9%86%D9%87-%D8%A7%D8%B1%D8%B2%DB%8C%D8%A7%D8%A8%DB%8C-%D9%85%D8%AD%DB%8C%D8%B7-%D8%B2%DB%8C%D8%B3%D8%AA%DB%8C-%D8%AF%D8%A7%D8%B1%D8%AF-%D9%88-%D9%86%D9%87-%D8%AA%D8%B9%D8%B1%DB%8C%D9%81-%D8%AA%D8%A7%D9%85%DB%8C%D9%86-%D8%A2%D8%A8-%D8%B4%D8%B1%D8%A8", Canonical_URL="https://www.mehrnews.com/news/4954007/%D8%B7%D8%B1%D8%AD%DB%8C-%DA%A9%D9%87-%D9%86%D9%87-%D8%A7%D8%B1%D8%B2%DB%8C%D8%A7%D8%A8%DB%8C-%D9%85%D8%AD%DB%8C%D8%B7-%D8%B2%DB%8C%D8%B3%D8%AA%DB%8C-%D8%AF%D8%A7%D8%B1%D8%AF-%D9%88-%D9%86%D9%87-%D8%AA%D8%B9%D8%B1%DB%8C%D9%81-%D8%AA%D8%A7%D9%85%DB%8C%D9%86-%D8%A2%D8%A8-%D8%B4%D8%B1%D8%A8", Reliability="B", Primary_or_Secondary="Secondary investigative report quoting Inspection Organization material", Full_Text="Yes", PDF_Status="Landing Page Only", Notes="متن نامه سازمان بازرسی به‌صورت سند پیوست رسمی بازیابی نشد؛ خبر مهر به‌عنوان منبع ثانویه و با این محدودیت ثبت شد.") ,
rec(Source_ID="S945", Title="درخواست برخورد قضایی با صادرکنندگان مجوز سدسازی و انتقال آب", Date_Gregorian="2019-08-17", Year=1398, Language="فارسی", Source_Type="نامه نمایندگان/درخواست قضایی", Publisher="بازنشر نامه نمایندگان در WaterResources", Province="چهارمحال‌وبختیاری / خوزستان / اصفهان", Project="بهشت‌آباد / کوهرنگ ۳ / خرسان ۳", Basin_Source="کارون", Basin_Destination="زاینده‌رود / فلات مرکزی", Topic="شکایت، دادستانی، سازمان حفاظت محیط‌زیست، اصل ۵۰", Position_Stance="مخالف", Summary="بازنشر متن نامه تعدادی از نمایندگان به رئیس قوه قضاییه برای رسیدگی به پروژه‌های بهشت‌آباد، خرسان و گلاب و اشاره به شکایت تشکل‌های محیط‌زیستی از بانیان کوهرنگ ۳.", Key_Claim="وجود درخواست رسمی نمایندگان برای ورود دادستانی گزارش شده، اما متن ثبت‌شده فعلاً بازنشر ثانویه است و شماره پرونده/کیفرخواست بازیابی نشد.", URL="https://waterresources.blogfa.com/post/1131", Canonical_URL="https://waterresources.blogfa.com/post/1131", Reliability="C", Primary_or_Secondary="Secondary reproduction of a reported parliamentary letter", Full_Text="Yes", PDF_Status="Landing Page Only", Notes="برای جلوگیری از ادعای بیش از سند، به‌عنوان متن اصلی شکایت/کیفرخواست ثبت نشده است؛ اصل نامه و نتیجه رسیدگی در گزارش شکاف‌ها آمده است.") ,
rec(Source_ID="S946", Title="جعفر شریف امامی، نوار ۳ — تاریخچه طرح کوهرنگ", Title_EN="Jafar Sharif-Emami, Tape 3 — Kouhrang project history", Date_Gregorian="1982-05-13", Year=1361, Language="فارسی", Source_Type="مصاحبه تاریخی/تاریخ شفاهی", Publisher="Iran Oral History / Harvard Iranian Oral History Project", Interviewee="جعفر شریف امامی", Person="جعفر شریف امامی", Position="مدیر بنگاه آبیاری؛ نخست‌وزیر و رئیس سنا", Organization="Harvard Center for Middle Eastern Studies", Province="چهارمحال‌وبختیاری / اصفهان", Project="کوهرنگ ۱", Basin_Source="سرشاخه کارون/کوهرنگ", Basin_Destination="زاینده‌رود", Topic="تاریخچه تصمیم‌گیری، سد و تونل، شاهد عینی", Position_Stance="تاریخی/روایتی", Summary="روایت دست‌اول شریف امامی از طرح انتقال شاخه‌ای از کارون از کوهرنگ به جلگه اصفهان، سابقه صفوی، اختلافات اجرایی و تشکیل شرکت طرح.", Key_Claim="این مصاحبه تاریخ شکل‌گیری نهادی و فنی انتقال کوهرنگ ۱ را از زبان یکی از مدیران محوری ثبت می‌کند.", URL="https://iranhistory.net/sharif-emami3/", Canonical_URL="https://iranhistory.net/sharif-emami3/", Archive_URL="https://cmes.fas.harvard.edu/projects/iohp", Document_ID="IOHP-Sharif-Emami-Tape3", Reliability="A", Primary_or_Secondary="Primary oral history", Full_Text="Yes", Notes="صفحه نوار ۳ و مجموعه مادر هاروارد کنترل شد؛ در دیتاست فعلی رکورد شریف امامی نبود.") ,
rec(Source_ID="S947", Title="عیسی کلانتری: سرانجام انتقال آب به بهشت‌آباد چه شد؟", Date_Gregorian="2020-10-03", Year=1399, Language="فارسی", Source_Type="مصاحبه مستقیم", Publisher="آرمان ملی / بازنشر Sarpoosh", Interviewee="عیسی کلانتری", Person="عیسی کلانتری", Position="رئیس سازمان حفاظت محیط‌زیست", Organization="سازمان حفاظت محیط‌زیست", Province="چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد", Basin_Source="کارون/بهشت‌آباد", Basin_Destination="زاینده‌رود", Topic="مجوز محیط‌زیستی، توقف پروژه، سیاست دولت", Position_Stance="مشروط/انتقادی", Summary="کلانتری در مصاحبه می‌گوید مجوز محیط‌زیستی بهشت‌آباد صادر نشده و وزارت نیرو پروژه را متوقف کرده تا مجوز اخذ شود.", Key_Claim="موضع رسمی رئیس وقت سازمان حفاظت محیط‌زیست، اخذ مجوز را پیش‌شرط ادامه بهشت‌آباد معرفی می‌کند.", URL="https://www.sarpoosh.com/politics/domestic-policy/domestic-policy990700572.html", Canonical_URL="https://www.sarpoosh.com/politics/domestic-policy/domestic-policy990700572.html", Archive_URL="https://www.pishkhan.com/Archive/1399/07/13990712/ArmanMeli6511411097109495550118356.pdf", Reliability="B", Primary_or_Secondary="Primary interview reproduced by secondary host", Full_Text="Yes", Notes="متن کامل پرسش و پاسخ در خطوط مربوط به بهشت‌آباد قابل بازیابی است؛ عنوان/مصاحبه در بانک موجود نبود.") ,
rec(Source_ID="S948", Title="بحران آب و سدسازی در ایران در گفتگو با نیک‌آهنگ کوثر", Title_EN="Iran's water crisis and dam building — interview with Nikahang Kowsar", Date_Gregorian="2024-03-07", Year=1402, Language="فارسی", Source_Type="مصاحبه مستقیم", Publisher="Kayhan London", Interviewee="نیک‌آهنگ کوثر", Person="نیک‌آهنگ کوثر", Position="زمین‌شناس و تحلیلگر آب", Province="خوزستان / چهارمحال‌وبختیاری / اصفهان", Project="کوهرنگ ۲ / کوهرنگ ۳ / بهشت‌آباد / چشمه‌لنگان / خدنگستان", Basin_Source="کارون/دز", Basin_Destination="زاینده‌رود", Topic="تاریخچه انتقال، سدسازی، پیامد پایین‌دست", Position_Stance="انتقادی", Summary="گفتگوی تفصیلی درباره انتقال‌های کوهرنگ، بهشت‌آباد، خدنگستان و چشمه‌لنگان و پیامدهای کاهش دبی و شوری کارون و خوزستان.", Key_Claim="مصاحبه، تداوم تاریخی انتقال از کارون/دز به زاینده‌رود و تعارض آن با پایداری پایین‌دست را از منظر زمین‌شناس بررسی می‌کند.", URL="https://kayhan.london/1402/12/17/343888/", Canonical_URL="https://kayhan.london/1402/12/17/343888/", Reliability="B", Primary_or_Secondary="Primary interview", Full_Text="Yes", Notes="مصاحبه نیک‌آهنگ کوثر با رکورد تحلیلی/یادداشت S470 تکرار نیست؛ قالب و محتوای مصاحبه مستقل است.") ,
rec(Source_ID="S949", Title="مهدی قمشی: خوزستان آب اضافه‌ای برای انتقال به دیگر استان‌ها ندارد", Date_Gregorian="2008-07-17", Year=1387, Language="فارسی", Source_Type="مصاحبه مستقیم", Publisher="گزارش مصاحبه بازنشرشده در Dezhpol News", Interviewee="مهدی قمشی", Person="مهدی قمشی", Position="عضو هیأت علمی دانشگاه شهید چمران اهواز", Organization="دانشگاه شهید چمران اهواز", Province="خوزستان / چهارمحال‌وبختیاری / اصفهان", Project="انتقال کارون / بهشت‌آباد / کوهرنگ ۳", Basin_Source="کارون", Basin_Destination="زاینده‌رود", Topic="آب مازاد، کشاورزی، شوری، تالاب‌ها", Position_Stance="مخالف/کارشناسی", Summary="گفتگوی تفصیلی با قمشی درباره نبود آب مازاد خوزستان، پیامدهای اشتغال و کشاورزی، افزایش شوری کارون و اثر انتقال بر تالاب‌ها و آبخوان چهارمحال‌وبختیاری.", Key_Claim="دیدگاه کارشناسی مخالف انتقال، بر اولویت نیازهای داخل حوضه کارون و مدیریت تقاضا تأکید دارد.", URL="https://dezhpol-news.blogfa.com/post/2261", Canonical_URL="https://dezhpol-news.blogfa.com/post/2261", Reliability="C", Primary_or_Secondary="Secondary reproduction of direct interview", Full_Text="Yes", Notes="منبع رسانه‌ای/وبلاگی است؛ به‌دلیل نبود صفحه اصلی روزنامه، اعتبار C داده شد و برای استفاده تحلیلی باید با اصل مصاحبه تطبیق شود.") ,
]

DUPLICATES = [
    {"candidate": "گزارش مرکز پژوهش‌های مجلس، شماره ۸۹۳۵", "existing": "S252", "reason": "رکورد موجود؛ فقط پیوند رسمی rc.majlis.ir همچنان برای بازیابی اصل سند به‌عنوان شکاف نگه داشته شد."},
    {"candidate": "گزارش مرکز پژوهش‌های مجلس، شماره مسلسل ۱۲۴۹۲", "existing": "S175", "reason": "رکورد موجود؛ عنوان و موضوع تکراری است."},
    {"candidate": "رأی ۳۷۲۸۹۴ دیوان عدالت اداری", "existing": "S228", "reason": "رأی قضایی هدف از قبل در بانک موجود است."},
    {"candidate": "مصاحبه حمید چیت‌چیان", "existing": "S300/S375", "reason": "مصاحبه/اظهارنظر وزیر نیرو پیش‌تر ثبت شده است."},
    {"candidate": "مصاحبه ناصر کرمی", "existing": "S434", "reason": "مصاحبه مستقیم درباره انتقال کارون به زاینده‌رود پیش‌تر ثبت شده است."},
    {"candidate": "کارون شدن زاینده‌رود / نیک‌آهنگ کوثر", "existing": "S470", "reason": "یادداشت تحلیلی موجود است؛ مصاحبه Kayhan London (S948) قالب و متن مستقل دارد."},
]

GAPS = [
    {"موضوع": "اصل ۹۰ مجلس", "وضعیت": "پیدا نشد", "توضیح": "در جست‌وجوی مستقیم ICANA و rc.majlis.ir متن گزارش/شماره پرونده اصل ۹۰ درباره زاینده‌رود، کارون، بهشت‌آباد یا کوهرنگ ۳ بازیابی نشد؛ خبرها و نامه‌های نمایندگان جایگزین سند اصل ۹۰ محسوب نشدند."},
    {"موضوع": "کیفرخواست یا رأی کیفری مجریان انتقال", "وضعیت": "پیدا نشد", "توضیح": "برای پروژه‌های هدف، کیفرخواست یا دادنامه کیفری قابل احراز در پایگاه رسمی عمومی پیدا نشد؛ ادعاهای شبکه‌های اجتماعی کنار گذاشته شد."},
    {"موضوع": "اصل نامه سازمان بازرسی کل کشور", "وضعیت": "فقط بازتاب ثانویه", "توضیح": "مهر و چند بازنشر به نامه شهریور ۱۳۹۰ استناد می‌کنند، اما PDF/شماره ثبت رسمی نامه بازیابی نشد؛ S944 با درجه B و این محدودیت ثبت شد."},
    {"موضوع": "اصل نامه قضایی نمایندگان ۱۳۹۸", "وضعیت": "فقط بازنشر ثانویه", "توضیح": "متن نامه و اشاره به شکایت تشکل‌ها در S945 موجود است، اما اصل مکاتبه، شماره دبیرخانه و نتیجه رسیدگی دادستانی پیدا نشد."},
    {"موضوع": "اصل صورت‌جلسه ۹ بندی شورای عالی آب", "وضعیت": "بازتاب خبری موجود", "توضیح": "S226 متن بازتاب‌یافته را دارد؛ نسخه اصل صورت‌جلسه هنوز بازیابی نشده است."},
]

def norm(s):
    return " ".join((s or "").replace("‌", " ").lower().split())

def main():
    data = json.loads((ROOT / "data.json").read_text(encoding="utf-8"))
    existing_ids = {x["Source_ID"] for x in data["sources"]}
    if any(x["Source_ID"] in existing_ids for x in ADDED):
        raise SystemExit("One or more targeted records already exist; refusing to duplicate.")
    # Guard against title/DOI/URL collisions before mutation.
    existing_text = "\n".join(norm(" ".join(str(x.get(k) or "") for k in ("Title", "Title_EN", "DOI", "URL"))) for x in data["sources"])
    for x in ADDED:
        keys = [norm(x.get("DOI")), norm(x.get("Title")), norm(x.get("Canonical_URL"))]
        if any(k and k in existing_text for k in keys):
            raise SystemExit(f"Possible duplicate detected: {x['Source_ID']} {x['Title']}")
    data["sources"].extend(ADDED)
    data["sources"].sort(key=lambda x: int(x["Source_ID"][1:]))
    data["meta"]["generatedAt"] = TODAY
    data["meta"]["totalSources"] = len(data["sources"])
    data["meta"]["lastSourceId"] = ADDED[-1]["Source_ID"]
    dash = data["meta"]["dashboard"]
    dash["Total sources"] = len(data["sources"]); dash["Last Source ID"] = ADDED[-1]["Source_ID"]
    for grade in "ABCD": dash[f"Reliability {grade}"] = sum(x.get("Reliability") == grade for x in data["sources"])
    dash["Persian-language sources"] = sum("فارسی" in (x.get("Language") or "") for x in data["sources"])
    dash["English-language sources"] = sum("English" in (x.get("Language") or "") or "انگلیسی" in (x.get("Language") or "") for x in data["sources"])
    dash["Sources with direct PDF links"] = sum(x.get("PDF_Status") == "Direct Link Available" for x in data["sources"])
    # Keep all collection views synchronized.
    for name in data["collections"]:
        data["collections"][name] = [i for i in data["collections"][name] if i in existing_ids]
    data["collections"]["legalIds"] += ["S936", "S937", "S945"]
    data["collections"]["academicIds"] += ["S938", "S939", "S940", "S941", "S942", "S943"]
    data["collections"]["parliamentIds"] += ["S945"]
    data["collections"]["newsIds"] += ["S944", "S946", "S947", "S948", "S949"]
    # Timeline/PDF index are denormalized views in data.json.
    for x in ADDED:
        data["timeline"].append({k: x.get(k) for k in ("Source_ID", "Date_Gregorian", "Date_Persian", "Year", "Project", "Title", "Source_Type", "Publisher", "Reliability", "URL")})
        if x.get("PDF_URL"):
            data["pdfIndex"].append({k: x.get(k) for k in ("Source_ID", "Title", "Publisher", "PDF_Status", "PDF_URL", "Local_PDF", "SHA256", "Reliability")})
    data["searchLog"].append({"Query":"جست‌وجوی هدفمند اسناد حقوقی، فنی، بازرسی، قضایی و مصاحبه‌ای زاینده‌رود/کارون","Language":"فارسی/English","Database / Search Engine":"Ramsar RSIS, IWMI archive, AgeconSearch, Crossref, Oxford, Springer, Mehr, Iranian Oral History/Harvard, media archives","Project":"زاینده‌رود / کارون / بهشت‌آباد / کوهرنگ","Source Type":"اسناد و مصاحبه‌های هدفمند","Search Date":TODAY,"Number of useful results":len(ADDED),"Notes":"تطبیق عنوان، DOI، URL و رکوردهای موجود؛ تکراری‌ها و شکاف‌های سندی جداگانه گزارش شدند."})
    data["changeLog"].append({"Change_ID":"C035","Date":TODAY,"Action":"TARGETED DOCUMENT SEARCH","Source_ID / Item":"S936–S949","Reason":"Add non-duplicate legal, treaty, technical, inspection, judicial-letter and interview records scoped to Iran/Zayandeh-Rud/Karun","Details":"Added 14 validated records; recorded 6 duplicates skipped and 5 documentary gaps in reports/targeted_document_search.*."})
    # Recompute project counts from the existing project labels.
    for p in data["projects"]:
        p["Source_Count"] = sum(p["Project"] in (x.get("Project") or "") for x in data["sources"])
    (ROOT / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    # Update offline fallback by replacing the single JSON assignment.
    app = (ROOT / "app.js").read_text(encoding="utf-8")
    marker = "const OFFLINE_FALLBACK = "
    start = app.index(marker) + len(marker); end = app.index(";\n\nconst $", start)
    (ROOT / "app.js").write_text(app[:start] + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + app[end:], encoding="utf-8")
    update_xlsx(data)
    write_reports()

def append(ws, vals):
    ws.append(vals)

def update_xlsx(data):
    path = ROOT / "Data.xlsx"; wb = load_workbook(path)
    for x in ADDED:
        vals = [x.get(f) for f in FIELDS]
        for s in ("ALL_SOURCES", "NEW_SOURCES"):
            append(wb[s], vals)
        if x["Source_Type"].startswith("مقاله") or x["Source_Type"].startswith("گزارش فنی") or x["Source_Type"].startswith("مقاله زمین") or x["Source_Type"].startswith("مقاله هیدرو") or x["Source_Type"].startswith("مقاله ژئو"):
            append(wb["ACADEMIC"], vals)
        if "نامه نمایندگان" in x["Source_Type"] or "قانون" in x["Source_Type"]:
            append(wb["PARLIAMENT"], vals)
        if x["Source_Type"].startswith("مصاحبه") or "گزارش بازرسی" in x["Source_Type"] or "مصاحبه تاریخی" in x["Source_Type"]:
            append(wb["NEWS_MEDIA"], vals)
        append(wb["TIMELINE"], [x.get(k) for k in ("Source_ID","Date_Gregorian","Date_Persian","Project","Title","Source_Type","Publisher","Reliability","URL")])
        if x.get("PDF_URL"):
            append(wb["PDF_INDEX"], [x.get(k) for k in ("Source_ID","Title","Publisher","PDF_Status","PDF_URL","Local_PDF","SHA256","Reliability")])
    wb["README"]["A1"] = "Water Transfer Isfahan — TARGETED DOCUMENT SEARCH"
    wb["README"]["B2"] = len(data["sources"]); wb["README"]["B3"] = f"S001–{data['meta']['lastSourceId']} (continuous, no gaps)"; wb["README"]["B4"] = f"Through {TODAY}"
    # Dashboard labels are stable; update values by label.
    labels = {wb["DASHBOARD"].cell(r,1).value:r for r in range(2, wb["DASHBOARD"].max_row+1)}
    for k,v in data["meta"]["dashboard"].items():
        if k in labels: wb["DASHBOARD"].cell(labels[k],2).value = v
    wb["SEARCH_LOG"].append(["جست‌وجوی هدفمند اسناد حقوقی، فنی، بازرسی، قضایی و مصاحبه‌ای زاینده‌رود/کارون","فارسی/English","Ramsar, IWMI, AgeconSearch, Crossref, Springer, Oxford, Mehr, Harvard Oral History, media archives","زاینده‌رود / کارون / بهشت‌آباد / کوهرنگ","اسناد و مصاحبه‌های هدفمند",TODAY,len(ADDED),"تطبیق عنوان، DOI، URL و رکوردهای موجود؛ تکراری‌ها و شکاف‌ها جداگانه گزارش شدند."])
    wb["CHANGE_LOG"].append(["C035",TODAY,"TARGETED DOCUMENT SEARCH","S936–S949","Add non-duplicate scoped records","Added 14 validated records; duplicate and gap reports written."])
    if "AllSourcesClean" in wb["ALL_SOURCES"].tables: wb["ALL_SOURCES"].tables["AllSourcesClean"].ref=f"A1:AK{wb['ALL_SOURCES'].max_row}"
    wb.save(path)

def write_reports():
    out = {"date":TODAY,"added":[x for x in ADDED],"duplicates_skipped":DUPLICATES,"gaps":GAPS}
    (ROOT/"reports/targeted_document_search.json").write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    wb=Workbook(); ws=wb.active; ws.title="Summary"; ws.sheet_view.rightToLeft=True
    ws.append(["Metric","Value"]); ws.append(["Search date",TODAY_FA]); ws.append(["Added",len(ADDED)]); ws.append(["Duplicates skipped",len(DUPLICATES)]); ws.append(["Documentary gaps",len(GAPS)])
    for title, rows, headers in [("Added",ADDED,FIELDS),("Duplicates",DUPLICATES,["candidate","existing","reason"]),("Gaps",GAPS,["موضوع","وضعیت","توضیح"])]:
        sh=wb.create_sheet(title); sh.sheet_view.rightToLeft=True; sh.append(headers)
        for row in rows: sh.append([row.get(h) for h in headers])
        sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
        for c in sh[1]: c.font=Font(bold=True)
    wb.save(ROOT/"reports/targeted_document_search.xlsx")

if __name__ == "__main__": main()
