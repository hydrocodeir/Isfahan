#!/usr/bin/env python3
"""Merge verified deep-literature results into JSON, Excel and offline UI."""

from __future__ import annotations

import html
import json
import re
from copy import copy
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from deep_literature_search import (
    REPORTS, ROOT, classify_scope, deduplicate, existing_match, normalize, save_json,
)
from build_validated_outputs import SOURCE_FIELDS, append_styled


TODAY = "2026-09-13"
TODAY_FA = "1405/06/22"


def candidate_key(item: dict) -> str:
    return normalize(item.get("doi")) or "title:" + normalize(item.get("title"))


def combine_candidates() -> tuple[list[dict], list[dict]]:
    main = json.loads((REPORTS / "deep_literature_unvalidated.json").read_text(encoding="utf-8"))["candidates"]
    supplement_payload = json.loads((REPORTS / "crossref_abstract_search.json").read_text(encoding="utf-8"))
    supplement = supplement_payload["records"]

    # Re-evaluate relevance with current rules. For bibliographic discovery,
    # matching in the abstract is accepted and recorded explicitly.
    retained = []
    for item in main + supplement:
        haystack = item.get("title") or ""
        if item.get("source_database") == "Crossref bibliographic":
            haystack += " " + (item.get("abstract") or "")
        scope, hits = classify_scope(haystack)
        if scope:
            item["scope"], item["term_hits"] = scope, hits
            retained.append(item)
    return deduplicate(retained), supplement_payload.get("log") or []


def project_for(item: dict) -> tuple[str, str | None, str | None, str | None]:
    text = normalize(" ".join((item.get("title") or "", item.get("abstract") or "")))
    projects = []
    mappings = [
        ("کوهرنگ ۳", ("koohrang iii", "kouhrang iii", "kuhrang iii", "koohrang 3", "kouhrang 3", "kuhrang 3", "کوهرنگ ۳", "کوهرنگ 3")),
        ("کوهرنگ", ("koohrang", "kouhrang", "kuhrang", "کوهرنگ")),
        ("بهشت‌آباد", ("behesht abad", "beheshtabad", "بهشت آباد")),
        ("گلاب", ("golab", "گلاب")),
        ("چشمه‌لنگان", ("cheshmeh langan", "چشمه لنگان")),
        ("خرسان ۳", ("khersan 3", "khersan iii", "خرسان ۳", "خرسان 3")),
        ("ماندگان", ("mandegan", "ماندگان")),
        ("بن-بروجن", ("ben borujen", "بن بروجن")),
        ("زاینده‌رود", ("zayandeh rud", "zayandeh rood", "zayanderud", "zayandehrud", "زاینده رود")),
        ("حوضه کارون", ("karun", "karoun", "کارون")),
    ]
    for label, aliases in mappings:
        if any(alias in text for alias in aliases):
            projects.append(label)
    if not projects:
        projects.append("انتقال بین‌حوضه‌ای")
    project = " / ".join(dict.fromkeys(projects))
    if "حوضه کارون" in projects and "زاینده‌رود" in projects:
        basin_source, basin_dest = "حوضه کارون", "حوضه زاینده‌رود"
    elif "حوضه کارون" in projects:
        basin_source, basin_dest = "حوضه کارون", None
    elif "زاینده‌رود" in projects:
        basin_source, basin_dest = None, "حوضه زاینده‌رود"
    else:
        basin_source = basin_dest = None
    topic = "ادبیات جهانی انتقال آب بین‌حوضه‌ای" if item.get("scope", "").startswith("Global") else "پژوهش علمی حوضه‌ها و طرح‌های مرتبط"
    return project, basin_source, basin_dest, topic


def clean_abstract(value: str | None) -> str | None:
    if not value:
        return None
    text = html.unescape(re.sub(r"<[^>]+>", " ", value))
    text = " ".join(text.split())
    return text[:1800] if text else None


def language_for(item: dict) -> str:
    value = (item.get("language") or "").lower()
    title = item.get("title") or ""
    if value in {"fa", "fas", "per", "persian"}:
        return "فارسی"
    if value in {"zh", "zho", "chi"}:
        return "چینی"
    if value in {"pt", "por"}:
        return "پرتغالی"
    return "فارسی" if re.search(r"[\u0600-\u06ff]", title) else "English"


def record_from(item: dict, source_id: str) -> dict:
    title = " ".join((item.get("title") or "").split())
    lang = language_for(item)
    project, basin_source, basin_dest, topic = project_for(item)
    doi = item.get("doi")
    canonical = f"https://doi.org/{doi}" if doi else item.get("landing_url")
    pdf = item.get("pdf_url")
    summary = clean_abstract(item.get("abstract"))
    databases = ", ".join(item.get("discovered_in") or [item.get("source_database")])
    biblio = item.get("biblio") or {}
    notes = [f"کشف و اعتبارسنجی در مرور عمیق ادبیات {TODAY_FA}", f"نمایه‌ها: {databases}"]
    if item.get("match_location"):
        notes.append("عبارت مرتبط در عنوان/چکیده/متادیتای کتاب‌شناختی یافت شد")
    volume_issue = ", ".join(filter(None, (
        f"Vol. {biblio.get('volume')}" if biblio.get("volume") else None,
        f"Issue {biblio.get('issue')}" if biblio.get("issue") else None,
        f"pp. {biblio.get('first_page')}-{biblio.get('last_page')}" if biblio.get("first_page") else None,
    )))
    if volume_issue:
        notes.append(volume_issue)

    record = {field: None for field in SOURCE_FIELDS}
    record.update({
        "Source_ID": source_id,
        "Title": title,
        "Title_EN": title if not re.search(r"[\u0600-\u06ff]", title) else None,
        "Date_Gregorian": item.get("publication_date"),
        "Year": item.get("year"),
        "Language": lang,
        "Source_Type": "مقاله علمی",
        "Publisher": item.get("journal") or item.get("publisher"),
        "Author": "، ".join(item.get("authors") or []) or None,
        "Organization": item.get("publisher"),
        "Province": "اصفهان / چهارمحال‌وبختیاری / خوزستان" if item.get("scope") == "Iran basins/projects" else None,
        "Project": project,
        "Basin_Source": basin_source,
        "Basin_Destination": basin_dest,
        "Topic": topic,
        "Position_Stance": "علمی / تحلیلی",
        "Summary": summary or "مقاله علمی مرتبط که وجود و متادیتای آن در نمایه دانشگاهی مستقل تأیید شده است.",
        "URL": canonical,
        "Canonical_URL": canonical,
        "DOI": doi,
        "Document_ID": f"DOI:{doi}" if doi else f"DOAJ:{item.get('source_id')}",
        "Reliability": "A",
        "Primary_or_Secondary": "Primary",
        "Full_Text": "Yes" if pdf else "Unknown",
        "PDF_Status": "Direct Link Available" if pdf else "Landing Page Only",
        "PDF_URL": pdf,
        "Notes": "؛ ".join(notes) + ".",
    })
    return record


def update_meta(data: dict, records: list[dict]) -> None:
    data["meta"]["generatedAt"] = TODAY
    data["meta"]["totalSources"] = len(data["sources"])
    data["meta"]["lastSourceId"] = data["sources"][-1]["Source_ID"]
    dashboard = data["meta"]["dashboard"]
    dashboard["Total sources"] = len(data["sources"])
    dashboard["Persian-language sources"] = sum("فارسی" in (x.get("Language") or "") for x in data["sources"])
    dashboard["English-language sources"] = sum("English" in (x.get("Language") or "") or "انگلیسی" in (x.get("Language") or "") for x in data["sources"])
    for grade in "ABCD":
        dashboard[f"Reliability {grade}"] = sum(x.get("Reliability") == grade for x in data["sources"])
    dashboard["Sources with direct PDF links"] = sum(x.get("PDF_Status") == "Direct Link Available" for x in data["sources"])
    dashboard["Last Source ID"] = data["sources"][-1]["Source_ID"]
    dashboard["Deep literature articles added"] = len(records)
    dashboard["Deep literature Persian articles"] = sum(x["Language"] == "فارسی" for x in records)
    dashboard["Deep literature English articles"] = sum(x["Language"] == "English" for x in records)


def update_excel(records: list[dict], data: dict) -> None:
    path = ROOT / "Data.xlsx"
    wb = load_workbook(path)
    for record in records:
        values = [record.get(field) for field in SOURCE_FIELDS]
        for sheet in ("ALL_SOURCES", "NEW_SOURCES", "ACADEMIC"):
            append_styled(wb[sheet], values)
        append_styled(wb["TIMELINE"], [record.get(key) for key in (
            "Source_ID", "Date_Gregorian", "Date_Persian", "Project", "Title", "Source_Type", "Publisher", "Reliability", "URL"
        )])
        if record.get("PDF_URL"):
            append_styled(wb["PDF_INDEX"], [record.get(key) for key in (
                "Source_ID", "Title", "Publisher", "PDF_Status", "PDF_URL", "Local_PDF", "SHA256", "Reliability"
            )])

    ws = wb["ALL_SOURCES"]
    if "AllSourcesClean" in ws.tables:
        ws.tables["AllSourcesClean"].ref = f"A1:AK{ws.max_row}"
    wb["README"]["A1"] = "Water Transfer Isfahan — DEEP SCHOLARLY LITERATURE R6"
    wb["README"]["B2"] = len(data["sources"])
    wb["README"]["B3"] = f"S001–{data['meta']['lastSourceId']} (continuous, no gaps)"
    wb["README"]["B4"] = f"Through {TODAY}"
    wb["DASHBOARD"]["A1"] = "DEEP SCHOLARLY LITERATURE R6 DASHBOARD"
    labels = {wb["DASHBOARD"].cell(row, 1).value: row for row in range(2, wb["DASHBOARD"].max_row + 1)}
    for label, value in data["meta"]["dashboard"].items():
        if label in labels:
            wb["DASHBOARD"].cell(labels[label], 2).value = value
        elif label.startswith("Deep literature"):
            append_styled(wb["DASHBOARD"], [label, value, "Validated scholarly discovery R6"])
    append_styled(wb["SEARCH_LOG"], [
        "مرور عمیق فارسی و انگلیسی: انتقال بین‌حوضه‌ای، زاینده‌رود، کارون و طرح‌های وابسته",
        "فارسی/English", "Crossref, DOAJ, OpenAlex (limited), publisher pages, Persian scholarly web indexes",
        "همه حوضه‌ها و طرح‌ها", "مقاله علمی", TODAY, len(records),
        "حذف تکراری بر پایه DOI/عنوان؛ پذیرش فقط با رکورد نمایه معتبر و DOI یا صفحه علمی پایدار.",
    ])
    append_styled(wb["CHANGE_LOG"], [
        "C031", TODAY, "DEEP LITERATURE MERGE", f"{records[0]['Source_ID']}–{records[-1]['Source_ID']}",
        "Systematic Persian/English scholarly discovery",
        f"Added {len(records)} verified unique scholarly articles from Crossref/DOAJ and supplementary bibliographic searches.",
    ])
    wb.save(path)


def write_report(rows: list[dict], records: list[dict], search_log: list[dict]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "خلاصه"
    summary.sheet_view.rightToLeft = True
    metrics = [
        ("تاریخ جست‌وجو", TODAY_FA),
        ("رکورد یکتای بررسی‌شده", len(rows)),
        ("از قبل موجود", sum(x["status"] == "existing" for x in rows)),
        ("افزوده‌شده", len(records)),
        ("نگه‌داری برای بررسی بیشتر", sum(x["status"] == "hold" for x in rows)),
        ("مقالات فارسی افزوده‌شده", sum(x["Language"] == "فارسی" for x in records)),
        ("مقالات انگلیسی افزوده‌شده", sum(x["Language"] == "English" for x in records)),
    ]
    summary.append(["شاخص", "مقدار"])
    for metric in metrics:
        summary.append(metric)

    for title, status in (("افزوده‌شده", "added"), ("ازقبل‌موجود", "existing"), ("نیازمندبررسی", "hold")):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        headers = ["وضعیت", "عنوان", "نویسندگان", "سال", "مجله", "DOI", "لینک", "دامنه", "پایگاه‌ها", "شناسه داشبورد", "علت"]
        ws.append(headers)
        for row in rows:
            if row["status"] != status:
                continue
            ws.append([row.get(key) for key in ("status", "title", "authors", "year", "journal", "doi", "url", "scope", "databases", "dashboard_id", "reason")])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths = [16, 65, 45, 10, 35, 35, 55, 30, 25, 18, 50]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width
    wb.save(REPORTS / "deep_literature_report.xlsx")


def update_offline_fallback(data: dict) -> None:
    path = ROOT / "app.js"
    text = path.read_text(encoding="utf-8")
    marker = "const OFFLINE_FALLBACK = "
    start = text.index(marker) + len(marker)
    end = text.index(";\n\nconst $", start)
    packed = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    path.write_text(text[:start] + packed + text[end:], encoding="utf-8", newline="\n")


def main() -> None:
    candidates, supplemental_log = combine_candidates()
    data = json.loads((ROOT / "data.json").read_text(encoding="utf-8"))
    existing_sources = data["sources"]
    old_validation = json.loads((REPORTS / "deep_literature_search.json").read_text(encoding="utf-8"))["candidates"]
    validated = {candidate_key(x): x for x in old_validation}

    selected, report = [], []
    for item in candidates:
        match_id, match_score = existing_match(item, existing_sources)
        old = validated.get(candidate_key(item)) or {}
        databases = sorted(set(item.get("discovered_in") or [item.get("source_database")]))
        crossref_record = bool(item.get("doi")) and any("Crossref" in db for db in databases)
        doaj_page = "DOAJ" in databases and bool(item.get("landing_url"))
        old_eligible = old.get("decision") == "eligible"
        if match_id:
            status, reason = "existing", f"تکراری با {match_id} (امتیاز عنوان {match_score:.3f})"
        elif crossref_record:
            status, reason = "added", "DOI و متادیتای مقاله مستقیماً در Crossref ثبت شده است."
        elif old_eligible or doaj_page:
            status, reason = "added", "رکورد مقاله در DOAJ و DOI/صفحه علمی آن معتبر است."
        else:
            status, reason = "hold", "DOI حل‌شونده یا صفحه علمی پایدار برای پذیرش خودکار تأیید نشد."
        dashboard_id = None
        if status == "added":
            dashboard_id = f"S{len(existing_sources) + len(selected) + 1:03d}"
            selected.append(record_from(item, dashboard_id))
        report.append({
            "status": status, "title": item.get("title"), "authors": "، ".join(item.get("authors") or []),
            "year": item.get("year"), "journal": item.get("journal"), "doi": item.get("doi"),
            "url": f"https://doi.org/{item['doi']}" if item.get("doi") else item.get("landing_url"),
            "scope": item.get("scope"), "databases": ", ".join(databases),
            "dashboard_id": dashboard_id or match_id, "reason": reason,
        })

    if not selected:
        raise SystemExit("No new validated records selected.")
    data["sources"].extend(selected)
    for record in selected:
        data["timeline"].append({key: record.get(key) for key in (
            "Source_ID", "Date_Gregorian", "Date_Persian", "Project", "Title", "Source_Type", "Publisher", "Reliability", "URL", "Year"
        )})
        data["collections"]["academicIds"].append(record["Source_ID"])
        if record.get("PDF_URL"):
            data["pdfIndex"].append({key: record.get(key) for key in (
                "Source_ID", "Title", "Publisher", "PDF_Status", "PDF_URL", "Local_PDF", "SHA256", "Reliability"
            )})
    data["searchLog"].append({
        "Query": "مرور عمیق فارسی و انگلیسی انتقال بین‌حوضه‌ای، زاینده‌رود، کارون و طرح‌های مرتبط",
        "Language": "فارسی/English", "Database / Search Engine": "Crossref, DOAJ, OpenAlex, publisher/Persian scholarly indexes",
        "Project": "همه حوضه‌ها و طرح‌ها", "Source Type": "مقاله علمی", "Search Date": TODAY,
        "Number of useful results": len(selected), "Notes": "حذف تکراری DOI/عنوان و کنترل لینک/متادیتای علمی.",
    })
    data["changeLog"].append({
        "Change_ID": "C031", "Date": TODAY, "Action": "DEEP LITERATURE MERGE",
        "Source_ID / Item": f"{selected[0]['Source_ID']}–{selected[-1]['Source_ID']}",
        "Reason": "Systematic Persian/English scholarly discovery",
        "Details": f"Added {len(selected)} verified unique articles; audit report under reports/.",
    })
    update_meta(data, selected)
    save_json(ROOT / "data.json", data)
    update_excel(selected, data)
    update_offline_fallback(data)
    write_report(report, selected, supplemental_log)
    save_json(REPORTS / "deep_literature_merge_results.json", {
        "date": TODAY, "added": len(selected), "first_id": selected[0]["Source_ID"], "last_id": selected[-1]["Source_ID"],
        "existing": sum(x["status"] == "existing" for x in report), "hold": sum(x["status"] == "hold" for x in report),
        "persian": sum(x["Language"] == "فارسی" for x in selected), "english": sum(x["Language"] == "English" for x in selected),
        "rows": report,
    })
    print(json.dumps({
        "candidates": len(candidates), "added": len(selected), "existing": sum(x["status"] == "existing" for x in report),
        "hold": sum(x["status"] == "hold" for x in report), "first": selected[0]["Source_ID"], "last": selected[-1]["Source_ID"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
