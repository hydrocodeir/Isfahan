#!/usr/bin/env python3
"""Curate the verified Persian challenge-focused Karun/Zayandeh-Rud sources.

This pass also corrects a false-positive tunnel article (S942), enriches the
previously incomplete Beheshtabad paper (S096), and repairs legal-sheet routing
for records added in the preceding pass.
"""
from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
TODAY = "2026-09-13"
TODAY_FA = "1405/06/22"
FIELDS = [
    "Source_ID", "Title", "Title_EN", "Date_Gregorian", "Date_Persian", "Year", "Language",
    "Source_Type", "Publisher", "Author", "Interviewee", "Person", "Position", "Organization",
    "Province", "Location", "Project", "Basin_Source", "Basin_Destination", "Topic",
    "Position_Stance", "Summary", "Key_Claim", "Quote", "URL", "Canonical_URL", "Archive_URL",
    "DOI", "Document_ID", "Reliability", "Primary_or_Secondary", "Full_Text", "PDF_Status",
    "PDF_URL", "Local_PDF", "SHA256", "Notes",
]
TIMELINE_JSON_FIELDS = (
    "Source_ID", "Date_Gregorian", "Date_Persian", "Year", "Project", "Title",
    "Source_Type", "Publisher", "Reliability", "URL",
)
TIMELINE_XLSX_FIELDS = (
    "Source_ID", "Date_Gregorian", "Date_Persian", "Project", "Title",
    "Source_Type", "Publisher", "Reliability", "URL",
)
PDF_FIELDS = (
    "Source_ID", "Title", "Publisher", "PDF_Status", "PDF_URL", "Local_PDF", "SHA256", "Reliability",
)


def rec(**kwargs):
    row = {field: None for field in FIELDS}
    row.update(kwargs)
    return row


CORRECTED_S942 = rec(
    Source_ID="S942",
    Title="چالش‌های زمین‌شناختی پروژه‌های انتقال آب بین‌حوضه‌ای (مطالعه موردی: طرح انتقال آب بهشت‌آباد به فلات مرکزی)",
    Title_EN="Geological challenges of inter-basin water-transfer projects: the Beheshtabad transfer to Iran's Central Plateau",
    Date_Gregorian="2012",
    Date_Persian="1391",
    Year=1391,
    Language="فارسی",
    Source_Type="مقاله کنفرانسی/زمین‌شناسی مهندسی",
    Publisher="همایش ملی انتقال آب بین‌حوضه‌ای (چالش‌ها و فرصت‌ها)، دانشگاه آزاد اسلامی شهرکرد / بازنشر متن کامل",
    Author="سید نعیم امامی",
    Person="سید نعیم امامی",
    Position="استادیار مرکز تحقیقات کشاورزی و منابع طبیعی چهارمحال‌وبختیاری",
    Organization="مرکز تحقیقات کشاورزی و منابع طبیعی چهارمحال‌وبختیاری",
    Province="چهارمحال‌وبختیاری / اصفهان",
    Location="سد و تونل بهشت‌آباد تا زاینده‌رود",
    Project="بهشت‌آباد",
    Basin_Source="بهشت‌آباد/کارون",
    Basin_Destination="زاینده‌رود / فلات مرکزی",
    Topic="زمین‌شناسی مهندسی، گسل، کارست، آبخوان، نشت تونل، رانش",
    Position_Stance="انتقادی/فنی",
    Summary="بر پایه بازدید صحرایی و مرور گزارش‌های مشاور، خطر فرسایش و لغزش مخزن، عبور تونل از پهنه‌های گسلی و کارستی، هجوم آب و اختلال آبخوان‌های مبدأ را بررسی می‌کند.",
    Key_Claim="مسیر ۶۵ کیلومتری با ۱۰ سامانه گسلی و دست‌کم ۲۵ گسل اصلی، کارست گسترده و آبخوان‌های تحت فشار روبه‌رو است و حفاری می‌تواند چشمه‌ها، چاه‌ها و تعادل هیدروژئولوژیک را مختل کند.",
    URL="https://waterresources.blogfa.com/post/62",
    Canonical_URL="https://waterresources.blogfa.com/post/62",
    Reliability="B",
    Primary_or_Secondary="Primary conference paper; full-text secondary mirror",
    Full_Text="Yes",
    PDF_Status="Full text HTML mirror",
    Notes="جایگزین رکورد قبلی S942 شد؛ مقاله قبلی درباره تونل انتقال آب کرمان بود و به‌اشتباه به بهشت‌آباد نسبت داده شده بود. متن حاضر مستقیماً پروژه بهشت‌آباد، کارون و زاینده‌رود را نام می‌برد؛ صفحه ناشر اصلی بازیابی نشد.",
)

ENRICHED_S096 = rec(
    Source_ID="S096",
    Title="بررسی چالش‌های طرح‌های انتقال آب بین‌حوضه‌ای (مورد مطالعاتی: طرح انتقال آب بهشت‌آباد)",
    Title_EN="Investigating the Challenges of Inter-Basin Water Transfer Projects (Case Study: Beheshtabad Water Transfer Plan)",
    Date_Gregorian="2024-12-19",
    Date_Persian="1403/09/29",
    Year=1403,
    Language="فارسی/English abstract",
    Source_Type="مقاله علمی/مطالعه موردی",
    Publisher="نشریه آب و توسعه پایدار، دوره ۱۱، شماره ۳، صص ۲۷–۴۰ / دانشگاه فردوسی مشهد",
    Author="فرشاد علی‌پور نصیرمحله؛ رسول میرعباسی نجف‌آبادی؛ علی ترابی حقیقی",
    Person="فرشاد علی‌پور نصیرمحله؛ رسول میرعباسی نجف‌آبادی؛ علی ترابی حقیقی",
    Position="پژوهشگران مهندسی منابع آب",
    Organization="دانشگاه شهرکرد؛ دانشگاه اولو",
    Province="چهارمحال‌وبختیاری / اصفهان",
    Project="بهشت‌آباد",
    Basin_Source="بهشت‌آباد/کارون",
    Basin_Destination="اصفهان / حوضه زاینده‌رود",
    Topic="فنی، محیط‌زیستی، اقتصادی، اجتماعی، سیاسی، امنیتی، حکمرانی",
    Position_Stance="تحلیلی/میان‌رشته‌ای",
    Summary="مطالعه موردی چالش‌های بهشت‌آباد نشان می‌دهد تمرکز صرف بر اصلاحات فنی کافی نیست و آثار محیط‌زیستی، اقتصادی، اجتماعی و سیاسی-امنیتیِ مبدأ و مقصد باید یکجا ارزیابی شود.",
    Key_Claim="نحوه مواجهه با پیامدهای احتمالی بهشت‌آباد یکی از چالش‌های مهم حکمرانی آب است و نبود راهکار اجتماعی-سیاسی می‌تواند تعارضات را تشدید کند.",
    URL="https://doi.org/10.22067/jwsd.v11i3.2410-1375",
    Canonical_URL="https://doi.org/10.22067/jwsd.v11i3.2410-1375",
    Archive_URL="https://oulurepo.oulu.fi/handle/10024/59474",
    DOI="10.22067/jwsd.v11i3.2410-1375",
    Document_ID="DOI:10.22067/jwsd.v11i3.2410-1375",
    Reliability="A",
    Primary_or_Secondary="Primary peer-reviewed article",
    Full_Text="Yes",
    PDF_Status="Direct Link Available",
    PDF_URL="https://jwsd.um.ac.ir/article_46223_254aeaad72ea9fae24715d6f62dc71a5.pdf",
    Notes="عنوان فارسی، نویسندگان، DOI صحیح (2410-1375، نه 1375-2410)، تاریخ انتشار و صفحات با متن مقاله و رکورد OuluREPO تطبیق شد.",
)

ADDED = [
    rec(
        Source_ID="S950", Title="بحران آب در حوضه زاینده‌رود؛ مؤلفه‌ها و راهکارها",
        Title_EN="The water crisis in the Zayandeh-Rud Basin: components and solutions",
        Date_Gregorian="2014", Date_Persian="زمستان 1392", Year=1392, Language="فارسی",
        Source_Type="مقاله تحلیلی/فنی", Publisher="فصلنامه دریچه، شماره ۳۱، صص ۲۷–۴۲",
        Author="لطف‌الله ضیایی", Person="لطف‌الله ضیایی",
        Position="کارشناس و عضو کمیته آب اتاق بازرگانی اصفهان", Organization="اتاق بازرگانی اصفهان",
        Province="اصفهان / خوزستان / چهارمحال‌وبختیاری", Project="زاینده‌رود / بهشت‌آباد / انتقال کارون و دز",
        Basin_Source="کارون/دز", Basin_Destination="زاینده‌رود",
        Topic="بیلان آب، حقابه، کشاورزی، صنعت، انتقال آب، گزینه‌های احیا", Position_Stance="موافق مشروط/تحلیلی",
        Summary="تحلیلی از کسری آب، حقابه‌های تاریخی، افت آبخوان و خشکی گاوخونی که انتقال آبِ مازاد کارون/دز را در کنار مدیریت مصرف پیشنهاد می‌کند.",
        Key_Claim="یک دیدگاه موافق انتقال است که آن را تنها پس از احراز مازاد، رعایت معیارهای یونسکو و همراهی با اصلاح مصرف قابل دفاع می‌داند؛ بنابراین شاهد مهمی برای ثبت طیف موافقان است.",
        URL="https://darichejournal.ir/wp-content/uploads/2020/split/%D9%84%D8%B7%D9%81%20%D8%A7%D9%84%D9%84%D9%87%20%D8%B6%DB%8C%D8%A7%DB%8C%DB%8C/31.pdf",
        Canonical_URL="https://darichejournal.ir/wp-content/uploads/2020/split/%D9%84%D8%B7%D9%81%20%D8%A7%D9%84%D9%84%D9%87%20%D8%B6%DB%8C%D8%A7%DB%8C%DB%8C/31.pdf",
        Archive_URL="https://darichejournal.ir/all-articles", Reliability="B",
        Primary_or_Secondary="Primary analytical article; not identified as peer reviewed", Full_Text="Yes",
        PDF_Status="Direct Link Available",
        PDF_URL="https://darichejournal.ir/wp-content/uploads/2020/split/%D9%84%D8%B7%D9%81%20%D8%A7%D9%84%D9%84%D9%87%20%D8%B6%DB%8C%D8%A7%DB%8C%DB%8C/31.pdf",
        Notes="متن کامل ۱۶ صفحه‌ای و آرشیو شماره ۳۱ کنترل شد؛ موضع نویسنده در Summary/Position_Stance صریحاً به‌عنوان دیدگاه تحلیلی ثبت شده، نه اجماع علمی.",
    ),
    rec(
        Source_ID="S951", Title="برنامه احیای حوضه آبریز زاینده‌رود و گاوخونی با رویکرد اصلاح حکمرانی آب",
        Title_EN="Zayandeh-Rud and Gavkhouni Basin Restoration Plan: a water-governance reform approach",
        Date_Gregorian="2021-11", Date_Persian="آبان 1400", Year=1400, Language="فارسی",
        Source_Type="گزارش سیاستی/برنامه حکمرانی", Publisher="انتشار مستقل نویسنده، ویرایش پنجم",
        Author="سجاد انتشاری", Person="سجاد انتشاری", Position="پژوهشگر مدیریت منابع آب",
        Organization="دانشگاه صنعتی اصفهان / IHE Delft (سوابق نویسنده)",
        Province="اصفهان / چهارمحال‌وبختیاری / خوزستان", Project="زاینده‌رود / گاوخونی / انتقال کارون",
        Basin_Source="کارون", Basin_Destination="زاینده‌رود/گاوخونی",
        Topic="حکمرانی، مشارکت ذی‌نفعان، شفافیت، حقابه، تقاضای القایی", Position_Stance="اصلاح‌گر/انتقادی",
        Summary="برنامه ۴۰ صفحه‌ای حاصل مرور نقشه‌راه‌ها و مشورت با ذی‌نفعان سه استان است و بحران را بیش از کمبود سازه، مسئله حکمرانی، اعتماد و تخصیص می‌داند.",
        Key_Claim="انتقال‌های قبلی به‌دلیل تخصیص و توسعه جدید کمبود را پایدار کرده‌اند؛ هر انتقال تازه باید پس از شفافیت داده، تثبیت حقوق، مدیریت تقاضا و مذاکره بین‌حوضه‌ای بررسی شود.",
        URL="https://sajadenteshari.ir/wp-content/uploads/2021/11/%D8%A8%D8%B1%D9%86%D8%A7%D9%85%D9%87-%D8%A7%D8%AD%DB%8C%D8%A7%DB%8C-%D8%B2%D8%A7%DB%8C%D9%86%D8%AF%D9%87-%D8%B1%D9%88%D8%AF-17-1.pdf",
        Canonical_URL="https://sajadenteshari.ir/wp-content/uploads/2021/11/%D8%A8%D8%B1%D9%86%D8%A7%D9%85%D9%87-%D8%A7%D8%AD%DB%8C%D8%A7%DB%8C-%D8%B2%D8%A7%DB%8C%D9%86%D8%AF%D9%87-%D8%B1%D9%88%D8%AF-17-1.pdf",
        Reliability="B", Primary_or_Secondary="Primary policy proposal; stakeholder-informed, not an enacted government plan",
        Full_Text="Yes", PDF_Status="Direct Link Available",
        PDF_URL="https://sajadenteshari.ir/wp-content/uploads/2021/11/%D8%A8%D8%B1%D9%86%D8%A7%D9%85%D9%87-%D8%A7%D8%AD%DB%8C%D8%A7%DB%8C-%D8%B2%D8%A7%DB%8C%D9%86%D8%AF%D9%87-%D8%B1%D9%88%D8%AF-17-1.pdf",
        Notes="به‌عنوان پیشنهاد سیاستی مستقل ثبت شده و نباید با مصوبه یا برنامه رسمی دولت اشتباه شود.",
    ),
    rec(
        Source_ID="S952", Title="پیامدهای اجتماعی طرح‌های توسعه و انتقال آب با تأکید بر مفهوم عدالت (مطالعه موردی استان چهارمحال‌وبختیاری)",
        Title_EN="Social consequences of water development and transfer with emphasis on justice: Chaharmahal and Bakhtiari Province",
        Date_Gregorian="2024-03-01", Date_Persian="1402/12/11", Year=1402, Language="فارسی",
        Source_Type="مقاله علمی/جامعه‌شناسی", Publisher="مطالعات سیاسی-اجتماعی تاریخ و فرهنگ ایران، ۲(۴)، ۲۰۲–۲۲۵",
        Author="امیر جعفری آزاد؛ خلیل میرزایی؛ سیف‌الله سیف‌اللهی",
        Province="چهارمحال‌وبختیاری / اصفهان / خوزستان", Project="انتقال کارون / بهشت‌آباد / کوهرنگ",
        Basin_Source="سرشاخه‌های کارون در چهارمحال‌وبختیاری", Basin_Destination="زاینده‌رود/فلات مرکزی",
        Topic="عدالت آبی، معیشت، نابرابری، ارزیابی اجتماعی، تعارض استانی", Position_Stance="انتقادی/اجتماعی",
        Summary="با استفاده از اسناد، مصاحبه، شبکه‌های اجتماعی و پیمایش، پیامدهای انتقال آب بر معیشت، احساس تبعیض و تعارضات منطقه‌ای در استان مبدأ را تحلیل می‌کند.",
        Key_Claim="ارزیابی‌های اجتماعی و محیط‌زیستیِ صوری یا ناکافی، هزینه‌های انتقال را به جوامع مبدأ منتقل کرده و مسئله عدالت و اعتماد را تشدید می‌کند.",
        URL="https://doi.org/10.61838/kman.jspsich.2.4.10", Canonical_URL="https://doi.org/10.61838/kman.jspsich.2.4.10",
        Archive_URL="https://www.journalspsich.com/index.php/journalspsich/article/view/92",
        DOI="10.61838/kman.jspsich.2.4.10", Document_ID="DOI:10.61838/kman.jspsich.2.4.10",
        Reliability="A", Primary_or_Secondary="Primary peer-reviewed article", Full_Text="Yes",
        PDF_Status="Direct Link Available", PDF_URL="https://www.journalspsich.com/index.php/journalspsich/article/download/92/77/437",
        Notes="عنوان، نویسندگان، تاریخ‌های دریافت/پذیرش، صفحات و DOI با PDF و صفحه مجله تطبیق شد.",
    ),
    rec(
        Source_ID="S953", Title="گفت‌وگو با پرویز کردوانی: وقتی آب را از کارون بگیرند و کشاورزان کشت نکنند ریزگرد بلند می‌شود",
        Date_Gregorian="2017-03-05", Date_Persian="1395/12/15", Year=1395, Language="فارسی",
        Source_Type="مصاحبه مستقیم/بازنشر", Publisher="ایرنا/عصر ایران؛ بازنشر ملیون ایران",
        Interviewee="پرویز کردوانی", Person="پرویز کردوانی", Position="استاد و پژوهشگر کویر و بیابان",
        Province="خوزستان / اصفهان / چهارمحال‌وبختیاری", Project="انتقال کارون به زاینده‌رود",
        Basin_Source="کارون", Basin_Destination="زاینده‌رود/فلات مرکزی",
        Topic="ریزگرد، کشاورزی، تالاب، مدیریت تقاضا، تقاضای القایی", Position_Stance="مخالف مدیریت‌نشده/مشروط",
        Summary="مصاحبه تفصیلی درباره ارتباط کاهش آب کارون با رهاشدن اراضی و ریزگرد و نیز خطر افزایش تقاضا پس از انتقال آب.",
        Key_Claim="انتقال بدون الزام صرفه‌جویی و بدون تأمین حقابه پایین‌دست مانند وام‌گرفتن بحران است و می‌تواند هم مبدأ و هم مقصد را آسیب‌پذیرتر کند.",
        URL="https://melliun.org/iran/117736", Canonical_URL="https://melliun.org/iran/117736",
        Reliability="B", Primary_or_Secondary="Secondary reproduction of a direct IRNA/Asriran interview", Full_Text="Yes",
        PDF_Status="Landing Page Only", Notes="صفحه بازنشر، تاریخ و متن کامل سؤال‌وجواب را دارد؛ اصل صفحه ایرنا/عصر ایران در جست‌وجوی جاری بازیابی نشد.",
    ),
    rec(
        Source_ID="S954", Title="مشاور استاندار خوزستان: طرح «کوهرنگ سه» تنش زیست‌محیطی خوزستان را دوچندان می‌کند",
        Date_Gregorian="2018-02-16", Date_Persian="1396/11/27", Year=1396, Language="فارسی",
        Source_Type="مصاحبه مستقیم", Publisher="ایران‌وایر", Author="کوروش بهرامی",
        Interviewee="مهدی قمشی", Person="مهدی قمشی", Position="عضو هیئت علمی دانشگاه شهید چمران و مشاور استاندار خوزستان",
        Organization="دانشگاه شهید چمران اهواز", Province="خوزستان / چهارمحال‌وبختیاری / اصفهان",
        Project="کوهرنگ ۳ / بهشت‌آباد / ونک-سولگان", Basin_Source="کارون بزرگ", Basin_Destination="زاینده‌رود/فلات مرکزی",
        Topic="اثر تجمعی انتقال‌ها، حقابه محیط‌زیست، صنایع آب‌بر، گزینه جایگزین", Position_Stance="مخالف/کارشناسی",
        Summary="گفت‌وگوی مستقیم درباره اثر تجمعی برداشت‌های متعدد از کارون و دز، نه صرفاً حجم یک پروژه، و اثر آن بر کشاورزی، شرب، آبزی‌پروری و محیط‌زیست خوزستان.",
        Key_Claim="ارزیابی باید مجموع کوهرنگ ۳، بهشت‌آباد و ونک-سولگان را بسنجد؛ صنایع آب‌بر نیز باید به پساب یا آب شور/دریا متکی شوند.",
        URL="https://iranwire.com/fa/features/24927/", Canonical_URL="https://iranwire.com/fa/features/24927/",
        Reliability="B", Primary_or_Secondary="Primary direct interview", Full_Text="Yes", PDF_Status="Landing Page Only",
        Notes="با مصاحبه قدیمی قمشی در S949 تکراری نیست؛ تاریخ، مصاحبه‌گر و محتوای مستقل دارد.",
    ),
    rec(
        Source_ID="S955", Title="زاینده‌رود؛ مشکل امروز، مصیبت فردا", Date_Gregorian="2014", Date_Persian="1393", Year=1393,
        Language="فارسی", Source_Type="مقاله تحلیلی/کارشناسی", Publisher="ایران امروز؛ بازنشر ملیون ایران",
        Author="خسرو بندری", Person="خسرو بندری", Position="کارشناس باسابقه طرح‌های آب، خاک و کشاورزی",
        Province="اصفهان / چهارمحال‌وبختیاری / خوزستان", Project="بهشت‌آباد / کوهرنگ ۱ و ۲ / زاینده‌رود",
        Basin_Source="کارون/دز", Basin_Destination="زاینده‌رود",
        Topic="تاریخچه، ظرفیت اکولوژیک، آمایش، صنعت آب‌بر، ارزیابی اجتماعی", Position_Stance="مخالف انتقال جدید/اصلاح‌گر",
        Summary="مروری کارشناسی بر تاریخ پروژه‌های انتقال و رشد مصارف مقصد که توقف بهشت‌آباد تا بازبینی جامع کارون، دز و زاینده‌رود را پیشنهاد می‌کند.",
        Key_Claim="مطالعات بدون ارزیابی اکولوژیک، اجتماعی و جمعیت‌شناختی هر دو حوضه ناقص است و تزریق آب می‌تواند رشد جمعیت و صنعت فراتر از ظرفیت طبیعی را تثبیت کند.",
        URL="https://melliun.org/iran/47407/amp", Canonical_URL="https://melliun.org/iran/47407/amp",
        Reliability="B", Primary_or_Secondary="Primary expert analysis on a secondary host", Full_Text="Yes", PDF_Status="Landing Page Only",
        Notes="تاریخ دقیق در بازنشر نمایش داده نمی‌شود؛ سال ۱۳۹۳ از ارجاعات زمانی داخل متن و آرشیو بازنشر استنباط و با عدم قطعیت ثبت شد.",
    ),
    rec(
        Source_ID="S956", Title="بررسی سازوکارهای برساخت مسائل محیط‌زیستی (مورد مطالعه: کارزار مردمی توقف تونل بهشت‌آباد)",
        Title_EN="Mechanisms of constructing environmental issues: the public campaign to stop the Beheshtabad Tunnel",
        Date_Gregorian="2021", Date_Persian="1400", Year=1400, Language="فارسی",
        Source_Type="مقاله علمی/جامعه‌شناسی محیط‌زیست", Publisher="برنامه‌ریزی رفاه و توسعه اجتماعی، ۱۲(۴۹)، ۲۱۷–۲۷۴",
        Author="بهرنگ ضابطیان؛ مرضیه موسوی", Province="چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد",
        Basin_Source="بهشت‌آباد/کارون", Basin_Destination="زاینده‌رود",
        Topic="کارزار مردمی، برساخت مسئله، کنش جمعی، تعارض اجتماعی", Position_Stance="تحلیلی/اجتماعی",
        Summary="مقاله فرایند شکل‌گیری و صورت‌بندی مخالفت اجتماعی با تونل بهشت‌آباد را به‌عنوان یک مسئله محیط‌زیستی و کارزار مردمی تحلیل می‌کند.",
        Key_Claim="مشروعیت اجتماعی پروژه فقط تابع محاسبات فنی نیست؛ نحوه تعریف مسئله، دسترسی به رسانه و مشارکت گروه‌های محلی بر منازعه اثر می‌گذارد.",
        URL="https://doi.org/10.22054/qjsd.2021.57194.2077", Canonical_URL="https://doi.org/10.22054/qjsd.2021.57194.2077",
        Archive_URL="https://qjsd.atu.ac.ir/article_13765.html", DOI="10.22054/qjsd.2021.57194.2077",
        Document_ID="DOI:10.22054/qjsd.2021.57194.2077", Reliability="A",
        Primary_or_Secondary="Primary peer-reviewed article", Full_Text="Landing/metadata verified", PDF_Status="Landing Page Only",
        Notes="DOI به صفحه رسمی مجله دانشگاه علامه طباطبائی resolve شد؛ متن کامل عمومی در جست‌وجوی جاری بازیابی نشد.",
    ),
    rec(
        Source_ID="S957", Title="سند تحول دولت مردمی ـ فصل آب: کاهش انتقال بین‌حوضه‌ای از منابع داخلی",
        Date_Gregorian="2022-02-27", Date_Persian="1400/12/08", Year=1400, Language="فارسی",
        Source_Type="سند بالادستی/خط‌مشی دولت", Publisher="ریاست جمهوری / سامانه قوانین و مقررات کشور",
        Author="دولت سیزدهم", Organization="ریاست جمهوری", Province="ملی",
        Project="انتقال بین‌حوضه‌ای / کارون / زاینده‌رود", Basin_Source="حوضه‌های آبریز داخلی ایران", Basin_Destination="حوضه‌های آبریز داخلی ایران",
        Topic="کاهش انتقال، حقوق بالادست، سدسازی، آب غیرمتعارف، حکمرانی حوضه‌ای", Position_Stance="کاهشی/تنظیمی",
        Summary="در فصل آب، مدیریت حوضه‌ای را جایگزین مدیریت استانی می‌داند و کاهش انتقال داخلی با رعایت حقوق بالادست، توقف سدهای مضر پایین‌دست و استفاده از آب غیرمتعارف را مقرر می‌کند.",
        Key_Claim="یک خط‌مشی اجرایی ملی است که ارزیابی انتقال‌های کارون–زاینده‌رود را به حقوق مبدأ، حیات پایین‌دست و گزینه‌های جایگزین پیوند می‌دهد.",
        URL="https://nezamat.ir/%D8%B3%D9%86%D8%AF-%D8%AA%D8%AD%D9%88%D9%84-%D8%AF%D9%88%D9%84%D8%AA-%D9%85%D8%B1%D8%AF%D9%85%DB%8C/",
        Canonical_URL="https://nezamat.ir/%D8%B3%D9%86%D8%AF-%D8%AA%D8%AD%D9%88%D9%84-%D8%AF%D9%88%D9%84%D8%AA-%D9%85%D8%B1%D8%AF%D9%85%DB%8C/",
        Document_ID="158483-1400/12/08", Reliability="A", Primary_or_Secondary="Primary government policy text; legal-system mirror",
        Full_Text="Yes", PDF_Status="Landing Page Only", Notes="سند، قانون مصوب مجلس نیست؛ خط‌مشی ابلاغی قوه مجریه ذیل اصل ۱۳۴ است و با همین وصف ثبت شده است.",
    ),
    rec(
        Source_ID="S958", Title="انتقال آب کارون: کابوس خوزستان، رؤیای اصفهان", Date_Gregorian="2013-12-09",
        Date_Persian="1392/09/18", Year=1392, Language="فارسی", Source_Type="یادداشت تحلیلی/سیاست عمومی",
        Publisher="بخش تعاملی الف؛ بازنشر WaterResources", Author="روزبه کردونی",
        Province="خوزستان / چهارمحال‌وبختیاری / اصفهان", Project="بهشت‌آباد / انتقال کارون",
        Basin_Source="کارون و دز", Basin_Destination="زاینده‌رود/فلات مرکزی",
        Topic="هزینه اقتصادی، برق‌آبی، حقابه، تعارض اجتماعی، امنیت", Position_Stance="انتقادی/شرط‌گذار",
        Summary="با اتکا به گزارش ۱۳۸۷ مرکز پژوهش‌های مجلس، آثار احتمالی بهشت‌آباد بر کمبود آینده خوزستان، نیروگاه‌های برق‌آبی و تعارض اجتماعی را مرور می‌کند.",
        Key_Claim="حتی با ادعای منفعت مقصد، تصمیم باید هزینه انرژی، حقوق ذی‌نفعان مبدأ، نیاز محیط‌زیستی و خطر امنیتی‌شدن اختلاف را هم‌زمان بسنجد.",
        URL="https://waterresources.blogfa.com/post/49", Canonical_URL="https://waterresources.blogfa.com/post/49",
        Reliability="C", Primary_or_Secondary="Secondary analytical article citing a parliamentary report", Full_Text="Yes", PDF_Status="Landing Page Only",
        Notes="منبع تحلیلی ثانویه است؛ ارقام آن به‌عنوان داده قطعی جایگزین گزارش اصلی مجلس (S252) نیست. نام نویسنده از سربرگ متن و نام بازنشرکننده از صفحه تفکیک شد.",
    ),
]

DUPLICATES = [
    {"candidate": "بررسی چالش‌های طرح‌های انتقال آب بین‌حوضه‌ای، علی‌پور و همکاران (۱۴۰۳)", "existing": "S096", "reason": "رکورد موجود بود؛ به‌جای افزودن تکراری، عنوان فارسی، نویسندگان، DOI صحیح، تاریخ و صفحات تکمیل شد."},
    {"candidate": "مصاحبه مهدی قمشی درباره نبود آب مازاد خوزستان", "existing": "S949", "reason": "نسخه قدیمی بازنشرشده موجود است؛ فقط مصاحبه مستقل ایران‌وایر با تاریخ/پرسش‌های متفاوت در S954 افزوده شد."},
    {"candidate": "گزارش مرکز پژوهش‌های مجلس درباره بهشت‌آباد، شماره ۸۹۳۵", "existing": "S252", "reason": "یادداشت S958 به آن استناد می‌کند اما خود گزارش دوباره افزوده نشد."},
    {"candidate": "مصاحبه‌های ناصر کرمی، حمید چیت‌چیان و نیک‌آهنگ کوثر", "existing": "S434; S300/S375; S948", "reason": "صفحات بازنشری تکراری کنار گذاشته شدند."},
    {"candidate": "مقاله چالش‌های بهشت‌آباد در پویش خانواده آب", "existing": "S096", "reason": "این صفحه معرفی همان مقاله علمی است و منبع جداگانه محسوب نشد."},
]

GAPS = [
    {"موضوع": "پیامدهای اجتماعی و امنیتی انتقال آب بین‌حوضه‌ای؛ داوودی دهاقانی و عامری (۱۳۹۸)", "وضعیت": "کتابشناسی تأیید؛ صفحه اصل بازیابی نشد", "توضیح": "عنوان، مجله ۷(۲۵)، صفحات ۵۱–۷۶ و DOR در فهرست منابع مقاله S096 دیده شد، اما صفحه ناشر/متن عمومی معتبر پیدا نشد؛ فعلاً به بانک قطعی افزوده نشد."},
    {"موضوع": "پیش‌بینی نشت آب زیرزمینی به تونل بهشت‌آباد؛ ریاحی‌پور و خلیلی (۱۴۰۳)", "وضعیت": "کتابشناسی تأیید؛ لینک اصل بازیابی نشد", "توضیح": "در منابع مقاله S096 به دوازدهمین همایش ملی محیط‌زیست، انرژی و منابع طبیعی ارجاع شده، اما صفحه همایش یا PDF قابل احراز پیدا نشد."},
    {"موضوع": "اصل گزارش/پرونده کمیسیون اصل ۹۰ درباره کارون–زاینده‌رود", "وضعیت": "پیدا نشد", "توضیح": "خبر و نامه نمایندگان به‌جای گزارش رسمی اصل ۹۰ پذیرفته نشد؛ شماره پرونده و متن امضاشده لازم است."},
    {"موضوع": "کیفرخواست یا دادنامه کیفری پروژه‌های بهشت‌آباد/کوهرنگ ۳", "وضعیت": "پیدا نشد", "توضیح": "ادعاهای رسانه‌ای بدون شماره پرونده یا رأی رسمی وارد دیتاست نشدند."},
    {"موضوع": "اصل گزارش بازرسی و اصل نامه قضایی نمایندگان", "وضعیت": "فقط بازنشر ثانویه", "توضیح": "S944 و S945 نگه داشته شدند، اما نسخه امضاشده/شماره دبیرخانه و نتیجه رسیدگی هنوز بازیابی نشده است."},
]


def norm(value):
    return " ".join(str(value or "").replace("‌", " ").replace("ي", "ی").replace("ك", "ک").lower().split())


def copy_row_style(ws, src_row, dst_row):
    for col in range(1, ws.max_column + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format


def append_styled(ws, values):
    old_last = ws.max_row
    ws.append(values)
    if old_last >= 2:
        copy_row_style(ws, old_last, ws.max_row)


def row_by_id(ws, source_id):
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == source_id:
            return row
    return None


def upsert_full_row(ws, source):
    row = row_by_id(ws, source["Source_ID"])
    if row is None:
        append_styled(ws, [source.get(field) for field in FIELDS])
    else:
        for col, field in enumerate(FIELDS, 1):
            ws.cell(row, col).value = source.get(field)


def upsert_compact_row(ws, source, fields):
    row = row_by_id(ws, source["Source_ID"])
    values = [source.get(field) for field in fields]
    if row is None:
        append_styled(ws, values)
    else:
        for col, value in enumerate(values, 1):
            ws.cell(row, col).value = value


def ensure_collection(collection, ids):
    seen = set(collection)
    for source_id in ids:
        if source_id not in seen:
            collection.append(source_id)
            seen.add(source_id)


def update_excel(data):
    path = ROOT / "Data.xlsx"
    wb = load_workbook(path)
    by_id = {source["Source_ID"]: source for source in data["sources"]}

    for source in (ENRICHED_S096, CORRECTED_S942):
        for sheet in ("ALL_SOURCES", "NEW_SOURCES", "ACADEMIC", "OFFICIAL_LEGAL", "PARLIAMENT", "NEWS_MEDIA"):
            if sheet in wb.sheetnames and row_by_id(wb[sheet], source["Source_ID"]):
                upsert_full_row(wb[sheet], source)
        upsert_compact_row(wb["TIMELINE"], source, TIMELINE_XLSX_FIELDS)
        if source.get("PDF_URL"):
            upsert_compact_row(wb["PDF_INDEX"], source, PDF_FIELDS)
        elif row_by_id(wb["PDF_INDEX"], source["Source_ID"]):
            wb["PDF_INDEX"].delete_rows(row_by_id(wb["PDF_INDEX"], source["Source_ID"]), 1)

    academic_ids = {"S950", "S951", "S952", "S956"}
    news_ids = {"S953", "S954", "S955", "S958"}
    legal_ids = {"S936", "S937", "S945", "S957"}
    parliament_ids = {"S936", "S945"}
    for source in ADDED:
        upsert_full_row(wb["ALL_SOURCES"], source)
        upsert_full_row(wb["NEW_SOURCES"], source)
        upsert_compact_row(wb["TIMELINE"], source, TIMELINE_XLSX_FIELDS)
        if source["Source_ID"] in academic_ids:
            upsert_full_row(wb["ACADEMIC"], source)
        if source["Source_ID"] in news_ids:
            upsert_full_row(wb["NEWS_MEDIA"], source)
        if source["Source_ID"] in legal_ids:
            upsert_full_row(wb["OFFICIAL_LEGAL"], source)
        if source["Source_ID"] in parliament_ids:
            upsert_full_row(wb["PARLIAMENT"], source)
        if source.get("PDF_URL"):
            upsert_compact_row(wb["PDF_INDEX"], source, PDF_FIELDS)

    # Repair the three legal rows missed by the previous routing logic.
    for source_id in ("S936", "S937", "S945"):
        upsert_full_row(wb["OFFICIAL_LEGAL"], by_id[source_id])

    wb["README"]["A1"] = "Water Transfer Isfahan — DEEP PERSIAN CHALLENGES CURATION"
    wb["README"]["B2"] = len(data["sources"])
    wb["README"]["B3"] = f"S001–{data['meta']['lastSourceId']} (continuous, no gaps)"
    wb["README"]["B4"] = f"Through {TODAY}"

    dashboard_rows = {wb["DASHBOARD"].cell(row, 1).value: row for row in range(2, wb["DASHBOARD"].max_row + 1)}
    for label, value in data["meta"]["dashboard"].items():
        if label in dashboard_rows:
            wb["DASHBOARD"].cell(dashboard_rows[label], 2).value = value

    append_styled(wb["SEARCH_LOG"], [
        "جست‌وجوی عمیق فارسی: چالش‌های انتقال بین‌حوضه‌ای کارون–زاینده‌رود و بهشت‌آباد",
        "فارسی", "DOI resolver, مجلات دانشگاهی، اسناد دولت، آرشیو رسانه، گزارش‌های فنی و سیاستی",
        "کارون / زاینده‌رود / بهشت‌آباد / کوهرنگ", "مقاله، مصاحبه، سند بالادستی، گزارش فنی و سیاستی",
        TODAY, len(ADDED), "۹ رکورد جدید؛ S096 تکمیل؛ S942 جایگزین؛ تکراری‌ها و شکاف‌های سندی جدا گزارش شدند.",
    ])
    append_styled(wb["CHANGE_LOG"], [
        "C036", TODAY, "DEEP PERSIAN CHALLENGES + CORRECTION", "S096; S942; S950–S958",
        "Add verified non-duplicate Persian challenge sources and remove false-positive scope assignment",
        "Added 9; enriched S096; replaced S942; repaired OFFICIAL_LEGAL routing for S936/S937/S945; reports written.",
    ])
    if "AllSourcesClean" in wb["ALL_SOURCES"].tables:
        wb["ALL_SOURCES"].tables["AllSourcesClean"].ref = f"A1:AK{wb['ALL_SOURCES'].max_row}"
    wb.save(path)


def write_reports():
    report = {
        "date": TODAY,
        "scope": "Iran only; Karun and Zayandeh-Rud basins, especially Beheshtabad/Kouhrang",
        "added": ADDED,
        "corrected": [
            {"Source_ID": "S942", "action": "replaced false-positive Kerman tunnel article", "new_record": CORRECTED_S942},
            {"Source_ID": "S096", "action": "enriched metadata and corrected DOI order", "new_record": ENRICHED_S096},
        ],
        "duplicates_skipped": DUPLICATES,
        "not_added_or_gaps": GAPS,
    }
    (ROOT / "reports/deep_persian_challenges.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.rightToLeft = True
    summary.append(["Metric", "Value"])
    summary.append(["Search date", TODAY_FA])
    summary.append(["Added", len(ADDED)])
    summary.append(["Corrected/enriched", 2])
    summary.append(["Duplicates skipped", len(DUPLICATES)])
    summary.append(["Not added / documentary gaps", len(GAPS)])
    sections = [
        ("Added", ADDED, FIELDS),
        ("Corrected", report["corrected"], ["Source_ID", "action"]),
        ("Duplicates", DUPLICATES, ["candidate", "existing", "reason"]),
        ("Gaps", GAPS, ["موضوع", "وضعیت", "توضیح"]),
    ]
    for title, rows, headers in sections:
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
    wb.save(ROOT / "reports/deep_persian_challenges.xlsx")


def main():
    data_path = ROOT / "data.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    by_id = {source["Source_ID"]: source for source in data["sources"]}
    if any(source["Source_ID"] in by_id for source in ADDED):
        raise SystemExit("One or more S950-S958 records already exist; refusing a non-idempotent rerun.")

    # Exact collision guards against the existing database after excluding the rows being updated.
    comparison = [source for source in data["sources"] if source["Source_ID"] not in {"S096", "S942"}]
    existing_dois = {norm(source.get("DOI")) for source in comparison if source.get("DOI")}
    existing_urls = {norm(source.get("Canonical_URL") or source.get("URL")) for source in comparison if source.get("Canonical_URL") or source.get("URL")}
    existing_titles = {norm(source.get("Title")) for source in comparison if source.get("Title")}
    for source in ADDED:
        if source.get("DOI") and norm(source["DOI"]) in existing_dois:
            raise SystemExit(f"Duplicate DOI: {source['Source_ID']} {source['DOI']}")
        if norm(source.get("Canonical_URL") or source.get("URL")) in existing_urls:
            raise SystemExit(f"Duplicate URL: {source['Source_ID']}")
        if norm(source["Title"]) in existing_titles:
            raise SystemExit(f"Duplicate title: {source['Source_ID']} {source['Title']}")

    for replacement in (ENRICHED_S096, CORRECTED_S942):
        index = next(i for i, source in enumerate(data["sources"]) if source["Source_ID"] == replacement["Source_ID"])
        data["sources"][index] = replacement
    data["sources"].extend(ADDED)
    data["sources"].sort(key=lambda source: int(source["Source_ID"][1:]))

    source_ids = {source["Source_ID"] for source in data["sources"]}
    for name, collection in data["collections"].items():
        data["collections"][name] = list(dict.fromkeys(source_id for source_id in collection if source_id in source_ids))
    ensure_collection(data["collections"]["academicIds"], ["S096", "S942", "S950", "S951", "S952", "S956"])
    ensure_collection(data["collections"]["legalIds"], ["S936", "S937", "S945", "S957"])
    ensure_collection(data["collections"]["parliamentIds"], ["S936", "S945"])
    ensure_collection(data["collections"]["newsIds"], ["S953", "S954", "S955", "S958"])

    replacements = {source["Source_ID"]: source for source in (ENRICHED_S096, CORRECTED_S942)}
    data["timeline"] = [
        {field: replacements[row["Source_ID"]].get(field) for field in TIMELINE_JSON_FIELDS}
        if row.get("Source_ID") in replacements else row
        for row in data["timeline"]
    ]
    data["timeline"].extend({field: source.get(field) for field in TIMELINE_JSON_FIELDS} for source in ADDED)
    data["timeline"].sort(key=lambda row: int(row["Source_ID"][1:]))

    data["pdfIndex"] = [row for row in data["pdfIndex"] if row.get("Source_ID") not in {"S096", "S942"}]
    for source in (ENRICHED_S096, CORRECTED_S942, *ADDED):
        if source.get("PDF_URL"):
            data["pdfIndex"].append({field: source.get(field) for field in PDF_FIELDS})
    data["pdfIndex"].sort(key=lambda row: int(row["Source_ID"][1:]))

    data["searchLog"].append({
        "Query": "جست‌وجوی عمیق فارسی: چالش‌های انتقال بین‌حوضه‌ای کارون–زاینده‌رود و بهشت‌آباد",
        "Language": "فارسی", "Database / Search Engine": "DOI resolver, مجلات دانشگاهی، اسناد دولت، آرشیو رسانه، گزارش‌های فنی و سیاستی",
        "Project": "کارون / زاینده‌رود / بهشت‌آباد / کوهرنگ", "Source Type": "مقاله، مصاحبه، سند بالادستی، گزارش فنی و سیاستی",
        "Search Date": TODAY, "Number of useful results": len(ADDED),
        "Notes": "۹ رکورد جدید؛ S096 تکمیل؛ S942 جایگزین؛ تکراری‌ها و شکاف‌های سندی جدا گزارش شدند.",
    })
    data["changeLog"].append({
        "Change_ID": "C036", "Date": TODAY, "Action": "DEEP PERSIAN CHALLENGES + CORRECTION",
        "Source_ID / Item": "S096; S942; S950–S958",
        "Reason": "Add verified non-duplicate Persian challenge sources and remove false-positive scope assignment",
        "Details": "Added 9; enriched S096; replaced S942; repaired legal routing; separate duplicate/gap report written.",
    })
    for project in data["projects"]:
        project["Source_Count"] = sum(project["Project"] in (source.get("Project") or "") for source in data["sources"])

    meta = data["meta"]
    meta["generatedAt"] = TODAY
    meta["totalSources"] = len(data["sources"])
    meta["lastSourceId"] = ADDED[-1]["Source_ID"]
    dashboard = meta["dashboard"]
    dashboard["Total sources"] = len(data["sources"])
    dashboard["Last Source ID"] = ADDED[-1]["Source_ID"]
    dashboard["Persian-language sources"] = sum("فارسی" in (source.get("Language") or "") for source in data["sources"])
    dashboard["English-language sources"] = sum("English" in (source.get("Language") or "") or "انگلیسی" in (source.get("Language") or "") for source in data["sources"])
    for grade in "ABCD":
        dashboard[f"Reliability {grade}"] = sum(source.get("Reliability") == grade for source in data["sources"])
    dashboard["Sources with direct PDF links"] = sum(source.get("PDF_Status") == "Direct Link Available" for source in data["sources"])
    dashboard["Persian challenge sources added in R6"] = len(ADDED)
    dashboard["Corrected / enriched records in R6"] = 2

    data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    app_path = ROOT / "app.js"
    app = app_path.read_text(encoding="utf-8")
    marker = "const OFFLINE_FALLBACK = "
    start = app.index(marker) + len(marker)
    end = app.index(";\n\nconst $", start)
    app_path.write_text(app[:start] + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + app[end:], encoding="utf-8")
    update_excel(data)
    write_reports()
    print(json.dumps({
        "totalSources": len(data["sources"]), "lastSourceId": meta["lastSourceId"],
        "added": [source["Source_ID"] for source in ADDED], "corrected": ["S096", "S942"],
        "pdfIndex": len(data["pdfIndex"]), "timeline": len(data["timeline"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
