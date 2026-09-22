"""Merge the reviewed Mandegan manifest into JSON, XLSX and the offline dashboard."""
from pathlib import Path
from copy import deepcopy
from urllib.parse import urlsplit, unquote
import hashlib
import json
import re
import jdatetime
from openpyxl import load_workbook
from add_deep_persian_challenges import (append_styled, upsert_full_row,
    upsert_compact_row, FIELDS, TIMELINE_JSON_FIELDS, TIMELINE_XLSX_FIELDS, PDF_FIELDS)

ROOT = Path(__file__).resolve().parents[1]
TODAY = '2026-09-22'
SHEETS = {'academicIds':'ACADEMIC','thesisIds':'THESES','legalIds':'OFFICIAL_LEGAL',
    'parliamentIds':'PARLIAMENT','newsIds':'NEWS_MEDIA','socialIds':'SOCIAL_MEDIA'}

def norm(value):
    return re.sub(r'\s+', ' ', (value or '').replace('ي','ی').replace('ك','ک').replace('\u200c',' ')).strip().casefold()

def urlkey(value):
    if not value: return ''
    u = urlsplit(value)
    host = u.netloc.lower().removeprefix('www.').replace('tasnimnews.com','tasnimnews.ir')
    path = unquote(u.path).rstrip('/')
    patterns = {'tasnimnews.ir':r'/fa/news/\d+/\d+/\d+/(\d+)',
        'mehrnews.com':r'/news/(\d+)', 'rokna.net':r'(?:news-|/)(\d{7})(?:-|$)',
        'khabarfarsi.com':r'/ua?/(\d+)', 'independentpersian.com':r'/node/(\d+)'}
    match = re.search(patterns[host],path) if host in patterns else None
    return host + '/' + match.group(1) if match else host + path.removesuffix('/amp')

def write_json(path, value):
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def main():
    data = json.loads((ROOT/'data.json').read_text(encoding='utf-8'))
    before = deepcopy(data)
    manifest = json.loads((ROOT/'reports/mandegan_curated.json').read_text(encoding='utf-8'))
    added = [item['record'] for item in manifest['rows']]
    if any(s['Source_ID'] == 'S959' for s in data['sources']):
        raise SystemExit('S959 already present; refusing a second merge. Run validate_mandegan.py instead.')
    assert len(data['sources']) == 958
    seen_urls = {urlkey(s.get('Canonical_URL') or s.get('URL')) for s in data['sources']}
    seen_titles = {norm(s['Title']) for s in data['sources']}
    seen_dois = {norm(s['DOI']) for s in data['sources'] if s.get('DOI')}
    for s in added:
        uk, tk = urlkey(s['URL']), norm(s['Title'])
        assert uk not in seen_urls, ('Duplicate URL',s['Source_ID'],uk)
        assert tk not in seen_titles, ('Duplicate title',s['Source_ID'])
        assert not s.get('DOI') or norm(s['DOI']) not in seen_dois
        seen_urls.add(uk); seen_titles.add(tk)
        if s.get('DOI'): seen_dois.add(norm(s['DOI']))
        if s.get('Date_Persian') and len(s['Date_Persian']) == 10:
            assert jdatetime.date(*map(int,s['Date_Persian'].split('/'))).togregorian().isoformat() == s['Date_Gregorian']
    by_id = {s['Source_ID']:s for s in data['sources']}
    enriched = []
    for sid, fields in manifest['enrich_existing'].items():
        by_id[sid].update(fields)
        enriched.append(by_id[sid])
    data['sources'].extend(added)
    for item in manifest['rows']:
        for group in item['collections']: data['collections'][group].append(item['record']['Source_ID'])
    replacements = {s['Source_ID']:s for s in enriched}
    data['timeline'] = [{f:replacements[r['Source_ID']].get(f) for f in TIMELINE_JSON_FIELDS}
        if r['Source_ID'] in replacements else r for r in data['timeline']]
    data['timeline'].extend({f:s.get(f) for f in TIMELINE_JSON_FIELDS} for s in added)
    data['pdfIndex'].extend({f:s.get(f) for f in PDF_FIELDS} for s in added if s.get('PDF_URL'))
    for project in data['projects']:
        project['Source_Count'] = sum(project['Project'] in (s.get('Project') or '') for s in data['sources'])
    query_log = json.loads((ROOT/'reports/mandegan_web_searches.json').read_text(encoding='utf-8'))
    searches = []
    new_keys = {urlkey(s['URL']) for s in added}
    for batch in query_log['searches']:
        searches.append({'Query':' ; '.join(batch['queries']), 'Language':'فارسی / English',
            'Database / Search Engine':'Web search; publisher pages; public archives', 'Project':'ماندگان',
            'Source Type':'خبری / حقوقی / علمی / فنی', 'Search Date':TODAY,
            'Number of useful results':len({urlkey(r['url']) for r in batch['results']} & new_keys),
            'Notes':'تعداد برابر پیوندهای منطبق با منابع منتخب این نوبت است؛ نتایج خام در reports/mandegan_web_searches.json. جست‌وجو مستقیم Google در دسترس نبود؛ نتایج این ردیف ممکن است با ردیف‌های دیگر همپوشانی داشته باشند.'})
    data['searchLog'].extend(searches)
    change = {'Change_ID':'C037','Date':TODAY,'Action':'MANDEGAN SOURCE UPDATE',
        'Source_ID / Item':'S959–S985; S216; S249',
        'Reason':'Add newly located Mandegan news, conference paper and administrative evidence with access limits',
        'Details':'27 added; 2 enriched; 4 partial-access reports labelled; 2 PDFs and 3 documentary images preserved; 70 search queries logged.'}
    data['changeLog'].append(change)
    meta = data['meta']; dashboard = meta['dashboard']
    meta.update(generatedAt=TODAY,totalSources=len(data['sources']),lastSourceId=added[-1]['Source_ID'])
    dashboard.update({'Total sources':len(data['sources']),'Last Source ID':added[-1]['Source_ID'],
        'Persian-language sources':sum('فارسی' in (s.get('Language') or '') for s in data['sources']),
        'English-language sources':sum('English' in (s.get('Language') or '') or 'انگلیسی' in (s.get('Language') or '') for s in data['sources']),
        'Sources with direct PDF links':sum(s.get('PDF_Status')=='Direct Link Available' for s in data['sources']),
        'Locally retrieved PDFs':sum(bool(s.get('Local_PDF')) for s in data['sources']),
        'Mandegan sources added 2026-09-22':len(added), 'Mandegan partial-access additions':4})
    for grade in 'ABCD': dashboard[f'Reliability {grade}'] = sum(s.get('Reliability')==grade for s in data['sources'])

    wb = load_workbook(ROOT/'Data.xlsx')
    for source in enriched:
        for ws in wb:
            if [c.value for c in ws[1]] == FIELDS and any(ws.cell(r,1).value == source['Source_ID'] for r in range(2,ws.max_row+1)):
                upsert_full_row(ws,source)
        upsert_compact_row(wb['TIMELINE'],source,TIMELINE_XLSX_FIELDS)
    for item in manifest['rows']:
        source = item['record']
        for sheet in ['ALL_SOURCES','NEW_SOURCES']+[SHEETS[g] for g in item['collections']]:
            upsert_full_row(wb[sheet],source)
        upsert_compact_row(wb['TIMELINE'],source,TIMELINE_XLSX_FIELDS)
        if source.get('PDF_URL'): upsert_compact_row(wb['PDF_INDEX'],source,PDF_FIELDS)
    for sheet, rows in [('SEARCH_LOG',searches),('CHANGE_LOG',[change])]:
        headers = [c.value for c in wb[sheet][1]]
        for row in rows: append_styled(wb[sheet],[row.get(h) for h in headers])
    for sheet, rows, key in [('DASHBOARD',[{'Metric':k,'Value':v} for k,v in dashboard.items()],'Metric'),
                             ('PROJECTS',data['projects'],'Project')]:
        ws = wb[sheet]
        for row in rows:
            target = next((r for r in range(2,ws.max_row+1) if ws.cell(r,1).value == row[key]),ws.max_row+1)
            ws.cell(target,1).value=row[key]
            ws.cell(target,2).value=row.get('Value',row.get('Source_Count'))
    wb['README']['A1']='Water Transfer Isfahan — MANDEGAN UPDATE'
    wb['README']['B2']=len(data['sources']); wb['README']['B3']='S001–S985 (continuous, no gaps)'; wb['README']['B4']=f'Through {TODAY}'
    for ws in wb:
        if ws.auto_filter.ref: ws.auto_filter.ref=ws.dimensions
        for table in ws.tables.values(): table.ref=ws.dimensions
    staged_xlsx = ROOT/'reports/_tmp_mandegan_Data.xlsx'
    wb.save(staged_xlsx)
    # Reopen the staged workbook before replacing any canonical data.
    check = load_workbook(staged_xlsx,read_only=True)
    assert check['ALL_SOURCES'].max_row == len(data['sources'])+1
    check.close()
    app = (ROOT/'app.js').read_text(encoding='utf-8')
    marker = 'const OFFLINE_FALLBACK = '
    start=app.index(marker)+len(marker); end=app.index(';\n\nconst $',start)
    patched_app = app[:start]+json.dumps(data,ensure_ascii=False,separators=(',',':'))+app[end:]
    assert patched_app[patched_app.index(';\n\nconst $'):] == app[end:]
    write_json(ROOT/'data.json',data)
    (ROOT/'app.js').write_text(patched_app,encoding='utf-8')
    staged_xlsx.replace(ROOT/'Data.xlsx')
    audit={'date':TODAY,'before_count':len(before['sources']),'after_count':len(data['sources']),
        'added_ids':[s['Source_ID'] for s in added],
        'enriched':[{ 'id':sid,'before':next(s for s in before['sources'] if s['Source_ID']==sid),'after':by_id[sid]} for sid in replacements],
        'unchanged_existing_records':sum(s==by_id[s['Source_ID']] for s in before['sources']),
        'query_count':query_log['queries'],'raw_unique_result_urls':len({r['url'] for batch in query_log['searches'] for r in batch['results']}),
        'limitations':manifest['limitations']+query_log['access_limitations']}
    write_json(ROOT/'reports/mandegan_merge_audit.json',audit)
    print(json.dumps({k:audit[k] for k in ['before_count','after_count','unchanged_existing_records','query_count','raw_unique_result_urls']},ensure_ascii=False))

if __name__=='__main__': main()
