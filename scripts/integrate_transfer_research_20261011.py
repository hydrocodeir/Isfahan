"""Stage the reviewed transfer registry and merge its evidence without replacing old records."""
import collections, copy, hashlib, json, pathlib, re, shutil, sys
from urllib.parse import urlsplit, unquote, parse_qsl, urlencode
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/transfer_research_2026_10_11'
TODAY = '2026-10-11'
PKG = ROOT / 'research/water-transfer-review-2026-10-11.json'
def read(p): return json.loads(p.read_text('utf-8-sig'))
def save(p, v): p.write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n', 'utf-8')
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(v): return re.sub(r'[^\w]', '', str(v or '').casefold().replace('ي','ی').replace('ك','ک'))
def urlkey(v):
    u=urlsplit(unquote(v or '')); host=u.netloc.casefold().removeprefix('www.').replace('tasnimnews.com','tasnimnews.ir')
    if host=='academicjournals.org' and re.search(r'/article-[^/]+/',u.path):
        return re.sub(r'/article-[^/]+/','/article/',host+u.path).rstrip('/')
    m=re.search(r'/news/(?:\d{4}/\d{2}/\d{2}/)?(\d+)',u.path)
    if m: return host+'/news/'+m[1]
    return host+u.path.rstrip('/')+('?' + urlencode(sorted((k,v) for k,v in parse_qsl(u.query) if k not in ['lang','utm_source','utm_medium'])) if u.query else '')

DOIS={'S04':'10.22067/jwsd.v11i3.2410-1375','S24':'10.5897/AJAR11.1410',
      'S26':'10.1061/(ASCE)WR.1943-5452.0000818','S27':'10.22059/jwim.2021.325508.887',
      'S28':'10.1016/j.jenvman.2024.120703','S29':'10.1007/s10668-023-03589-z',
      'S30':'10.1016/j.jhydrol.2013.03.021','S31':'10.1016/j.jhydrol.2013.03.021'}
LABELS={'T01':'کوهرنگ ۱','T02':'کوهرنگ ۲','T03':'کوهرنگ ۳','T04':'چشمه‌لنگان',
        'T05':'خدنگستان','T06':'بهشت‌آباد','T07':'خرسان ۳','T08':'گلاب','T09':'بن-بروجن',
        'T10':'سبزکوه','T11':'غدیر خوزستان','T12':'کارون به جراحی','T13':'آبریز / مارون ۲',
        'T14':'تنگ‌سرخ بشار به شیراز','T15':'کوثر','T16':'آب دریا به اصفهان','T17':'دز–قمرود'}
ACADEMIC={'S01','S04','S07','S24','S25','S26','S27','S28','S29','S30','S31','S33','S34','S36','S37','S44'}
ENGLISH={'S01','S07','S13','S15','S16','S17','S22','S23','S24','S26','S28','S29','S30','S31','S36','S37','S43'}

def prepare():
    OUT.mkdir(exist_ok=True); (OUT/'backup').mkdir(exist_ok=True); (OUT/'staged').mkdir(exist_ok=True)
    data=read(ROOT/'data.json'); package=read(PKG)
    assert 'transferResearch' not in data, 'This research registry is already integrated'
    for name in ['data.json','app.js','Data.xlsx']:
        dest=OUT/'backup'/name
        if dest.exists():
            if name!='app.js': assert digest(dest)==digest(ROOT/name), 'Baseline changed: '+name
        else: shutil.copy2(ROOT/name,dest)
    used={sid for r in package['projects']+package['studies'] for sid in r['source_ids']}
    bylocal={s['id']:s for s in package['sources']}; parents={p['id']:p for p in package['projects']}
    fields=list(data['sources'][0]); added=[]; mapping=[]; source_map={}
    next_id=max(int(s['Source_ID'][1:]) for s in data['sources'])+1
    # Resolve alternate landing pages / PDFs for the same DOI as one work,
    # including old records whose DOI cell is empty but whose URL is known.
    doi_anchors={}
    for sid,doi in DOIS.items():
        if sid not in used:continue
        for old in data['sources']:
            if norm(doi)==norm(old.get('DOI')) or urlkey(bylocal[sid]['url']) in {urlkey(old.get(k)) for k in ['URL','Canonical_URL','PDF_URL']}:
                doi_anchors[norm(doi)]=old; break
    for sid in sorted(used,key=lambda s:int(s[1:])):
        source=bylocal[sid]; doi=DOIS.get(sid)
        match=doi_anchors.get(norm(doi)) if doi else None; reason='DOI / alternate URL' if match else None
        for old in data['sources']+added:
            if match:break
            if doi and norm(doi)==norm(old.get('DOI')): match=old; reason='DOI'; break
            if urlkey(source['url']) in {urlkey(old.get(k)) for k in ['URL','Canonical_URL','PDF_URL']}:
                match=old; reason='URL'; break
            if norm(source['title']) in {norm(old.get('Title')),norm(old.get('Title_EN'))}:
                match=old; reason='title'; break
        if match:
            source_map[sid]=match['Source_ID']; mapping.append({'research_source_id':sid,'source_id':match['Source_ID'],'action':'reuse','match':reason}); continue
        linked=[r for r in package['studies'] if sid in r['source_ids']]
        pids={p['id'] for p in package['projects'] if sid in p['source_ids']}
        for r in linked: pids.update(r['project_ids'])
        ordered=sorted(pids,key=lambda p:int(p[1:])); ps=[parents[p] for p in ordered]
        typ=source['type']; is_academic=sid in ACADEMIC
        stype='مقاله / پژوهش علمی' if is_academic else 'رکورد پایان‌نامه' if sid=='S35' else 'گزارش پژوهشی / نهادی' if sid in {'S42','S43','S45'} else 'خبر / اعلام مقام مسئول' if 'statement' in typ else 'رزومه / سابقه مشاوره' if 'cv' in typ else 'معرفی پروژه / مطالعات مشاور'
        date=str(source.get('date') or '').translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789'))
        exact=re.fullmatch(r'(\d{4})(?:[-/](\d{2})(?:[-/](\d{2}))?)?',date)
        dg=dp=year=None
        if exact:
            year=int(exact[1]); date=date.replace('/','-')
            if year>1700: dg=date
            else: dp=date
        pdf=source['url'] if urlsplit(source['url']).path.lower().endswith('.pdf') else None
        partial=any(x in source['access'] for x in ['نمایه','ناموفق','جلد و فهرست','چکیده','رکورد','متادیتا'])
        primary=typ.startswith('primary_') or typ in {'government_research_report_mirror','government_document_mirror','publication_pdf_mirror'}
        authors='؛ '.join(dict.fromkeys(r['performer'] for r in linked)) or None
        r=dict.fromkeys(fields)
        r.update(Source_ID=f'S{next_id:03}',Title=source['title'],Date_Gregorian=dg,Date_Persian=dp,Year=year,
                 Language='English' if sid in ENGLISH else 'فارسی',Source_Type=stype,
                 Publisher=urlsplit(source['url']).netloc.removeprefix('www.'),Author=authors,
                 Organization='؛ '.join(dict.fromkeys(x['commissioner_or_funder'] for x in linked)) or None,
                 Province='؛ '.join(dict.fromkeys(p for x in ps for p in x['provinces'])),
                 Project='؛ '.join(LABELS[p] for p in ordered),Basin_Source='؛ '.join(dict.fromkeys(x['origin'] for x in ps)) or None,
                 Basin_Destination='؛ '.join(dict.fromkeys(x['destination'] for x in ps)) or None,
                 Topic='مطالعات و تحقیقات انتقال آب',Position_Stance='تحلیلی / خنثی',
                 Summary=' '.join(dict.fromkeys(x['scope'] for x in linked)) or source['access'],
                 URL=source['url'],Canonical_URL=source['url'],DOI=doi,
                 Reliability='B' if primary and not partial else 'C',Primary_or_Secondary='Primary' if primary else 'Secondary',
                 Full_Text='Partial' if partial else 'Yes',PDF_Status='Direct Link Available' if pdf else 'Landing Page Only',PDF_URL=pdf,
                 Notes=f"مرور انتقال آب {TODAY}؛ شناسه شاهد {sid}. عنوان توصیفی مرور است. تاریخ شاهد: {source.get('date') or 'نامشخص'}. دسترسی: {source['access']} وضعیت طرح به تاریخ شاهد وابسته است؛ کارفرمای نامشخص و انتساب از رزومه، قرارداد رسمی محسوب نمی‌شود.")
        added.append(r); source_map[sid]=r['Source_ID']; next_id+=1
        mapping.append({'research_source_id':sid,'source_id':r['Source_ID'],'action':'add'})
        group='academicIds' if is_academic else 'thesisIds' if sid=='S35' else 'legalIds' if sid in {'S42','S45'} else 'newsIds' if 'statement' in typ else 'socialIds' if typ=='company_social_profile' else None
        if group: data['collections'][group].append(r['Source_ID'])
    data['sources'].extend(added)
    for s in added:
        data['timeline'].append({k:s.get(k) for k in data['timeline'][0]})
        if s['PDF_URL']: data['pdfIndex'].append({k:s.get(k) for k in data['pdfIndex'][0]})
    projects=copy.deepcopy(package['projects']); studies=copy.deepcopy(package['studies'])
    for p in projects:
        p['research_id']=p['id']; p['id']='WT-'+p['id']; p['project_tag']=LABELS[p['research_id']]
        p['source_ids']=list(dict.fromkeys(source_map[x] for x in p['source_ids']))
    for s in studies:
        s['research_id']=s['id']; s['id']='WT-'+s['id']; s['project_ids']=['WT-'+p for p in s['project_ids']]
        s['source_ids']=list(dict.fromkeys(source_map[x] for x in s['source_ids']))
        s['cooperation_type']='international' if re.match(r'^ایران[–-]',s['cooperation']) else 'unconfirmed' if s['research_id'] in {'R01','R04','R20'} else 'domestic'
    data['transferProjects']=projects
    data['transferResearch']={k:copy.deepcopy(package[k]) for k in ['category','research_date','geographic_scope','method','counts','leads']}
    data['transferResearch'].update(studies=studies,source_ids=list(dict.fromkeys(source_map.values())),source_mapping=source_map,
        source_evidence=[dict(s,source_id=source_map.get(s['id']),included=s['id'] in used) for s in package['sources']],
        counting_note='۳۰ رکورد مطالعاتی، پژوهشی و مشاوره‌ای؛ شمار قراردادهای مستقل رسمی نیست. ۵ سرنخ تأییدنشده جدا نگهداری می‌شود.')
    data['collections']['transferStudyIds']=list(dict.fromkeys(sid for s in studies for sid in s['source_ids']))
    names={p['Project'] for p in data['projects']}
    for name in LABELS.values():
        if name not in names: data['projects'].append({'Project':name,'Source_Count':0}); names.add(name)
    for p in data['projects']: p['Source_Count']=sum(p['Project'] in (s.get('Project') or '') for s in data['sources'])
    count=len(data['sources']); last=data['sources'][-1]['Source_ID']; dash=data['meta']['dashboard']
    data['meta'].update(generatedAt=TODAY,totalSources=count,lastSourceId=last,geographicScope=package['geographic_scope'])
    data['transferResearch']['counts']['integrated_evidence_sources']=len(set(source_map.values()))
    dash.update({'Total sources':count,'Last Source ID':last,'Transfer projects catalogued':len(projects),'Transfer study/research records':len(studies),
                 'Transfer review evidence sources':len(set(source_map.values())),'Transfer review sources added':len(added),
                 'Persian-language sources':sum('فارسی' in (s.get('Language') or '') for s in data['sources']),
                 'English-language sources':sum(any(t in (s.get('Language') or '') for t in ['English','انگلیسی']) for s in data['sources']),
                 'Sources with direct PDF links':sum(s.get('PDF_Status')=='Direct Link Available' for s in data['sources']),
                 'Locally retrieved PDFs':sum(bool(s.get('Local_PDF')) for s in data['sources'])})
    for grade in 'ABCD':dash['Reliability '+grade]=sum(s.get('Reliability')==grade for s in data['sources'])
    search={'Query':'انتقال آب، خوزستان، اصفهان، چهارمحال و بختیاری، کهگیلویه و بویراحمد؛ Water transfer; studies; consultancy; feasibility',
            'Language':'فارسی / English','Database / Search Engine':'Web search; official project websites; university repositories',
            'Project':'حوضه‌های زاینده‌رود، کارون و استان‌های مرتبط','Source Type':'طرح / مطالعات / پژوهش / مشاوره',
            'Search Date':TODAY,'Number of useful results':len(added),
            'Notes':f"۱۷ طرح و ۳۰ رکورد مطالعه؛ {len(added)} منبع جدید، {sum(m['action']=='reuse' for m in mapping)} شاهد با استفاده از رکورد موجود؛ ۵ سرنخ جدا. دسترسی مستقیم Google ناموفق بود. جزئیات: research/water-transfer-review-2026-10-11.md"}
    change={'Change_ID':f"C{max(int(c['Change_ID'][1:]) for c in data['changeLog'])+1:03}",'Date':TODAY,'Action':'TRANSFER STUDIES REGISTRY MERGE',
            'Source_ID / Item':f"{added[0]['Source_ID']}–{last}" if added else 'transferResearch',
            'Reason':'User requested main dataset and dashboard integration',
            'Details':f"{len(added)} sources added; {count-len(added)} old sources preserved; 17 project profiles; 30 research records; 5 pending leads separated."}
    data['searchLog'].append(search); data['changeLog'].append(change)
    wb=load_workbook(ROOT/'Data.xlsx'); ops={}; cells=[]; new_sheets={}
    def append(name, records):
        if not records:return
        headers=[c.value for c in wb[name][1]]
        ops[name]={'start_row':wb[name].max_row+1,'columns':len(headers),'rows':[[r.get(k) for k in headers] for r in records]}
    for sheet in ['ALL_SOURCES','NEW_SOURCES','TIMELINE']:append(sheet,added)
    for group,sheet in [('academicIds','ACADEMIC'),('thesisIds','THESES'),('legalIds','OFFICIAL_LEGAL'),('newsIds','NEWS_MEDIA'),('socialIds','SOCIAL_MEDIA')]:
        append(sheet,[s for s in added if s['Source_ID'] in data['collections'][group]])
    append('PDF_INDEX',[s for s in added if s['PDF_URL']]); append('SEARCH_LOG',[search]); append('CHANGE_LOG',[change])
    for sheet, pairs in [('DASHBOARD',dash.items()),('PROJECTS',[(p['Project'],p['Source_Count']) for p in data['projects']])]:
        ws=wb[sheet]; lookup={ws.cell(i,1).value:i for i in range(2,ws.max_row+1)}; n=ws.max_row
        for key,value in pairs:
            if key not in lookup:n+=1;lookup[key]=n
            row=lookup[key]
            for col,v in [('A',key),('B',value)]:
                if ws[f'{col}{row}'].value!=v:cells.append({'sheet':sheet,'cell':f'{col}{row}','value':v})
    for cell,v in [('A1','Water Transfer — FOUR PROVINCES & STUDIES UPDATE'),('B2',count),('B3',f'S001–{last} (continuous, no gaps)'),('B4',f'Through {TODAY}')]:cells.append({'sheet':'README','cell':cell,'value':v})
    def tab(name, records, columns):
        rows=[[label for key,label in columns]]
        for r in records:
            row=[]
            for key,label in columns:
                v=r.get(key); row.append('؛ '.join(map(str,v)) if isinstance(v,list) else v)
            rows.append(row)
        new_sheets[name]={'rows':rows,'columns':len(columns)}
    projectcols=[('id','شناسه طرح'),('name_fa','نام طرح'),('origin','مبدأ'),('destination','مقصد'),('provinces','استان‌های طرح'),('downstream_or_shared','پایین‌دست / حوضه مشترک'),('transfer_type','نوع انتقال'),('status_evidence','شاهد وضعیت'),('status_as_of','تاریخ شاهد'),('source_ids','شناسه منابع')]
    studycols=[('id','شناسه مطالعه'),('title_fa','عنوان مطالعه'),('subcategory','زیرگروه'),('project_ids','طرح‌های مرتبط'),('performer','مجری / مشاور'),('commissioner_or_funder','کارفرما / حامی'),('period','دوره'),('cooperation','همکاری'),('scope','موضوع و حدود مطالعه'),('evidence','نوع شاهد'),('source_ids','شناسه منابع')]
    tab('TRANSFER_PROJECTS',projects,projectcols); tab('TRANSFER_STUDIES',studies,studycols)
    for name,records in [('TRANSFER_PROJECTS',projects),('TRANSFER_STUDIES',studies)]:
        new_sheets[name]['rows'][0].append('پیوند منابع'); new_sheets[name]['columns']+=1
        srcby={s['Source_ID']:s for s in data['sources']}
        for row,r in zip(new_sheets[name]['rows'][1:],records):row.append('\n'.join(srcby[sid]['URL'] for sid in r['source_ids']))
    save(OUT/'staged/data.json',data);save(OUT/'workbook_changes.json',{'append':ops,'cells':cells,'new_sheets':new_sheets})
    save(OUT/'manifest.json',{'date':TODAY,'before_count':count-len(added),'after_count':count,'added':len(added),'mapping':mapping,
        'pending_leads':package['leads'],'excluded_evidence':[s['id'] for s in package['sources'] if s['id'] not in used],
        'baseline_sha256':{n:digest(ROOT/n) for n in ['data.json','app.js','Data.xlsx']}})
    print(json.dumps({'added':len(added),'reused':sum(m['action']=='reuse' for m in mapping),'total':count,'mapping':mapping},ensure_ascii=False,indent=2))

if __name__=='__main__': prepare()
