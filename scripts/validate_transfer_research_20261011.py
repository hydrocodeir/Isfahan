"""Validate the staged registry and native workbook preservation, then apply it."""
import copy, json, pathlib, shutil
from openpyxl import load_workbook
from integrate_transfer_research_20261011 import ROOT, OUT, read, save, digest, urlkey, norm

def validate():
    data=read(OUT/'staged/data.json'); before=read(OUT/'backup/data.json'); manifest=read(OUT/'manifest.json'); edits=read(OUT/'workbook_changes.json')
    n=len(before['sources']); ids={s['Source_ID'] for s in data['sources']}; projects={p['id'] for p in data['transferProjects']}
    assert data['sources'][:n]==before['sources']
    assert len(ids)==len(data['sources'])==1042
    assert [s['Source_ID'] for s in data['sources']]==[f'S{i:03}' for i in range(1,len(ids)+1)]
    for group,members in data['collections'].items(): assert len(members)==len(set(members)) and set(members)<=ids, group
    assert len(data['transferProjects'])==17 and len(data['transferResearch']['studies'])==30
    for r in data['transferProjects']+data['transferResearch']['studies']:
        assert set(r['source_ids'])<=ids
        if 'project_ids' in r:assert set(r['project_ids'])<=projects
    assert {s['cooperation_type'] for s in data['transferResearch']['studies']}=={'domestic','international','unconfirmed'}
    assert len(data['transferResearch']['leads'])==5
    assert manifest['excluded_evidence']==['S39','S40','S41']
    for key in ['timeline','pdfIndex','searchLog','changeLog']:
        assert data[key][:len(before[key])]==before[key], key
    for key in ['timeline','pdfIndex']:
        assert len(data[key])==len({r['Source_ID'] for r in data[key]})
    for s in data['sources'][n:]:
        for old in before['sources']:
            assert not s['DOI'] or norm(s['DOI'])!=norm(old.get('DOI')), ('DOI duplicate',s['Source_ID'],old['Source_ID'])
            assert urlkey(s['URL']) not in {urlkey(old.get(k)) for k in ['URL','Canonical_URL','PDF_URL']},('URL duplicate',s['Source_ID'],old['Source_ID'])
    app=(ROOT/'app.js').read_text('utf-8'); marker='const OFFLINE_FALLBACK = '; start=app.index(marker)+len(marker);end=app.index(';\n\nconst $',start)
    staged=app[:start]+json.dumps(data,ensure_ascii=False,separators=(',',':'))+app[end:]
    (OUT/'staged/app.js').write_text(staged,'utf-8')
    assert json.loads(staged[start:staged.index(';\n\nconst $',start)])==data
    wb=load_workbook(OUT/'staged/Data.xlsx'); old=load_workbook(OUT/'backup/Data.xlsx')
    assert wb.sheetnames==old.sheetnames+['TRANSFER_PROJECTS','TRANSFER_STUDIES']
    allowed={(e['sheet'],e['cell']):e['value'] for e in edits['cells']}; checked=0
    for ws in old:
        new=wb[ws.title]
        for row in ws:
            for c in row:
                nc=new[c.coordinate]
                assert nc.value==allowed.get((ws.title,c.coordinate),c.value),(ws.title,c.coordinate,'value')
                assert nc._style==c._style,(ws.title,c.coordinate,'style')
                assert (nc.hyperlink.target if nc.hyperlink else None)==(c.hyperlink.target if c.hyperlink else None),(ws.title,c.coordinate,'hyperlink')
                assert (nc.comment.text if nc.comment else None)==(c.comment.text if c.comment else None),(ws.title,c.coordinate,'comment')
                checked+=1
        assert ws.freeze_panes==new.freeze_panes and set(ws.tables)==set(new.tables)
        assert list(ws.merged_cells.ranges)==list(new.merged_cells.ranges)
        assert str(ws.protection)==str(new.protection)
        assert ws.data_validations==new.data_validations
        assert {k:dict(v) for k,v in ws.column_dimensions.items()}=={k:dict(v) for k,v in new.column_dimensions.items()}
        for k,v in ws.row_dimensions.items():assert dict(v)==dict(new.row_dimensions[k]),(ws.title,k,'row dimensions')
    for sheet,op in edits['append'].items():
        actual=[[wb[sheet].cell(r,c).value for c in range(1,op['columns']+1)] for r in range(op['start_row'],op['start_row']+len(op['rows']))]
        expected=[[None if v=='' else v for v in row] for row in op['rows']]
        assert actual==expected, ('append mismatch',sheet)
    for sheet,op in edits['new_sheets'].items():
        ws=wb[sheet]; assert ws.max_row==len(op['rows']) and ws.max_column==op['columns']
        assert [list(row) for row in ws.values]==[[None if v=='' else v for v in row] for row in op['rows']]
        assert ws.freeze_panes=='C2' and ws.sheet_view.showGridLines is False
        assert len(ws.tables)==1
    assert wb['ALL_SOURCES'].max_row==1043
    assert wb['ALL_SOURCES'].tables['AllSourcesClean'].ref=='A1:AK1043'
    for e in edits['cells']:assert wb[e['sheet']][e['cell']].value==e['value']
    source_values=[list(r) for r in wb['ALL_SOURCES'].values]
    headers=source_values[0]
    assert source_values[n+1:]==[[None if s.get(k)=='' else s.get(k) for k in headers] for s in data['sources'][n:]]
    legacy_differences=[(i,k,v,s.get(k)) for i,(row,s) in enumerate(zip(source_values[1:n+1],data['sources'][:n])) for k,v in zip(headers,row) if v!=s.get(k)]
    old_values=[list(r) for r in old['ALL_SOURCES'].values]
    baseline_differences=[(i,k,v,s.get(k)) for i,(row,s) in enumerate(zip(old_values[1:],before['sources'])) for k,v in zip(headers,row) if v!=s.get(k)]
    assert legacy_differences==baseline_differences, 'Pre-existing Excel/JSON differences changed'
    wb.close();old.close()
    result={'status':'PASS','date':'2026-10-11','before_sources':n,'after_sources':len(ids),'added_sources':30,
            'reused_research_evidence':12,'project_profiles':17,'study_records':30,'pending_leads':5,
            'existing_records_preserved':n,'existing_excel_cells_checked':checked,'existing_style_changes':0,
            'json_offline_equal':True,'new_excel_json_rows_equal':True,
            'legacy_excel_json_differences_preserved':len(legacy_differences)}
    save(OUT/'validation.json',result);print(json.dumps(result,ensure_ascii=False,indent=2))

def apply():
    assert read(OUT/'validation.json')['status']=='PASS'
    assert read(OUT/'browser_validation.json')['status']=='PASS'
    manifest=read(OUT/'manifest.json')
    for name,sha in manifest['baseline_sha256'].items():assert digest(ROOT/name)==sha, 'File changed during staging: '+name
    for name in ['Data.xlsx','data.json','app.js']:shutil.copy2(OUT/'staged'/name,ROOT/name)
    print('Applied validated dataset, workbook and offline data.')

if __name__=='__main__': (apply if '--apply' in __import__('sys').argv else validate)()
