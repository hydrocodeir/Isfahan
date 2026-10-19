"""Stage the reviewed refresh and an explicit workbook edit plan."""
import copy, hashlib, json, pathlib, re, shutil, sys
from openpyxl import load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[1];OUT=ROOT/'reports/deep_refresh_2026_10_19';TODAY='2026-10-19'
def read(p):return json.loads(p.read_text('utf-8-sig'))
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n','utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def unique(a):return list(dict.fromkeys(a))
def fa(n):return str(n).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
def prepare():
 for d in ['backup','staged']:(OUT/d).mkdir(exist_ok=True)
 before=read(ROOT/'data.json');data=copy.deepcopy(before);pkg=read(OUT/'curated.json')
 assert data['meta']['generatedAt']!=TODAY,'Already integrated'
 for name in ['data.json','app.js','Data.xlsx']:
  dest=OUT/'backup'/name
  if dest.exists():assert sha(dest)==sha(ROOT/name),'Baseline changed: '+name
  else:shutil.copy2(ROOT/name,dest)
 oldby={s['Source_ID']:s for s in data['sources']};nextid=max(int(s[1:]) for s in oldby)+1;mapping={};added=[];enriched=[]
 for c in pkg['sources']:
  if c.get('reuse_source_id'):
   sid=c['reuse_source_id'];mapping[c['key']]=sid
   if sid=='S097':
    r=oldby[sid];fields={'DOI':c['record']['DOI'],'Author':c['record']['Author'],'Date_Gregorian':c['record']['Date_Gregorian'],
      'Full_Text':'Partial','Notes':(r.get('Notes') or '')+f"\nمرور {TODAY}: هویت مقاله با صفحه ناشر {c['record']['URL']} و DOI تطبیق داده شد. عنوان فارسی: {c['record']['Title']}. تاریخ از متادیتای ناشر تکمیل شد. پیوند قدیمی برای حفظ سابقه نگه داشته شد."}
    fields={k:v for k,v in fields.items() if r.get(k)!=v};r.update(fields);enriched.append({'source_id':sid,'fields':fields})
   continue
  r=copy.deepcopy(c['record']);r['Source_ID']=f'S{nextid:03}';mapping[c['key']]=r['Source_ID'];nextid+=1;added.append(r)
 data['sources'].extend(added)
 parent={p['id']:copy.deepcopy(p) for p in data['transferProjects']}
 for p in pkg['new_projects']:
  p=copy.deepcopy(p);p['source_ids']=[mapping[k] for k in p.pop('source_keys')]
  if p['id']=='WT-T18':p['source_ids'].append('S216')
  parent[p['id']]=p
 for c in pkg['sources']:
  sid=mapping[c['key']];r=next(s for s in data['sources'] if s['Source_ID']==sid)
  if sid not in oldby:
   ps=[parent[p] for p in c['project_ids']]
   r['Project']=r.get('Project') or '؛ '.join(p['project_tag'] for p in ps) or ('ایران مرکزی؛ انتقال آب دریا (زمینه‌ای)' if c['key']=='A20' else 'حوضه کارون؛ پژوهش هیدرولوژیک زمینه‌ای' if c['key']=='J02' else 'زاینده‌رود؛ حکمرانی و سیاست انتقال')
   r['Province']=r.get('Province') or '؛ '.join(unique(v for p in ps for v in p['provinces'])) or 'ملی / زمینه‌ای'
   r['Basin_Source']=r.get('Basin_Source') or '؛ '.join(unique(p['origin'] for p in ps)) or ('حوضه کارون' if c['key']=='J02' else None)
   r['Basin_Destination']=r.get('Basin_Destination') or '؛ '.join(unique(p['destination'] for p in ps)) or None
 data['transferProjects']=list(parent.values())
 st=data['transferResearch']['studies'];oldstudies=copy.deepcopy(st);by_study={s['id']:s for s in st}
 for update in pkg['study_updates']:
  s=by_study[update['id']];s.update(update['patch']);s['source_ids']=unique(s['source_ids']+[mapping[k] for k in update['source_keys']])
  s['scope']+=' در مرور ۲۰۲۶/۱۰/۱۹، سال و کارفرما از صفحه مستقل مشاور تکمیل شد؛ اصل قرارداد مشاهده نشده است.'
 newstudies=[];ni=max(int(s['research_id'][1:]) for s in st)+1
 for s in pkg['new_studies']:
  s=copy.deepcopy(s);s['research_id']=f'R{ni:02}';s['id']='WT-'+s['research_id'];ni+=1;s['category']='studies-and-research'
  s['source_ids']=unique(mapping[k] for k in s.pop('source_keys'))
  if s['title_fa'].startswith('تغییر زیستگاه'):s['provinces']=['چهارمحال و بختیاری']
  st.append(s);newstudies.append(s)
 data['transferResearch']['leads'].append({'id':'L06','title':'بررسی تکمیلی انتقال خرسان توسط دانشگاه تهران','reason':'در اظهارات نماینده یزد نقل شده؛ اصل نامه سفارش، شرح خدمات و قرارداد دانشگاه مشاهده نشده است.','needed':'نامه سازمان محیط زیست، واحد مجری، تاریخ و گزارش مطالعه','source_ids':[mapping['N02']],'source_scope':'canonical'})
 tr=data['transferResearch'];tr['research_date']=TODAY
 tr['counting_note']=f"{fa(len(st))} رکورد مطالعاتی، پژوهشی و مشاوره‌ای؛ شمار قراردادهای مستقل رسمی نیست. {fa(len(tr['leads']))} سرنخ تأییدنشده جدا نگهداری می‌شود."
 tr['refreshes']=tr.get('refreshes',[])+[{'date':TODAY,'added_sources':len(added),'new_study_records':len(newstudies),'enriched_studies':2,'enriched_sources':len(enriched),'new_project_profiles':len(pkg['new_projects']),'report':'reports/deep_refresh_2026_10_19/report.md','limitations':pkg['limitations']}]
 tr['source_ids']=unique(tr['source_ids']+list(mapping.values()))
 tr['source_evidence'].extend({'id':'20261019-'+c['key'],'title':c['record']['Title'],'url':c['record']['URL'],
   'date':c['record'].get('Date_Gregorian') or c['record'].get('Date_Persian'),'access':c['record']['Notes'],
   'source_id':mapping[c['key']],'included':True,'retrieved_on':TODAY} for c in pkg['sources'])
 tr['counts'].update(parent_projects=len(parent),study_and_related_consultancy_records=len(st),sources=len(tr['source_evidence']),
  unresolved_leads=len(tr['leads']),integrated_evidence_sources=len(tr['source_ids']))
 for s in added:
  key='thesisIds' if 'رساله' in s['Source_Type'] else 'academicIds' if any(t in s['Source_Type'] for t in ['مقاله','کنفرانس']) else 'parliamentIds' if s.get('Interviewee') else 'newsIds' if 'خبر' in s['Source_Type'] else None
  if key:data['collections'][key].append(s['Source_ID'])
  if s.get('Interviewee'):data['collections']['newsIds'].append(s['Source_ID'])
 data['collections']['transferStudyIds']=unique(i for s in st for i in s['source_ids'])
 for k in data['collections']:data['collections'][k]=unique(data['collections'][k])
 for e in enriched:
  sid=e['source_id'];s=oldby[sid]
  for r in data['timeline']:
   if r['Source_ID']==sid:
    for k in r:r[k]=s.get(k)
 for s in added:
  data['timeline'].append({k:s.get(k) for k in data['timeline'][0]})
  if s['PDF_URL']:data['pdfIndex'].append({k:s.get(k) for k in data['pdfIndex'][0]})
 existing_names={p['Project'] for p in data['projects']}
 for p in parent.values():
  if p['project_tag'] not in existing_names:data['projects'].append({'Project':p['project_tag'],'Source_Count':0});existing_names.add(p['project_tag'])
 for p in data['projects']:
  related={s['Source_ID'] for s in data['sources'] if p['Project'] in (s.get('Project') or '')}
  profile=next((r for r in parent.values() if r['project_tag']==p['Project']),None)
  if profile:
   related.update(profile['source_ids'])
   related.update(sid for s in st if profile['id'] in s['project_ids'] for sid in s['source_ids'])
  p['Source_Count']=len(related)
 count=len(data['sources']);last=data['sources'][-1]['Source_ID'];dash=data['meta']['dashboard']
 data['meta'].update(generatedAt=TODAY,totalSources=count,lastSourceId=last)
 dash.update({'Total sources':count,'Last Source ID':last,'Transfer projects catalogued':len(parent),'Transfer study/research records':len(st),
  'Transfer review evidence sources':len(tr['source_ids']),'Deep refresh sources added 2026-10-19':len(added),
  'Deep refresh study records added 2026-10-19':len(newstudies),
  'Persian-language sources':sum('فارسی' in (s.get('Language') or '') for s in data['sources']),
  'English-language sources':sum(any(x in (s.get('Language') or '') for x in ['English','انگلیسی']) for s in data['sources']),
  'Sources with direct PDF links':sum(s.get('PDF_Status')=='Direct Link Available' for s in data['sources']),
  'Locally retrieved PDFs':sum(bool(s.get('Local_PDF')) for s in data['sources'])})
 for g in 'ABCD':dash['Reliability '+g]=sum(s.get('Reliability')==g for s in data['sources'])
 web=read(OUT/'web_search_log.json')+read(OUT/'web_search_log_additional.json');registry=read(OUT/'registry_search.json')
 searches=[]
 for batch in web:
  for q in batch['queries']:searches.append({'Query':q,'Language':'فارسی' if re.search('[\u0600-\u06ff]',q) else 'English','Database / Search Engine':'Web search','Project':'چهار استان و طرح‌های مرتبط','Source Type':'مقاله / مطالعه / سند / خبر','Search Date':TODAY,'Number of useful results':None,'Notes':'پرس‌وجوی کشف؛ تعداد نتیجه مفید جداگانه برای هر Query محاسبه نشده؛ گزارش و تصمیم‌های ادغام در reports/deep_refresh_2026_10_19/.'})
 for batch in registry:searches.append({'Query':batch['query'],'Language':'فارسی' if re.search('[\u0600-\u06ff]',batch['query']) else 'English','Database / Search Engine':batch['engine'],'Project':'چهار استان و طرح‌های مرتبط','Source Type':'ادبیات علمی','Search Date':TODAY,'Number of useful results':None,'Notes':f"نتایج بازیابی‌شده پیش از پالایش: {len(batch.get('items',[]))}؛ "+(batch.get('error') or 'موفق؛ نتیجه کشف به معنی منبع پذیرفته‌شده نیست')})
 summary={'Query':'مرور عمیق انتقال آب — تطبیق نهایی منابع، مطالعات و پروژه‌ها','Language':'فارسی / English','Database / Search Engine':'Curated merge','Project':'زاینده‌رود؛ کارون؛ چهار استان و طرح‌های مرتبط','Source Type':'مرور و ادغام','Search Date':TODAY,'Number of useful results':len(added),'Notes':f'{len(added)} منبع جدید، {len(enriched)} منبع تکمیل‌شده، {len(newstudies)} رکورد مطالعه جدید، ۲ مطالعه تکمیل‌شده و ۸ پروفایل طرح اضافه شد. شش PDF جدید بازیابی شد.'}
 searches.append(summary);data['searchLog'].extend(searches)
 change={'Change_ID':f"C{max(int(c['Change_ID'][1:]) for c in data['changeLog'])+1:03}",'Date':TODAY,'Action':'DEEP SOURCE REFRESH AND STUDIES UPDATE','Source_ID / Item':f"{added[0]['Source_ID']}–{last}; S097",'Reason':'User requested deep search and addition of new sources','Details':f'{len(added)} new sources; S097 enriched; {len(newstudies)} research records; 2 consultancy records enriched; 8 project profiles; six primary PDFs. Full decisions and access logs retained.'}
 data['changeLog'].append(change)
 save(OUT/'staged/data.json',data)
 app=(ROOT/'app.js').read_text('utf-8');mark='const OFFLINE_FALLBACK = ';start=app.index(mark)+len(mark);end=app.index(';\n\nconst $',start)
 staged_app=app[:start]+json.dumps(data,ensure_ascii=False,separators=(',',':'))+app[end:]
 if 'const studyProjects =' not in staged_app:
  anchor='  aliases.forEach(([name, aa]) => { if (aa.some(a => raw.includes(norm(a)))) tags.push(name); });'
  assert anchor in staged_app
  staged_app=staged_app.replace(anchor,anchor+'\n  const studyProjects = (state.data?.transferResearch?.studies || []).filter(study => study.source_ids.includes(s.Source_ID)).flatMap(study => study.project_ids);\n  (state.data?.transferProjects || []).forEach(p => {\n    if (raw.includes(norm(p.project_tag)) || p.source_ids.includes(s.Source_ID) || studyProjects.includes(p.id)) tags.push(p.project_tag);\n  });',1)
 if 'const availableProjects =' not in staged_app:
  anchor='function populateFilters() {\n  const s = state.sources;'
  assert anchor in staged_app
  staged_app=staged_app.replace(anchor,anchor+'\n  const availableProjects = uniq([...canonicalProjects, ...(state.data.projects || []).map(p => p.Project), ...s.flatMap(x => x._projectTags)]);',1)
  staged_app=staged_app.replace('), canonicalProjects, "همه پروژه‌ها");','), availableProjects, "همه پروژه‌ها");')
 (OUT/'staged/app.js').write_text(staged_app,'utf-8')
 wb=load_workbook(ROOT/'Data.xlsx');operations={};cells=[]
 def append(sheet,records):
  if not records:return
  heads=[c.value for c in wb[sheet][1]]
  operations[sheet]={'start_row':wb[sheet].max_row+1,'columns':len(heads),'rows':[[r.get(k) for k in heads] for r in records]}
 for name in ['ALL_SOURCES','NEW_SOURCES','TIMELINE']:append(name,added)
 for group,sheet in [('academicIds','ACADEMIC'),('thesisIds','THESES'),('parliamentIds','PARLIAMENT'),('newsIds','NEWS_MEDIA')]:append(sheet,[s for s in added if s['Source_ID'] in data['collections'][group]])
 append('PDF_INDEX',[s for s in added if s['PDF_URL']]);append('SEARCH_LOG',searches);append('CHANGE_LOG',[change])
 for e in enriched:
  for ws in wb:
   if ws.cell(1,1).value!='Source_ID':continue
   heads={c.value:c.column_letter for c in ws[1]}
   for row in range(2,ws.max_row+1):
    if ws.cell(row,1).value==e['source_id']:
     for key,v in e['fields'].items():
      if key in heads:cells.append({'sheet':ws.title,'cell':heads[key]+str(row),'value':v})
 srcby={s['Source_ID']:s for s in data['sources']}
 def projectrow(p):return [p['id'],p['name_fa'],p['origin'],p['destination'],'؛ '.join(p['provinces']),'؛ '.join(p['downstream_or_shared']),p['transfer_type'],p['status_evidence'],p['status_as_of'],'؛ '.join(p['source_ids']),'\n'.join(srcby[i]['URL'] for i in p['source_ids'])]
 def studyrow(s):return [s['id'],s['title_fa'],s['subcategory'],'؛ '.join(s['project_ids']),s['performer'],s['commissioner_or_funder'],s['period'],s['cooperation'],s['scope'],s['evidence'],'؛ '.join(s['source_ids']),'\n'.join(srcby[i]['URL'] for i in s['source_ids'])]
 operations['TRANSFER_PROJECTS']={'start_row':wb['TRANSFER_PROJECTS'].max_row+1,'columns':11,'rows':[projectrow(parent[p['id']]) for p in pkg['new_projects']]}
 operations['TRANSFER_STUDIES']={'start_row':wb['TRANSFER_STUDIES'].max_row+1,'columns':12,'rows':[studyrow(s) for s in newstudies]}
 for updated in pkg['study_updates']:
  ws=wb['TRANSFER_STUDIES'];row=next(r for r in range(2,ws.max_row+1) if ws.cell(r,1).value==updated['id'])
  for col,v in enumerate(studyrow(by_study[updated['id']]),1):
   if ws.cell(row,col).value!=v:cells.append({'sheet':ws.title,'cell':ws.cell(row,col).coordinate,'value':v})
 for name,pairs in [('DASHBOARD',dash.items()),('PROJECTS',[(p['Project'],p['Source_Count']) for p in data['projects']])]:
  ws=wb[name];lookup={ws.cell(r,1).value:r for r in range(2,ws.max_row+1)};lastrow=ws.max_row
  for key,v in pairs:
   if key not in lookup:lastrow+=1;lookup[key]=lastrow
   row=lookup[key]
   for col,value in [('A',key),('B',v)]:
    if ws[f'{col}{row}'].value!=value:cells.append({'sheet':name,'cell':f'{col}{row}','value':value})
 for addr,value in [('A1','Water Transfer — DEEP SOURCE REFRESH'),('B2',count),('B3',f'S001–{last} (continuous, no gaps)'),('B4',f'Through {TODAY}')]:cells.append({'sheet':'README','cell':addr,'value':value})
 save(OUT/'workbook_changes.json',{'append':operations,'cells':cells,'new_sheets':{}})
 manifest={'date':TODAY,'before_sources':len(before['sources']),'after_sources':count,'added_source_ids':[s['Source_ID'] for s in added],
  'source_mapping':mapping,'enriched_sources':enriched,'new_study_ids':[s['id'] for s in newstudies],
  'updated_study_ids':[s['id'] for s in pkg['study_updates']],'new_project_ids':[p['id'] for p in pkg['new_projects']],
  'local_pdfs_added':sum(bool(s['Local_PDF']) for s in added),'web_queries_logged':sum(len(b['queries']) for b in web),
  'registry_queries':len(registry),'registry_successes':sum('items' in b for b in registry),
  'baseline_sha256':{name:sha(ROOT/name) for name in ['data.json','app.js','Data.xlsx']}}
 save(OUT/'manifest.json',manifest)
 print(json.dumps({k:v for k,v in manifest.items() if k not in ['source_mapping','baseline_sha256','enriched_sources']},ensure_ascii=False,indent=2))
if __name__=='__main__':prepare()
