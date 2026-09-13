#!/usr/bin/env python3
"""Apply the final publisher-PDF and correction-note cleanup after link QA."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PDF_URL = "https://jwsd.um.ac.ir/article_46223_254aeaad72ea9fae24715d6f62dc71a5.pdf"
S942_NOTE = "جایگزین رکورد قبلی S942 شد؛ مقاله قبلی درباره تونل انتقال آب کرمان بود و به‌اشتباه به بهشت‌آباد نسبت داده شده بود. متن حاضر مستقیماً پروژه بهشت‌آباد، کارون و زاینده‌رود را نام می‌برد؛ صفحه ناشر اصلی بازیابی نشد."


def main():
    data_path = ROOT / "data.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    by_id = {source["Source_ID"]: source for source in data["sources"]}
    by_id["S096"]["PDF_URL"] = PDF_URL
    by_id["S096"]["Notes"] = "عنوان فارسی، نویسندگان، DOI صحیح (2410-1375، نه صورت وارونه آن)، تاریخ انتشار و صفحات با متن مقاله و رکورد ناشر/OuluREPO تطبیق شد؛ PDF به لینک مستقیم ناشر منتقل شد."
    by_id["S942"]["Notes"] = S942_NOTE
    for row in data["pdfIndex"]:
        if row.get("Source_ID") == "S096":
            row["PDF_URL"] = PDF_URL
    data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    app_path = ROOT / "app.js"
    app = app_path.read_text(encoding="utf-8")
    marker = "const OFFLINE_FALLBACK = "
    start = app.index(marker) + len(marker)
    end = app.index(";\n\nconst $", start)
    app_path.write_text(app[:start] + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + app[end:], encoding="utf-8")

    wb = load_workbook(ROOT / "Data.xlsx")
    for sheet in ("ALL_SOURCES", "ACADEMIC", "NEW_SOURCES"):
        ws = wb[sheet]
        headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
        for row in range(2, ws.max_row + 1):
            source_id = ws.cell(row, 1).value
            if source_id == "S096":
                ws.cell(row, headers["PDF_URL"]).value = PDF_URL
                ws.cell(row, headers["Notes"]).value = by_id["S096"]["Notes"]
            elif source_id == "S942":
                ws.cell(row, headers["Notes"]).value = S942_NOTE
    ws = wb["PDF_INDEX"]
    headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == "S096":
            ws.cell(row, headers["PDF_URL"]).value = PDF_URL
    wb.save(ROOT / "Data.xlsx")


if __name__ == "__main__":
    main()
