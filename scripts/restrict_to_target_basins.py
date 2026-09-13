#!/usr/bin/env python3
"""Remove records outside the Iran Zayandeh-Rud/Karun dashboard scope."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from build_validated_outputs import append_styled
from deep_literature_search import ROOT, REPORTS, normalize, save_json


TODAY = "2026-09-13"
TODAY_FA = "1405/06/22"

TARGET_TERMS = (
    "zayandeh rud", "zayandeh rood", "zayanderud", "zayandehrud", "zayandeh river",
    "karun", "karoun", "gavkhuni", "gavkhouni", "gavkhooni",
    "kuhrang", "kouhrang", "koohrang", "behesht abad", "beheshtabad", "golab",
    "cheshmeh langan", "khodangestan", "khersan", "mandegan", "ben borujen",
    "pelasjan", "marbar", "bideh", "yalan", "pashangan", "goukan",
    "زاینده رود", "کارون", "گاوخونی", "گاخونی", "کوهرنگ", "بهشت آباد", "گلاب",
    "چشمه لنگان", "خدنگستان", "خرسان", "ماندگان", "بن بروجن", "پلاسجان",
    "ماربر", "بیده", "یلان", "پشندگان", "گوکان",
)

WATER_TERMS = (
    "water", "river", "basin", "watershed", "catchment", "aquifer", "groundwater", "hydro",
    "dam", "wetland", "swamp", "playa", "drought", "irrigation", "reservoir", "flood",
    "sediment", "pollution", "contamin", "water quality", "ecolog", "ecosystem", "fish",
    "microplastic", "streamflow", "runoff", "precipitation", "evapotranspiration", "subsidence",
    "water governance", "water conflict", "water scarcity", "water allocation", "water transfer",
    "آب", "رود", "حوضه", "آبریز", "آبخوان", "سد", "تالاب", "باتلاق", "پلایا", "خشکسالی",
    "آبیاری", "مخزن", "سیلاب", "رسوب", "آلود", "کیفیت آب", "بوم", "اکوسیستم", "ماهی",
    "ریزگرد", "فرونشست", "رواناب", "بارش", "تبخیر", "حکمرانی", "تعارض", "بحران آب",
    "تخصیص", "انتقال آب", "منابع آب", "حقابه", "حق آبه",
)

TARGET_FIELDS = ("Title", "Title_EN", "Project", "Basin_Source", "Basin_Destination", "Location")
SUBJECT_FIELDS = TARGET_FIELDS + ("Topic", "Summary", "Key_Claim", "Notes")

FALSE_FRIEND_PATTERNS = (
    "harta karun",  # Indonesian: treasure, not Iran's Karun River
    "archaeological landscape of the western gavkhuni",
    "archaeometallurgical evidence",
    "excavations of beheshtabad archaeological hill",
    "methods of representing implied author and implied reader",
    "protolith of amphibolites",
    "quartzofeldspathic schists",
    "zayanderud eclogites",
    "jurassic subduction initiation",
    "kuh e siah volcano",
    "magmatic arc gavkhuni",
    "geochemical and geological studies of oil shales",
    "khersan glacier territory",
    "alamkuh mountain central alborz",
)


def has_any(text: str, terms: tuple[str, ...]) -> bool:
    return any(term in text for term in terms)


def in_scope(record: dict) -> tuple[bool, str]:
    title_text = normalize(record.get("Title"))
    if has_any(title_text, FALSE_FRIEND_PATTERNS):
        return False, "هم‌نامی جغرافیایی یا موضوع باستان‌شناسی/سنگ‌شناسی نامرتبط با سامانه آبی هدف"
    project_text = normalize(" ".join(str(record.get(field) or "") for field in ("Project", "Basin_Source", "Basin_Destination")))
    target_text = normalize(" ".join(str(record.get(field) or "") for field in TARGET_FIELDS))
    subject_text = normalize(" ".join(str(record.get(field) or "") for field in SUBJECT_FIELDS))
    if not has_any(target_text, TARGET_TERMS):
        return False, "فاقد ارتباط مستقیم با زاینده‌رود، کارون، گاوخونی یا طرح‌های متصل به این دو حوضه"
    # A named dashboard project/basin is intrinsically in scope even when a
    # terse news or social-media title does not repeat the word "water".
    if has_any(project_text, TARGET_TERMS):
        return True, "در فیلد پروژه/حوضه مستقیماً به سامانه هدف متصل است"
    if not has_any(subject_text, WATER_TERMS):
        return False, "نام جغرافیایی هدف ذکر شده، اما موضوع منبع آب، حوضه آبریز یا پیامدهای مرتبط نیست"
    return True, "ارتباط مستقیم مکانی و موضوعی با حوضه‌های هدف"


def write_removed_report(removed: list[tuple[dict, str]]) -> None:
    previous_path = REPORTS / "out_of_scope_removed.json"
    previous = []
    if previous_path.exists():
        previous = json.loads(previous_path.read_text(encoding="utf-8")).get("records") or []
    additions = [{"removal_pass": "semantic false-positive cleanup", "reason": reason, **record} for record, reason in removed]
    combined, seen = [], set()
    for record in previous + additions:
        key = normalize(record.get("DOI")) or normalize(record.get("Canonical_URL") or record.get("URL")) or normalize(record.get("Title"))
        if key in seen:
            continue
        seen.add(key)
        combined.append(record)
    payload = {
        "date": TODAY,
        "criterion": "Direct target-basin/project term plus water/basin-related subject matter",
        "removed_count": len(combined),
        "records": combined,
    }
    save_json(REPORTS / "out_of_scope_removed.json", payload)

    wb = Workbook()
    ws = wb.active
    ws.title = "منابع حذف‌شده"
    ws.sheet_view.rightToLeft = True
    headers = ["شناسه قبلی", "عنوان", "نوع", "پروژه", "موضوع", "زبان", "ناشر", "DOI", "لینک", "علت حذف"]
    ws.append(headers)
    for record in combined:
        reason = record.get("reason")
        ws.append([
            record.get("Source_ID"), record.get("Title"), record.get("Source_Type"), record.get("Project"),
            record.get("Topic"), record.get("Language"), record.get("Publisher"), record.get("DOI"),
            record.get("Canonical_URL") or record.get("URL"), reason,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="9C0006")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in enumerate([16, 70, 22, 35, 45, 15, 30, 35, 60, 65], 1):
        ws.column_dimensions[chr(64 + col)].width = width
    wb.save(REPORTS / "out_of_scope_removed.xlsx")


def filter_excel(id_map: dict[str, str], data: dict, removed_count: int, change_id: str) -> None:
    path = ROOT / "Data.xlsx"
    wb = load_workbook(path)
    id_sheets = (
        "ALL_SOURCES", "TIMELINE", "ACADEMIC", "THESES", "OFFICIAL_LEGAL", "PARLIAMENT",
        "NEWS_MEDIA", "SOCIAL_MEDIA", "PDF_INDEX", "NEW_SOURCES",
    )
    for sheet_name in id_sheets:
        ws = wb[sheet_name]
        for row in range(ws.max_row, 1, -1):
            old_id = ws.cell(row, 1).value
            if old_id not in id_map:
                ws.delete_rows(row, 1)
            else:
                ws.cell(row, 1).value = id_map[old_id]

    for ws in (wb["ALL_SOURCES"], wb["ACADEMIC"], wb["NEW_SOURCES"]):
        for table in ws.tables.values():
            table.ref = f"A1:AK{ws.max_row}"

    wb["README"]["A1"] = "Water Transfer Isfahan — TARGET BASINS ONLY R6.1"
    wb["README"]["B2"] = len(data["sources"])
    wb["README"]["B3"] = f"S001–{data['meta']['lastSourceId']} (continuous, no gaps)"
    wb["README"]["B4"] = f"Through {TODAY}"
    wb["DASHBOARD"]["A1"] = "ZAYANDEH-RUD & KARUN TARGET BASINS R6.1"
    labels = {wb["DASHBOARD"].cell(row, 1).value: row for row in range(2, wb["DASHBOARD"].max_row + 1)}
    for label, value in data["meta"]["dashboard"].items():
        if label in labels:
            wb["DASHBOARD"].cell(labels[label], 2).value = value
        elif label == "Out-of-scope sources removed":
            append_styled(wb["DASHBOARD"], [label, value, "Removed by strict target-basin scope cleanup"])
    project_counts = {item["Project"]: item["Source_Count"] for item in data.get("projects") or []}
    for row in range(2, wb["PROJECTS"].max_row + 1):
        name = wb["PROJECTS"].cell(row, 1).value
        if name in project_counts:
            wb["PROJECTS"].cell(row, 2).value = project_counts[name]
    append_styled(wb["SEARCH_LOG"], [
        "پاک‌سازی دامنه: فقط ایران و حوضه‌های زاینده‌رود/کارون و طرح‌های متصل",
        "فارسی/English", "Dataset-wide scope audit", "زاینده‌رود / کارون", "همه منابع",
        TODAY, len(data["sources"]), f"{removed_count} منبع خارج از دامنه حذف و جداگانه گزارش شد.",
    ])
    append_styled(wb["CHANGE_LOG"], [
        change_id, TODAY, "TARGET-SCOPE CLEANUP", f"S001–{data['meta']['lastSourceId']}",
        "User required all records to concern the Zayandeh-Rud/Karun basin system in Iran",
        f"Removed {removed_count} out-of-scope records; renumbered retained records continuously.",
    ])
    wb.save(path)


def update_offline(data: dict) -> None:
    path = ROOT / "app.js"
    text = path.read_text(encoding="utf-8")
    marker = "const OFFLINE_FALLBACK = "
    start = text.index(marker) + len(marker)
    end = text.index(";\n\nconst $", start)
    packed = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    path.write_text(text[:start] + packed + text[end:], encoding="utf-8", newline="\n")


def main() -> None:
    data = json.loads((ROOT / "data.json").read_text(encoding="utf-8"))
    kept, removed = [], []
    for record in data["sources"]:
        keep, reason = in_scope(record)
        (kept if keep else removed).append(record if keep else (record, reason))

    print(json.dumps({
        "before": len(data["sources"]), "keep": len(kept), "remove": len(removed),
        "remove_preexisting": sum(int(x[0]["Source_ID"][1:]) <= 533 for x in removed),
        "remove_deep_search": sum(int(x[0]["Source_ID"][1:]) >= 534 for x in removed),
        "top_removed_projects": Counter(x[0].get("Project") for x in removed).most_common(12),
    }, ensure_ascii=False, indent=2))
    if "--dry-run" in sys.argv:
        return

    write_removed_report(removed)
    id_map = {record["Source_ID"]: f"S{index:03d}" for index, record in enumerate(kept, 1)}
    for record in kept:
        record["Source_ID"] = id_map[record["Source_ID"]]
    data["sources"] = kept
    data["timeline"] = [{key: record.get(key) for key in (
        "Source_ID", "Date_Gregorian", "Date_Persian", "Project", "Title", "Source_Type", "Publisher", "Reliability", "URL", "Year"
    )} for record in kept]
    for name, ids in data["collections"].items():
        data["collections"][name] = [id_map[source_id] for source_id in ids if source_id in id_map]
    data["pdfIndex"] = [{key: record.get(key) for key in (
        "Source_ID", "Title", "Publisher", "PDF_Status", "PDF_URL", "Local_PDF", "SHA256", "Reliability"
    )} for record in kept if record.get("PDF_URL")]

    for project in data.get("projects") or []:
        name = normalize(project.get("Project"))
        project["Source_Count"] = sum(name in normalize(record.get("Project")) for record in kept)
    data["meta"]["generatedAt"] = TODAY
    data["meta"]["totalSources"] = len(kept)
    data["meta"]["lastSourceId"] = kept[-1]["Source_ID"]
    dashboard = data["meta"]["dashboard"]
    dashboard["Total sources"] = len(kept)
    dashboard["Persian-language sources"] = sum("فارسی" in (x.get("Language") or "") for x in kept)
    dashboard["English-language sources"] = sum("English" in (x.get("Language") or "") or "انگلیسی" in (x.get("Language") or "") for x in kept)
    for grade in "ABCD":
        dashboard[f"Reliability {grade}"] = sum(x.get("Reliability") == grade for x in kept)
    dashboard["Sources with direct PDF links"] = sum(x.get("PDF_Status") == "Direct Link Available" for x in kept)
    dashboard["Last Source ID"] = kept[-1]["Source_ID"]
    retained_deep = [x for x in kept if "مرور عمیق ادبیات" in (x.get("Notes") or "")]
    dashboard["Deep literature articles added"] = len(retained_deep)
    dashboard["Deep literature Persian articles"] = sum(x.get("Language") == "فارسی" for x in retained_deep)
    dashboard["Deep literature English articles"] = sum(x.get("Language") == "English" for x in retained_deep)
    dashboard["Out-of-scope sources removed"] = int(dashboard.get("Out-of-scope sources removed") or 0) + len(removed)
    data["searchLog"].append({
        "Query": "پاک‌سازی دامنه: فقط ایران، زاینده‌رود، کارون و طرح‌های متصل",
        "Language": "فارسی/English", "Database / Search Engine": "Dataset-wide scope audit",
        "Project": "زاینده‌رود / کارون", "Source Type": "All", "Search Date": TODAY,
        "Number of useful results": len(kept), "Notes": f"{len(removed)} منبع خارج از دامنه حذف شد.",
    })
    change_numbers = [int(match.group(1)) for row in data["changeLog"] if (match := re.fullmatch(r"C(\d+)", str(row.get("Change_ID") or "")))]
    change_id = f"C{max(change_numbers, default=0) + 1:03d}"
    data["changeLog"].append({
        "Change_ID": change_id, "Date": TODAY, "Action": "TARGET-SCOPE CLEANUP",
        "Source_ID / Item": f"S001–{kept[-1]['Source_ID']}",
        "Reason": "Restrict all records to Iran's Zayandeh-Rud/Karun basin system",
        "Details": f"Removed {len(removed)} out-of-scope records and renumbered retained records.",
    })
    save_json(ROOT / "data.json", data)
    filter_excel(id_map, data, len(removed), change_id)
    update_offline(data)
    print(f"Removed {len(removed)} records; retained {len(kept)} as S001–{kept[-1]['Source_ID']}.")


if __name__ == "__main__":
    main()
