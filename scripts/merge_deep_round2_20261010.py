"""Stage and validate the reviewed 27-source import; preserve prior records."""
import collections, copy, hashlib, json, pathlib, re, shutil, sys, urllib.parse
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/merge_deep_round2_2026_10_10'
SRC = ROOT/'reports/deep_round2_2026_10_10'
TODAY = '2026-10-10'
SHEETS = {'academicIds':'ACADEMIC','legalIds':'OFFICIAL_LEGAL','parliamentIds':'PARLIAMENT','newsIds':'NEWS_MEDIA'}
def read(p): return json.loads(p.read_text('utf-8-sig'))
def save(p,d): p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n','utf-8')
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(v): return re.sub(r'[^\w]','',str(v or '').casefold().replace('ي','ی').replace('ك','ک'))
def urlkey(v):
    u=urllib.parse.urlsplit(urllib.parse.unquote(v or ''))
    host=u.netloc.lower().removeprefix('www.').replace('tasnimnews.com','tasnimnews.ir')
    m=re.search(r'/news/(?:\d{4}/\d{2}/\d{2}/)?(\d+)',u.path)
    if m:return host+'/news/'+m[1]
    query=urllib.parse.urlencode(sorted((k,v) for k,v in urllib.parse.parse_qsl(u.query) if k not in ['lang','utm_source','utm_medium']))
    return host+u.path.rstrip('/')+('?' + query if query else '')
def dates(c):
    cid=c['ID'];raw=c.get('Date','').translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789'))
    special={'A001':(None,'1391',1391),'A095':('2020',None,2020),'A110':('2015','1394',1394),
             'A122':('2009',None,2009),'R01':('2016-06-01',None,2016),'R02':(None,'1399',1399),
             'N01':('2018-10-05','1397/07/13',1397),'N02':('2014-06-14','1393/03/24',1393),
             'I01':('2014-06-03','1393/03/13',1393),'I02':('2018-09-03','1397/06/12',1397),
             'D01':('2024-08-28','1403/06/07',1403),'D02':('2014-05-20','1393/02/30',1393)}
    if cid in special:return special[cid]
    date=re.match(r'20\d{2}(?:[-/]\d{2}(?:[-/]\d{2})?)?',raw)[0].replace('/','-')
    return date,None,int(date[:4])

def prepare():
    OUT.mkdir(exist_ok=True);(OUT/'backup').mkdir(exist_ok=True);(OUT/'staged').mkdir(exist_ok=True)
    data=read(ROOT/'data.json');before=copy.deepcopy(data)
    assert len(data['sources'])==985 and data['meta']['lastSourceId']=='S985', 'Already merged or baseline changed'
    for name in ['data.json','app.js','Data.xlsx']:
        backup=OUT/'backup'/name
        if backup.exists():assert digest(backup)==digest(ROOT/name),'Baseline differs from saved backup'
        else:shutil.copy2(ROOT/name,backup)
    package=read(SRC/'candidates.json');candidates=[c for c in package['Candidates'] if c['Status']=='candidate_verified']
    assert len(candidates)==27
    fields=list(data['sources'][0]);added=[];mapping=[]
    for n,c in enumerate(candidates,986):
        r=dict.fromkeys(fields);cid=c['ID'];cat=c['Category'];fa=bool(re.search('[\u0600-\u06ff]',c['Title']))
        dg,dp,year=dates(c)
        project=c['Project'].replace('چشمه لنگان','چشمه‌لنگان').replace('بن–بروجن','بن-بروجن')
        if cid=='A122':project='انتقال کارون؛ فلات مرکزی'
        if cid=='R01':project='زاینده‌رود؛ بالادست سد'
        if cid=='I02':project='خرسان ۳؛ سبزکوه'
        types={'articles':'مقاله علمی','news':'خبر','interviews':'مصاحبه','official_reports':'گزارش رسمی/نهادی (زمینه‌ای)','documents':'سند رسمی/آگهی مناقصه' if cid=='D01' else 'نامه رسمی/متن بازنشرشده در رسانه'}
        local=(SRC/c['Local_File']).relative_to(ROOT).as_posix() if c.get('Local_File') else None
        pdf=c.get('PDF_URL') if local else None
        if cid=='A063':pdf='https://www.dpceonline.it/index.php/dpceonline/article/download/2480/2838'
        if cid in ['R01','R02','D01']:pdf=c['URL']
        groups=['academicIds'] if cat=='articles' else ['newsIds'] if cat in ['news','interviews'] else ['legalIds']
        if cid in ['N02','D02']:groups.append('parliamentIds')
        if cid=='D02':groups.append('newsIds')
        note=f"ورود در {TODAY} از بسته deep_round2_2026_10_10، شناسه {cid}. تاریخ منبع: {c.get('Date')}. {c['Notes']} سطح دسترسی: {c['Access']}"
        if c.get('Relation')=='context':note+=' ارتباط زمینه‌ای؛ این رکورد ارزیابی محیط‌زیستی یا مجوز طرح انتقال نیست.'
        if c.get('PDF_URL') and not local:note+=' پیوند پیشنهادی PDF به دریافت موفق نرسید: '+c['PDF_URL']
        r.update(Source_ID=f'S{n:03}',Title=c['Title'],Title_EN=c.get('Title_EN') or (c['Title'] if not fa else None),
                 Date_Gregorian=dg,Date_Persian=dp,Year=year,Language='فارسی' if fa else 'English',
                 Source_Type=types[cat],Publisher=c.get('Publisher'),Author='؛ '.join(c.get('Authors',[])) or None,
                 Interviewee=c.get('Person') if cat=='interviews' else None,Person=c.get('Person'),
                 Project=project,Topic=c.get('Title_FA') or c['Title'],Summary=c['Summary'],
                 URL=c['URL'],Canonical_URL=c['URL'],DOI=c.get('DOI') or None,Document_ID=c.get('Document_ID'),
                 Reliability='B' if cat in ['news','interviews'] or cid=='D02' else 'A',
                 Primary_or_Secondary='Secondary' if cat=='news' or cid=='D02' else 'Primary',
                 Full_Text='Yes' if local or cat!='articles' else 'Partial',
                 PDF_Status='Direct Link Available' if pdf else 'Landing Page Only',PDF_URL=pdf,Local_PDF=local,
                 SHA256=digest(ROOT/local) if local else None,Notes=note)
        if cid=='A122':r['Source_Type']='مقاله کنفرانسی'
        if cid=='A063':r['Primary_or_Secondary']='Secondary'
        if cid=='N01':r['Author']='مسعود بربر'
        if cid=='N02':r['Person']='حسن کامران';r['Position']='نماینده مجلس'
        if cid=='I01':r['Position']='رئیس شورای شهر فریدون‌شهر';r['Province']='اصفهان'
        if cid=='D01':r['Organization']='شرکت سهامی آب منطقه‌ای اصفهان';r['Location']='بیرگان، شهرستان کوهرنگ'
        for old in data['sources']+added:
            assert norm(r['Title']) not in {norm(old.get('Title')),norm(old.get('Title_EN'))},('duplicate title',cid,old['Source_ID'])
            assert urlkey(r['URL']) not in {urlkey(old.get('URL')),urlkey(old.get('Canonical_URL')),urlkey(old.get('PDF_URL'))},('duplicate URL',cid,old['Source_ID'])
            assert not r['DOI'] or norm(r['DOI'])!=norm(old.get('DOI')),('duplicate DOI',cid,old['Source_ID'])
        if local:
            assert (ROOT/local).read_bytes().startswith(b'%PDF')
            assert r['SHA256']==c['SHA256']
        added.append(r);mapping.append({'candidate_id':cid,'source_id':r['Source_ID'],'category':cat,'collections':groups})
        for g in groups:data['collections'][g].append(r['Source_ID'])
    data['sources'].extend(added)
    for r in added:
        data['timeline'].append({k:r[k] for k in before['timeline'][0]})
        if r['PDF_URL']:data['pdfIndex'].append({k:r[k] for k in before['pdfIndex'][0]})
    for p in ['سبزکوه','چغاخور','گندمان','بیرگان']:
        if p not in {x['Project'] for x in data['projects']}:data['projects'].append({'Project':p,'Source_Count':0})
    for p in data['projects']:p['Source_Count']=sum(p['Project'] in (r.get('Project') or '') for r in data['sources'])
    search={'Query':'بهشت‌آباد، کوهرنگ، گلاب، سبزکوه، خرسان، پشندگان، بیرگان؛ مقاله، مصاحبه، خبر، گزارش و سند',
            'Language':'فارسی / English','Database / Search Engine':'OpenAlex; Crossref; publisher pages; media archives; official repositories',
            'Project':'کارون و زاینده‌رود و طرح‌های مرتبط','Source Type':'مقاله / خبر / مصاحبه / گزارش / سند','Search Date':TODAY,
            'Number of useful results':27,'Notes':'۱۹ مقاله، ۲ خبر، ۲ مصاحبه، ۲ گزارش زمینه‌ای و ۲ سند وارد شد؛ ۶ سرنخ تاییدنشده کنار ماند. جزئیات و ۵۲ مدخل جست‌وجو: reports/deep_round2_2026_10_10/.'}
    change={'Change_ID':f"C{max(int(c['Change_ID'][1:]) for c in data['changeLog'])+1:03}",'Date':TODAY,
            'Action':'DEEP RESEARCH ROUND 2 MERGE','Source_ID / Item':'S986–S1012',
            'Reason':'User-authorized import of 27 reviewed research candidates',
            'Details':'27 added; 985 existing records preserved; 6 pending leads excluded; 10 local PDFs indexed; research overlap and access limits retained.'}
    data['searchLog'].append(search);data['changeLog'].append(change)
    data['meta'].update(generatedAt=TODAY,totalSources=1012,lastSourceId='S1012')
    dash=data['meta']['dashboard']
    dash.update({'Total sources':1012,'Last Source ID':'S1012',
        'Persian-language sources':sum('فارسی' in (s.get('Language') or '') for s in data['sources']),
        'English-language sources':sum('English' in (s.get('Language') or '') or 'انگلیسی' in (s.get('Language') or '') for s in data['sources']),
        'Sources with direct PDF links':sum(s.get('PDF_Status')=='Direct Link Available' for s in data['sources']),
        'Locally retrieved PDFs':sum(bool(s.get('Local_PDF')) for s in data['sources'])})
    for grade in 'ABCD':dash[f'Reliability {grade}']=sum(s.get('Reliability')==grade for s in data['sources'])
    save(OUT/'staged/data.json',data)
    app=(ROOT/'app.js').read_text('utf-8');marker='const OFFLINE_FALLBACK = ';start=app.index(marker)+len(marker);end=app.index(';\n\nconst $',start)
    patched=app[:start]+json.dumps(data,ensure_ascii=False,separators=(',',':'))+app[end:]
    (OUT/'staged/app.js').write_text(patched,'utf-8')
    wb=load_workbook(ROOT/'Data.xlsx');operations={}
    def append(sheet,records):
        if not records:return
        headers=[c.value for c in wb[sheet][1]]
        operations[sheet]={'start_row':wb[sheet].max_row+1,'columns':len(headers),'rows':[[r.get(k) for k in headers] for r in records]}
    for sheet in ['ALL_SOURCES','NEW_SOURCES']:append(sheet,added)
    for group,sheet in SHEETS.items():append(sheet,[r for r,m in zip(added,mapping) if group in m['collections']])
    append('TIMELINE',added);append('PDF_INDEX',[r for r in added if r['PDF_URL']]);append('SEARCH_LOG',[search]);append('CHANGE_LOG',[change])
    cells=[]
    for sheet,records in [('DASHBOARD',[{'k':k,'v':v} for k,v in dash.items()]),('PROJECTS',[{'k':p['Project'],'v':p['Source_Count']} for p in data['projects']])]:
        ws=wb[sheet];lookup={ws.cell(i,1).value:i for i in range(2,ws.max_row+1)};last=ws.max_row
        for item in records:
            if item['k'] not in lookup:last+=1;lookup[item['k']]=last
            row=lookup[item['k']]
            for col,value in [('A',item['k']),('B',item['v'])]:
                if row>ws.max_row or ws[f'{col}{row}'].value!=value:cells.append({'sheet':sheet,'cell':f'{col}{row}','value':value})
    for cell,value in [('A1','Water Transfer Isfahan — DEEP RESEARCH UPDATE'),('B2',1012),('B3','S001–S1012 (continuous, no gaps)'),('B4',f'Through {TODAY}')]:cells.append({'sheet':'README','cell':cell,'value':value})
    cells.append({'sheet':'DASHBOARD','cell':'C10','value':'Local_PDF exists and SHA-256 is recorded'})
    save(OUT/'workbook_changes.json',{'append':operations,'cells':cells})
    save(OUT/'manifest.json',{'date':TODAY,'before_count':985,'after_count':1012,'mapping':mapping,
        'excluded_pending':[c['ID'] for c in package['Candidates'] if c['Status']=='needs_verification'],
        'baseline_sha256':{n:digest(ROOT/n) for n in ['data.json','app.js','Data.xlsx']}})
    print('STAGED',len(added),'sources; workbook operations',list(operations))

def validate_and_commit():
    data=read(OUT/'staged/data.json');before=read(OUT/'backup/data.json');manifest=read(OUT/'manifest.json');changes=read(OUT/'workbook_changes.json')
    assert data['sources'][:985]==before['sources']
    assert [s['Source_ID'] for s in data['sources']]==[f'S{i:03}' for i in range(1,1013)]
    ids={s['Source_ID'] for s in data['sources']};by={s['Source_ID']:s for s in data['sources']}
    assert data['meta']['totalSources']==len(ids)==1012
    for group,members in data['collections'].items():assert len(members)==len(set(members)) and set(members)<=ids
    for key in ['timeline','pdfIndex']:
        assert len(data[key])==len({s['Source_ID'] for s in data[key]})
        assert data[key][:len(before[key])]==before[key]
    app=(OUT/'staged/app.js').read_text('utf-8');oldapp=(OUT/'backup/app.js').read_text('utf-8');mark='const OFFLINE_FALLBACK = ';start=app.index(mark)+len(mark);end=app.index(';\n\nconst $',start)
    assert json.loads(app[start:end])==data
    assert app[end:]==oldapp[oldapp.index(';\n\nconst $'):]
    wb=load_workbook(OUT/'staged/Data.xlsx');old=load_workbook(OUT/'backup/Data.xlsx')
    assert wb.sheetnames==old.sheetnames
    allowed={(c['sheet'],c['cell']):c['value'] for c in changes['cells']}
    old_values_checked=0;style_differences=[]
    for ws in old:
        new=wb[ws.title]
        for row in ws:
            for c in row:
                expected=allowed.get((ws.title,c.coordinate),c.value)
                assert new[c.coordinate].value==expected,(ws.title,c.coordinate,expected,new[c.coordinate].value)
                old_values_checked+=1
                # Compare visible style properties, independent of internal style IDs.
                nc=new[c.coordinate]
                if c.has_style and (copy.copy(c.font)!=copy.copy(nc.font) or copy.copy(c.fill)!=copy.copy(nc.fill) or copy.copy(c.alignment)!=copy.copy(nc.alignment) or c.number_format!=nc.number_format):
                    style_differences.append([ws.title,c.coordinate])
                assert (c.hyperlink.target if c.hyperlink else None)==(nc.hyperlink.target if nc.hyperlink else None),(ws.title,c.coordinate,'hyperlink')
        assert ws.freeze_panes==new.freeze_panes
        assert set(ws.tables)==set(new.tables)
        assert list(ws.merged_cells.ranges)==list(new.merged_cells.ranges)
    for sheet,op in changes['append'].items():
        ws=wb[sheet]
        actual=[[ws.cell(r,c).value for c in range(1,op['columns']+1)] for r in range(op['start_row'],op['start_row']+len(op['rows']))]
        assert actual==op['rows'],('appended rows mismatch',sheet)
    assert wb['ALL_SOURCES'].max_row==1013
    assert wb['ALL_SOURCES'].tables['AllSourcesClean'].ref=='A1:AK1013'
    for c in changes['cells']:assert wb[c['sheet']][c['cell']].value==c['value']
    assert not style_differences,('changed existing styles',len(style_differences),style_differences[:5])
    for s in data['sources'][985:]:
        if s['Local_PDF']:assert digest(ROOT/s['Local_PDF'])==s['SHA256']
        for key in ['Title','URL','Publisher','Project','Source_Type']:assert s[key]
    wb.close();old.close()
    for name,sha in manifest['baseline_sha256'].items():assert digest(ROOT/name)==sha,'Source changed during preparation'
    # All consistency checks passed; backups exist before canonical replacements.
    for name in ['Data.xlsx','data.json','app.js']:shutil.copy2(OUT/'staged'/name,ROOT/name)
    result={'status':'PASS','date':TODAY,'before':985,'after':1012,'added':27,'first_id':'S986','last_id':'S1012',
            'unchanged_existing_records':985,'existing_excel_cells_checked':old_values_checked,'existing_style_changes':0,
            'local_pdfs_added':10,'pending_excluded':manifest['excluded_pending'],
            'json_offline_equal':True,'new_excel_json_rows_equal':True,
            'canonical_sha256':{n:digest(ROOT/n) for n in ['data.json','app.js','Data.xlsx']}}
    save(OUT/'validation.json',result);print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__':
    (prepare if len(sys.argv)<2 or sys.argv[1]=='prepare' else validate_and_commit)()
