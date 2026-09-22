"""Check the Mandegan import against the curated manifest and the pre-import git data."""
import hashlib
import io
import json
import subprocess
from pathlib import Path
import jdatetime
from openpyxl import load_workbook
from merge_mandegan import urlkey, norm, SHEETS, FIELDS, TIMELINE_XLSX_FIELDS, PDF_FIELDS

ROOT = Path(__file__).resolve().parents[1]

def main():
    data = json.loads((ROOT/'data.json').read_text(encoding='utf-8'))
    curated = json.loads((ROOT/'reports/mandegan_curated.json').read_text(encoding='utf-8'))
    audit = json.loads((ROOT/'reports/mandegan_merge_audit.json').read_text(encoding='utf-8'))
    by_id = {s['Source_ID']:s for s in data['sources']}
    assert list(by_id) == [f'S{i:03d}' for i in range(1,986)]
    assert len(data['sources']) == data['meta']['totalSources'] == 985
    assert data['meta']['lastSourceId'] == 'S985'
    original = json.loads(subprocess.check_output(['git','show','HEAD:data.json'],cwd=ROOT).decode('utf-8'))
    assert len(original['sources']) == 958, 'Validation is tied to the pre-import revision'
    changed = [s['Source_ID'] for s in original['sources'] if s != by_id[s['Source_ID']]]
    assert changed == ['S216','S249']
    for old in original['sources']:
        expected = dict(old)
        expected.update(curated['enrich_existing'].get(old['Source_ID'],{}))
        assert by_id[old['Source_ID']] == expected
    app = (ROOT/'app.js').read_text(encoding='utf-8')
    marker='const OFFLINE_FALLBACK = '
    start=app.index(marker)+len(marker); end=app.index(';\n\nconst $',start)
    assert json.loads(app[start:end]) == data
    original_app = subprocess.check_output(['git','show','HEAD:app.js'],cwd=ROOT).decode('utf-8')
    assert app[end:] == original_app[original_app.index(';\n\nconst $'):]
    wb=load_workbook(ROOT/'Data.xlsx',read_only=True,data_only=True)
    all_rows={row[0]:dict(zip(FIELDS,row)) for row in list(wb['ALL_SOURCES'].values)[1:]}
    baseline_wb=load_workbook(io.BytesIO(subprocess.check_output(['git','show','HEAD:Data.xlsx'],cwd=ROOT)),read_only=True,data_only=True)
    baseline_rows={row[0]:dict(zip(FIELDS,row)) for row in list(baseline_wb['ALL_SOURCES'].values)[1:]}
    for sid, row in baseline_rows.items():
        assert all_rows[sid] == (by_id[sid] if sid in changed else row)
    existing_differences=[{'Source_ID':sid,'field':f,'json':s[f],'xlsx':all_rows[sid][f]}
        for sid,s in by_id.items() for f in FIELDS if s[f]!=all_rows[sid][f]]
    assert all(int(r['Source_ID'][1:])<959 for r in existing_differences)
    ids=set(by_id)
    for key, members in data['collections'].items():
        assert len(members)==len(set(members)) and set(members)<=ids
    for item in curated['rows']:
        s=item['record']; sid=s['Source_ID']
        assert by_id[sid]==s
        assert all_rows[sid]==s
        assert set(s)==set(FIELDS)
        for group in item['collections']:
            assert sid in data['collections'][group]
            rows={r[0]:dict(zip(FIELDS,r)) for r in list(wb[SHEETS[group]].values)[1:]}
            assert rows[sid]==s
        for sheet, fields in [('NEW_SOURCES',FIELDS),('TIMELINE',TIMELINE_XLSX_FIELDS)]+([('PDF_INDEX',PDF_FIELDS)] if s.get('PDF_URL') else []):
            rows={r[0]:dict(zip(fields,r)) for r in list(wb[sheet].values)[1:]}
            assert rows[sid]=={f:s.get(f) for f in fields}
        assert len([x for x in data['sources'] if urlkey(x.get('Canonical_URL') or x.get('URL'))==urlkey(s['URL'])])==1
        assert len([x for x in data['sources'] if norm(x['Title'])==norm(s['Title'])])==1
        if len(s['Date_Persian'])==10:
            assert jdatetime.date(*map(int,s['Date_Persian'].split('/'))).togregorian().isoformat()==s['Date_Gregorian']
        assert s['Date_Gregorian'] <= '2026-09-22'
        if s.get('Local_PDF'):
            raw=(ROOT/s['Local_PDF']).read_bytes()
            assert raw.startswith(b'%PDF') and hashlib.sha256(raw).hexdigest()==s['SHA256']
    for name in ['timeline','pdfIndex']:
        rows=data[name]
        assert len(rows)==len({r['Source_ID'] for r in rows})
        for r in rows:
            if r['Source_ID'] in {i['record']['Source_ID'] for i in curated['rows']} | set(changed):
                assert r=={f:by_id[r['Source_ID']].get(f) for f in r}
    assert len(data['timeline'])==985
    for sheet,key in [('SEARCH_LOG','searchLog'),('CHANGE_LOG','changeLog')]:
        rows=list(wb[sheet].values); headers=rows[0]
        count=62 if key=='searchLog' else 1
        assert [dict(zip(headers,r)) for r in rows[-count:]]==data[key][-count:]
    dashboard={r[0]:r[1] for r in list(wb['DASHBOARD'].values)[1:]}
    assert all(dashboard[k]==v for k,v in data['meta']['dashboard'].items())
    projects={r[0]:r[1] for r in list(wb['PROJECTS'].values)[1:]}
    assert all(projects[r['Project']]==r['Source_Count'] for r in data['projects'])
    doc_manifest=json.loads((ROOT/'reports/mandegan_documents/manifest.json').read_text(encoding='utf-8'))
    for doc in doc_manifest['documents']:
        assert hashlib.sha256((ROOT/doc['path']).read_bytes()).hexdigest()==doc['sha256']
    wb.close()
    result={'status':'PASS','sources':985,'added':27,'enriched':changed,'unchanged_prior_records':956,
        'imported_excel_json_rows_equal':True,'complete_json_offline_equal':True,
        'preexisting_excel_json_differences_preserved':existing_differences,
        'new_url_title_duplicates':0,'document_hashes_verified':len(doc_manifest['documents']),
        'partial_access_ids':[r['record']['Source_ID'] for r in curated['rows'] if r['record']['Full_Text']=='Partial'],
        'query_count':audit['query_count'],'raw_unique_result_urls':audit['raw_unique_result_urls']}
    (ROOT/'reports/mandegan_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({**result,'preexisting_excel_json_differences_preserved':len(existing_differences)},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
