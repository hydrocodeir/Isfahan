#!/usr/bin/env python3
"""Merge verified sources and generate human-readable audit reports."""

from __future__ import annotations

import csv
import json
import re
from copy import copy
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
TODAY = "2026-09-13"
TODAY_FA = "1405/06/22"

SOURCE_FIELDS = [
    "Source_ID", "Title", "Title_EN", "Date_Gregorian", "Date_Persian", "Year", "Language",
    "Source_Type", "Publisher", "Author", "Interviewee", "Person", "Position", "Organization",
    "Province", "Location", "Project", "Basin_Source", "Basin_Destination", "Topic",
    "Position_Stance", "Summary", "Key_Claim", "Quote", "URL", "Canonical_URL", "Archive_URL",
    "DOI", "Document_ID", "Reliability", "Primary_or_Secondary", "Full_Text", "PDF_Status",
    "PDF_URL", "Local_PDF", "SHA256", "Notes",
]

REASON_FA = {
    "Supplied DOI is not registered and no independent matching source was found.":
        "DOI ارائه‌شده در Crossref ثبت نشده و منبع هم‌عنوان مستقلی پیدا نشد.",
    "Supplied links are unavailable and no independent matching source was found.":
        "لینک‌های ارائه‌شده در دسترس نبودند و منبع هم‌عنوان مستقلی پیدا نشد.",
    "Supplied page is a soft-404 and no independent matching source was found.":
        "صفحه، ۴۰۴ نرم بود و منبع هم‌عنوان مستقلی پیدا نشد.",
    "Available pages/search results do not substantively match the supplied source title.":
        "صفحه‌ها یا نتایج جست‌وجو با عنوان/اطلاعات منبع تطبیق محتوایی کافی نداشتند.",
    "Indexed title found, but the recovered source URL was unavailable and no live alternative link was found.":
        "عنوان در نمایه جست‌وجو دیده شد، اما لینک بازیابی‌شده باز نشد و جایگزین زنده‌ای پیدا نشد.",
}


def load_json(path: Path):
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def save_json(path: Path, value) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def normalize(value: str | None) -> str:
    return " ".join(
        re.sub(r"[^\w\s]", " ", (value or "").replace("‌", " ").replace("ي", "ی").replace("ك", "ک").lower()).split()
    )


def clean_title(title: str, fallback: str) -> str:
    value = " ".join((title or fallback).split())
    value = re.sub(r"^مقاله\s+", "", value)
    value = re.sub(r"\s*[-|]\s*(ایرنا|ایمنا|خبرآنلاین|جامعه اندیشکده ها|سلامت نیوز).*$", "", value)
    value = re.sub(r"\s*-\s*خبرگزاری مهر.*$", "", value)
    value = re.sub(r"-موقع مجلات نور التخصصية$", "", value)
    if value.endswith("...") or len(value) < max(20, len(fallback) // 2):
        return fallback
    return value.strip(" -|")


def page_metadata(url: str) -> dict:
    try:
        response = requests.get(
            url,
            timeout=(10, 25),
            headers={"User-Agent": "Mozilla/5.0 Chrome/140 Safari/537.36", "Accept-Language": "fa,en;q=0.8"},
        )
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        def meta_values(*names):
            values = []
            for name in names:
                for tag in soup.find_all("meta", attrs={"name": name}) + soup.find_all("meta", attrs={"property": name}):
                    if tag.get("content"):
                        values.append(" ".join(tag["content"].split()))
            return values
        title_values = meta_values("citation_title", "og:title", "twitter:title")
        authors = meta_values("citation_author", "author")
        dates = meta_values("citation_publication_date", "article:published_time", "date", "datePublished")
        sites = meta_values("og:site_name", "citation_journal_title")
        for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
            text = script.string or script.get_text(" ", strip=True)
            match = re.search(r'"datePublished"\s*:\s*"([^"]+)"', text)
            if match:
                dates.append(match.group(1))
        date_value = None
        for candidate in dates:
            match = re.search(r"(19|20)\d{2}(?:[-/]\d{1,2}(?:[-/]\d{1,2})?)?", candidate)
            if match:
                date_value = match.group(0).replace("/", "-")
                break
        return {
            "title": title_values[0] if title_values else None,
            "authors": list(dict.fromkeys(authors)),
            "date": date_value,
            "site": sites[0] if sites else None,
            "body": " ".join(soup.get_text(" ", strip=True).split())[:20000],
        }
    except Exception:
        return {}


def evidence_domain(url: str) -> str:
    host = urlparse(url).netloc.lower()
    return host[4:] if host.startswith("www.") else host


def cleaned_record(source: dict, result: dict, new_id: str) -> dict:
    evidence = result["evidence"]
    url = evidence["url"]
    doi_check = result.get("doi_check") or {}
    metadata = page_metadata(url)
    if evidence.get("evidence_type") == "crossref_doi":
        metadata = {
            "title": doi_check.get("title"),
            "authors": doi_check.get("authors") or [],
            "date": str(doi_check.get("year")) if doi_check.get("year") else None,
            "site": doi_check.get("publisher"),
            "body": "",
        }
    actual_title = clean_title(metadata.get("title") or evidence.get("title") or "", source["Title"])
    domain = evidence_domain(url)
    scholarly = bool(
        evidence.get("evidence_type") == "crossref_doi"
        or domain in {"civilica.com", "noormags.ir", "link.springer.com", "doi.org"}
        or domain.endswith(".ac.ir")
    )
    publisher = source.get("Publisher")
    if domain == "civilica.com":
        publisher = "سیویلیکا (CIVILICA)"
    elif domain == "ijce.sbu.ac.ir":
        publisher = "Iranian Journal of Chemical Engineering (IJCE)"
    elif metadata.get("site") and scholarly:
        publisher = metadata["site"]

    verified_date = metadata.get("date")
    persian_date = None
    year_value = None
    input_gregorian = source.get("Date_Gregorian") or ""
    if verified_date:
        if len(verified_date) == 4:
            gregorian_date = verified_date
            year_value = int(verified_date)
        else:
            gregorian_date = verified_date
            if input_gregorian.startswith(verified_date[:7]):
                persian_date = source.get("Date_Persian")
                year_value = source.get("Year")
    else:
        gregorian_date = None

    body_norm = normalize(metadata.get("body"))
    authors = metadata.get("authors") or []
    author_value = "، ".join(authors) if authors else None
    # Keep named people only if the live body explicitly contains them.
    def verified_person(field):
        value = source.get(field)
        return value if value and normalize(value) in body_norm else None

    document_id = None
    civilica_match = re.search(r"civilica\.com/doc/(\d+)", url)
    if civilica_match:
        document_id = "CIVILICA-" + civilica_match.group(1)
    elif doi_check.get("status") == "registered":
        document_id = "DOI:" + source["DOI"]

    record = {field: None for field in SOURCE_FIELDS}
    record.update(
        {
            "Source_ID": new_id,
            "Title": actual_title,
            "Title_EN": actual_title if re.search(r"[A-Za-z]", actual_title) and not re.search(r"[\u0600-\u06ff]", actual_title) else None,
            "Date_Gregorian": gregorian_date,
            "Date_Persian": persian_date,
            "Year": year_value,
            "Language": source.get("Language"),
            "Source_Type": source.get("Source_Type"),
            "Publisher": publisher,
            "Author": author_value,
            "Interviewee": verified_person("Interviewee"),
            "Person": verified_person("Person"),
            "Position": None,
            "Organization": None,
            "Province": source.get("Province"),
            "Location": source.get("Location"),
            "Project": source.get("Project"),
            "Basin_Source": source.get("Basin_Source"),
            "Basin_Destination": source.get("Basin_Destination"),
            "Topic": source.get("Project") or actual_title,
            "Position_Stance": None,
            "Summary": f"وجود منبع و پیوند آن در ممیزی اینترنتی {TODAY_FA} تأیید شد. عنوان ثبت‌شده از صفحه/نمایه معتبر منبع گرفته شده است.",
            "Key_Claim": None,
            "Quote": None,
            "URL": url,
            "Canonical_URL": url,
            "Archive_URL": None,
            "DOI": source.get("DOI") if doi_check.get("status") == "registered" else None,
            "Document_ID": document_id,
            "Reliability": "A" if scholarly else "B",
            "Primary_or_Secondary": "Primary" if scholarly else "Secondary",
            "Full_Text": "Yes" if metadata.get("body") else "Unknown",
            "PDF_Status": "Landing Page Only",
            "PDF_URL": None,
            "Local_PDF": None,
            "SHA256": None,
            "Notes": (
                f"اعتبارسنجی وب {TODAY_FA}؛ شناسه ورودی {source['Source_ID']}. "
                "عنوان و لینک تأیید/اصلاح شد؛ نقل‌قول، ادعای کلیدی و اطلاعات توصیفیِ فاقد شاهد حذف شد."
            ),
        }
    )
    return record


def classify(record: dict) -> set[str]:
    domain = evidence_domain(record["Canonical_URL"])
    result = set()
    if (
        domain in {"civilica.com", "noormags.ir", "link.springer.com", "doi.org"}
        or domain.endswith(".ac.ir")
        or "مقاله" in (record.get("Source_Type") or "")
        or "Academic" in (record.get("Source_Type") or "")
    ):
        result.add("academicIds")
    media_domains = {
        "irna.ir", "imna.ir", "yjc.ir", "mehrnews.com", "khabaronline.ir", "salamatnews.com",
        "rokna.net", "iranthinktanks.com", "rooyeshnovin.org",
    }
    if domain in media_domains or any(token in (record.get("Source_Type") or "") for token in ("خبر", "مصاحبه", "گزارش")):
        result.add("newsIds")
    return result


def append_styled(ws, values) -> None:
    source_row = ws.max_row
    ws.append(values)
    target_row = ws.max_row
    if source_row >= 2:
        for col in range(1, len(values) + 1):
            src = ws.cell(source_row, col)
            dst = ws.cell(target_row, col)
            if src.has_style:
                dst._style = copy(src._style)
            dst.number_format = src.number_format
            dst.alignment = copy(src.alignment)
    for col, value in enumerate(values, 1):
        if isinstance(value, str) and value.startswith("http"):
            ws.cell(target_row, col).hyperlink = value
            ws.cell(target_row, col).style = "Hyperlink"


def update_excel(records: list[dict], merged: dict, id_map: dict[str, str]) -> None:
    path = ROOT / "Data.xlsx"
    wb = load_workbook(path)
    for record in records:
        values = [record.get(field) for field in SOURCE_FIELDS]
        append_styled(wb["ALL_SOURCES"], values)
        append_styled(wb["NEW_SOURCES"], values)
        classes = classify(record)
        if "academicIds" in classes:
            append_styled(wb["ACADEMIC"], values)
        if "thesisIds" in classes:
            append_styled(wb["THESES"], values)
        if "legalIds" in classes:
            append_styled(wb["OFFICIAL_LEGAL"], values)
        if "parliamentIds" in classes:
            append_styled(wb["PARLIAMENT"], values)
        if "newsIds" in classes:
            append_styled(wb["NEWS_MEDIA"], values)
        if "socialIds" in classes:
            append_styled(wb["SOCIAL_MEDIA"], values)
        timeline = [record.get(key) for key in (
            "Source_ID", "Date_Gregorian", "Date_Persian", "Project", "Title", "Source_Type", "Publisher", "Reliability", "URL"
        )]
        append_styled(wb["TIMELINE"], timeline)

    ws = wb["ALL_SOURCES"]
    if "AllSourcesClean" in ws.tables:
        ws.tables["AllSourcesClean"].ref = f"A1:AK{ws.max_row}"

    wb["README"]["A1"] = "Water Transfer Isfahan — VALIDATED WEB AUDIT R5"
    wb["README"]["B2"] = merged["meta"]["totalSources"]
    wb["README"]["B3"] = f"S001–{merged['meta']['lastSourceId']} (continuous, no gaps)"
    wb["README"]["B4"] = f"Through {TODAY}"

    dashboard = wb["DASHBOARD"]
    dashboard["A1"] = "VALIDATED WEB AUDIT R5 DASHBOARD"
    dash_values = merged["meta"]["dashboard"]
    labels = {dashboard.cell(row, 1).value: row for row in range(2, dashboard.max_row + 1)}
    for label, value in dash_values.items():
        if label in labels:
            dashboard.cell(labels[label], 2).value = value
    for label, value, definition in (
        ("Validated Antigravity sources", len(records), "Accepted after live URL/DOI/title verification"),
        ("Rejected Antigravity candidates", 488 - len(records), "Not merged; see validation report"),
    ):
        append_styled(dashboard, [label, value, definition])

    projects = wb["PROJECTS"]
    project_counts = {row["Project"]: row["Source_Count"] for row in merged["projects"]}
    for row in range(2, projects.max_row + 1):
        name = projects.cell(row, 1).value
        if name in project_counts:
            projects.cell(row, 2).value = project_counts[name]

    append_styled(
        wb["SEARCH_LOG"],
        [
            "ممیزی تک‌به‌تک ۴۸۸ رکورد new_data.json",
            "فارسی/انگلیسی",
            "Direct HTTP, Crossref, Yahoo index, targeted web search",
            "همه پروژه‌ها",
            "Mixed",
            TODAY,
            len(records),
            f"{len(records)} تأیید و {488-len(records)} رد؛ گزارش کامل در reports/.",
        ],
    )
    append_styled(
        wb["CHANGE_LOG"],
        [
            "C030", TODAY, "VALIDATED MERGE", f"S513–S{512+len(records)}",
            "Internet verification of Antigravity candidates",
            f"Added {len(records)} verified sources; rejected {488-len(records)}; unsupported generated claims removed.",
        ],
    )
    wb.save(path)


def report_rows(results: list[dict], id_map: dict[str, str]) -> list[dict]:
    rows = []
    for result in results:
        evidence = result.get("evidence") or {}
        checks = result.get("url_checks") or []
        searches = result.get("search_check") or []
        statuses = sorted({search.get("status", "") for search in searches if search.get("status")})
        doi_status = (result.get("doi_check") or {}).get("status")
        rows.append(
            {
                "شناسه ورودی": result["original_source_id"],
                "شناسه افزوده‌شده": id_map.get(result["original_source_id"]),
                "نتیجه": "تأیید" if result["verdict"] == "accepted" else "رد",
                "عنوان ورودی": result.get("title"),
                "عنوان شاهد": evidence.get("title"),
                "ناشر ورودی": result.get("publisher"),
                "لینک ورودی": result.get("input_url"),
                "لینک معتبر/بهترین شاهد": evidence.get("url"),
                "نوع شاهد": evidence.get("evidence_type"),
                "امتیاز تطبیق": evidence.get("score"),
                "وضعیت DOI": doi_status,
                "نتیجه بررسی لینک‌ها": ", ".join(check.get("outcome", "") for check in checks),
                "وضعیت جست‌وجوی مستقل": ", ".join(statuses) if statuses else "اجرا نشد (شاهد مستقیم کافی بود)",
                "علت تصمیم": REASON_FA.get(result.get("reason"), result.get("reason")),
            }
        )
    return rows


def style_report_sheet(ws, rows: list[dict], table_name: str) -> None:
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in (7, 8):
            cell = row[col - 1]
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    widths = [15, 18, 10, 55, 55, 28, 50, 50, 18, 14, 15, 30, 28, 60]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    if rows:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)


def write_reports(results: list[dict], id_map: dict[str, str]) -> None:
    rows = report_rows(results, id_map)
    REPORTS.mkdir(exist_ok=True)
    headers = list(rows[0])
    with (REPORTS / "rejected_sources.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(row for row in rows if row["نتیجه"] == "رد")

    wb = Workbook()
    summary = wb.active
    summary.title = "خلاصه"
    summary.sheet_view.rightToLeft = True
    summary.append(["شاخص", "مقدار"])
    summary.append(["تعداد بررسی‌شده", len(rows)])
    summary.append(["تأیید و افزوده‌شده", sum(row["نتیجه"] == "تأیید" for row in rows)])
    summary.append(["ردشده", sum(row["نتیجه"] == "رد" for row in rows)])
    summary.append(["تاریخ ممیزی", TODAY_FA])
    summary.append(["قاعده پذیرش", "DOI ثبت‌شده یا صفحه/نمایه زنده با تطبیق قوی عنوان؛ صفحات جست‌وجو و فهرست پذیرفته نشدند."])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 100
    for cell in summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for title, verdict, table_name in (("تأییدشده", "تأیید", "AcceptedSources"), ("ردشده", "رد", "RejectedSources")):
        subset = [row for row in rows if row["نتیجه"] == verdict]
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in subset:
            ws.append([row[header] for header in headers])
        style_report_sheet(ws, subset, table_name)
    wb.save(REPORTS / "source_validation_report.xlsx")


def main() -> None:
    old = load_json(ROOT / "data.json")
    if len(old["sources"]) != 512 or old["sources"][-1]["Source_ID"] != "S512":
        raise SystemExit("Refusing to merge: data.json is not the expected S001–S512 baseline.")
    new_data = load_json(ROOT / "new_data.json")
    audit = load_json(REPORTS / "source_validation_results.json")
    source_map = {source["Source_ID"]: source for source in new_data["sources"]}
    accepted_results = [row for row in audit["results"] if row["verdict"] == "accepted"]
    id_map = {row["original_source_id"]: f"S{513 + index:03d}" for index, row in enumerate(accepted_results)}
    records = [
        cleaned_record(source_map[row["original_source_id"]], row, id_map[row["original_source_id"]])
        for row in accepted_results
    ]

    merged = old
    merged["sources"].extend(records)
    merged["meta"]["generatedAt"] = TODAY
    merged["meta"]["totalSources"] = len(merged["sources"])
    merged["meta"]["lastSourceId"] = records[-1]["Source_ID"]
    dashboard = merged["meta"]["dashboard"]
    dashboard["Total sources"] = len(merged["sources"])
    dashboard["Persian-language sources"] = sum("فارسی" in (row.get("Language") or "") for row in merged["sources"])
    dashboard["English-language sources"] = sum(
        "English" in (row.get("Language") or "") or "انگلیسی" in (row.get("Language") or "")
        for row in merged["sources"]
    )
    for grade in "ABCD":
        dashboard[f"Reliability {grade}"] = sum(row.get("Reliability") == grade for row in merged["sources"])
    dashboard["Sources with direct PDF links"] = sum(row.get("PDF_Status") == "Direct Link Available" for row in merged["sources"])
    dashboard["Locally retrieved PDFs"] = sum(row.get("PDF_Status") == "Retrieved" for row in merged["sources"])
    dashboard["Last Source ID"] = records[-1]["Source_ID"]
    dashboard["Validated Antigravity sources"] = len(records)
    dashboard["Rejected Antigravity candidates"] = len(audit["results"]) - len(records)

    for project in merged["projects"]:
        name = normalize(project["Project"])
        project["Source_Count"] += sum(name in normalize(record.get("Project")) for record in records)

    for record in records:
        merged["timeline"].append(
            {key: record.get(key) for key in (
                "Source_ID", "Date_Gregorian", "Date_Persian", "Project", "Title", "Source_Type", "Publisher", "Reliability", "URL"
            )} | {"Year": record.get("Year")}
        )
        for collection in classify(record):
            merged["collections"][collection].append(record["Source_ID"])
    merged["searchLog"].append(
        {
            "Query": "ممیزی تک‌به‌تک ۴۸۸ رکورد new_data.json",
            "Language": "فارسی/انگلیسی",
            "Database / Search Engine": "Direct HTTP, Crossref, Yahoo index, targeted web search",
            "Project": "همه پروژه‌ها",
            "Source Type": "Mixed",
            "Search Date": TODAY,
            "Number of useful results": len(records),
            "Notes": f"{len(records)} تأیید و {len(audit['results'])-len(records)} رد؛ گزارش کامل در reports/.",
        }
    )
    merged["changeLog"].append(
        {
            "Change_ID": "C030",
            "Date": TODAY,
            "Action": "VALIDATED MERGE",
            "Source_ID / Item": f"S513–{records[-1]['Source_ID']}",
            "Reason": "Internet verification of Antigravity candidates",
            "Details": f"Added {len(records)} verified sources; rejected {len(audit['results'])-len(records)}; unsupported generated claims removed.",
        }
    )

    save_json(ROOT / "data.json", merged)
    write_reports(audit["results"], id_map)
    update_excel(records, merged, id_map)
    print(f"Merged {len(records)} verified records as S513–{records[-1]['Source_ID']}.")
    print(f"Rejected {len(audit['results']) - len(records)} records; reports written under reports/.")


if __name__ == "__main__":
    main()
